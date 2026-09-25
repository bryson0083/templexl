"""會話式 API：``Book`` 與 ``Sheet``。

``render()``（api.py）是一次性 facade——檔案進、檔案出。``Book`` 則把
「載入 → 渲染 → 結構操作 → 存檔」拆成呼叫端可掌握的步驟，服務兩類
``render()`` 表達不了的情境： 

1. 渲染後仍需以 openpyxl 做結構後處理（合併儲存格、刪欄），且不希望
   多一次 load/save 循環。
2. 逐工作表填入不同資料（例如「複製 Sample 工作表 N 份、每份不同資料」）。

**邊界原則**：本模組只包 openpyxl「做不到或做不對」的事（渲染、Excel 等效
刪欄）。openpyxl 原生能正確完成的操作（合併、寫值、``copy_worksheet``、
刪除工作表）一律經 ``Sheet.ws`` / ``Book.workbook`` 逃生門以原生 API 執行，
本套件不提供其包裝。

渲染編排（載入檢查、管線執行、warnings 收集、原子寫入）集中於本模組，
``api.render()`` 為建於其上的糖衣，確保兩個入口共用同一份編排邏輯。
"""
from __future__ import annotations

import logging
import os
import re
import tempfile
import uuid
import zipfile
from typing import Any, Dict, Iterator, List, Optional, Union

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .context import RenderContext
from .core.block_manager import BlockManager
from .core.container import ContainerManager
from .core.drawing_keeper import (
    capture_unsupported_drawings,
    describe_kept_anchors,
    restore_drawings,
)
from .core.parser import TemplateParser
from .core.renderer import TemplateRenderer
from .core.structural_ops import delete_column
from .core.table_sync import expected_autofilter_ref
from .exceptions import (
    FileFormatError,
    RenderError,
    StructuralEditError,
    TemplateError,
    TemplateNotFoundError,
    TemplateResourceError,
)
from .models.container import Container
from .report import RenderReport
from .result import BookRenderResult, SheetRenderResult

logger = logging.getLogger(__name__)

# 模板 zip 資源上限（解壓炸彈防護）。正常業務模板遠低於此；
# 如有特殊需求可於載入前調整這兩個模組常數。
MAX_DECOMPRESSED_SIZE = 2 * 1024 * 1024 * 1024  # 解壓後總大小上限：2GB
MAX_COMPRESSION_RATIO = 200  # 整體壓縮比上限（一般 xlsx 遠低於 20:1）

_TAG_PATTERN = re.compile(r"#?\{\{.*?\}\}")


class Sheet:
    """單一工作表的會話包裝。

    僅提供「templexl 才做得到」的操作（渲染、Excel 等效刪欄）；其餘一律
    經 :attr:`ws` 以 openpyxl 原生 API 操作。

    本物件為輕量包裝，每次自 :attr:`Book.sheets` 取用時建立；相等性
    以底層 worksheet 為準，因此重複取用得到等價（``==``）的物件。
    """

    __slots__ = ("_book", "_worksheet")

    def __init__(self, book: "Book", worksheet: Worksheet):
        self._book = book
        self._worksheet = worksheet

    @property
    def ws(self) -> Worksheet:
        """底層 openpyxl ``Worksheet``（逃生門）。"""
        return self._worksheet

    @property
    def title(self) -> str:
        """工作表名稱。"""
        return self._worksheet.title

    @property
    def book(self) -> "Book":
        """所屬的 :class:`Book`。"""
        return self._book

    def render(
        self,
        data: Optional[Dict[str, Any]] = None,
        *,
        escape_formulas: bool = True,
    ) -> SheetRenderResult:
        """只渲染本工作表（記憶體內，不落地）。

        掃描、標籤映射與渲染三階段皆限縮於本工作表，其他工作表的儲存格
        與標籤完全不被讀寫——這使「複製同一份模板工作表 N 份、每份填入
        不同資料」成為可能。

        本工作表已無任何標籤時（通常代表已渲染過）為 no-op：不改動儲存格、
        不拋例外，僅以 logger 發出 warning。

        Args:
            data: 渲染資料；鍵為標籤名稱，值為純量或 pandas DataFrame。
            escape_formulas: 是否中和 ``data`` 中危險前綴字串（防公式注入），
                預設開啟；語意與 ``render()`` 相同。

        Returns:
            SheetRenderResult: 含 ``sheet_name`` 與本工作表範圍的 ``warnings``。

        Raises:
            RenderError: 渲染過程發生錯誤。
        """
        workbook = self._book._require_open()
        worksheet = self._worksheet
        sheet_name = worksheet.title
        data = dict(data) if data else {}

        try:
            if not _has_tags(worksheet):
                logger.warning(
                    "工作表 '%s' 沒有可渲染的標籤（可能已渲染過），本次渲染為 no-op",
                    sheet_name,
                )
                return SheetRenderResult(sheet_name=sheet_name, warnings=[])

            render_context = RenderContext(
                process_id=str(uuid.uuid4()),
                template_path=self._book.template_path,
                output_path="",
                data=data,
            )
            _execute_render_pipeline(
                workbook, render_context, escape_formulas, sheet_names=[sheet_name]
            )
            _merge_shift_tracking(workbook, self._book._shift_tracking,
                                  render_context.shift_tracking)
            return SheetRenderResult(
                sheet_name=sheet_name,
                warnings=_collect_worksheet_warnings(worksheet),
            )
        except TemplateError:
            raise
        except Exception as e:
            raise RenderError(f"渲染過程發生未預期錯誤: {str(e)}", sheet_name=sheet_name)

    def delete_columns(self, idx: int) -> None:
        """刪除第 ``idx`` 欄（1-based），結果與 Excel「刪除整欄」等效。

        openpyxl 的 ``delete_cols`` 只搬動儲存格；本方法一併修正合併範圍、
        欄寬與列印範圍。工作表存在無法保證正確調整的特徵（公式、圖片、
        圖表、Excel 表格、條件格式、資料驗證、模板形狀、openpyxl 會丟棄的
        圖片如 EMF／WMF）時拒絕執行並拋出例外，不產出靜默損壞的結果。
        頁首頁尾圖片不綁欄位，不阻擋刪欄。

        刪除多欄請**由大到小**逐次呼叫，使座標全程以原始欄位計算：

        >>> for col in sorted(del_cols, reverse=True):   # doctest: +SKIP
        ...     sheet.delete_columns(col)

        Args:
            idx: 欲刪除的欄索引（1-based，A 欄為 1）。

        Raises:
            ValueError: ``idx`` 小於 1。
            StructuralEditError: 工作表存在無法保證正確調整的特徵。
        """
        self._book._require_open()

        # 保留內容（模板形狀、openpyxl 會丟棄的圖片）屬 Book 狀態，
        # delete_column() 只拿得到 openpyxl 工作表看不見，故檢查放在此處、
        # 委派前執行（design 決策 8）；頁首頁尾 VML 不綁欄位，不檢查。
        item = self._book._kept_drawings.get(self._worksheet)
        if item and item.anchors:
            kind, row, col = describe_kept_anchors(item)[0]
            coordinate = (f"{get_column_letter(col + 1)}{row + 1}"
                         if row is not None and col is not None else "未知位置")
            raise StructuralEditError(
                "刪除整欄",
                f"工作表含{kind}（錨點位置 {coordinate}），欄向錨點無法正確位移",
                self._worksheet.title,
            )

        delete_column(self._worksheet, idx)

    def __eq__(self, other: object) -> bool:
        if not isinstance(other, Sheet):
            return NotImplemented
        return self._book is other._book and self._worksheet is other._worksheet

    def __hash__(self) -> int:
        return hash((id(self._book), id(self._worksheet)))

    def __repr__(self) -> str:
        return f"<Sheet {self.title!r}>"


class SheetCollection:
    """:attr:`Book.sheets` 的存取器：支援索引與名稱兩種鍵。"""

    __slots__ = ("_book",)

    def __init__(self, book: "Book"):
        self._book = book

    def __getitem__(self, key: Union[int, str]) -> Sheet:
        workbook = self._book._require_open()
        if isinstance(key, bool) or not isinstance(key, (int, str)):
            raise TypeError(
                f"sheets 索引僅支援 int 或工作表名稱（str），收到 {type(key).__name__}"
            )
        if isinstance(key, int):
            try:
                worksheet = workbook.worksheets[key]
            except IndexError:
                raise IndexError(
                    f"工作表索引 {key} 超出範圍（共 {len(workbook.worksheets)} 張）"
                ) from None
        else:
            if key not in workbook.sheetnames:
                raise KeyError(
                    f"找不到工作表 '{key}'；現有工作表: {workbook.sheetnames}"
                )
            worksheet = workbook[key]
        return Sheet(self._book, worksheet)

    def __len__(self) -> int:
        return len(self._book._require_open().worksheets)

    def __iter__(self) -> Iterator[Sheet]:
        book = self._book
        for worksheet in book._require_open().worksheets:
            yield Sheet(book, worksheet)

    def __contains__(self, key: object) -> bool:
        if not isinstance(key, str):
            return False
        return key in self._book._require_open().sheetnames

    @property
    def names(self) -> List[str]:
        """工作表名稱清單（依實際順序）。"""
        return list(self._book._require_open().sheetnames)

    def __repr__(self) -> str:
        return f"<SheetCollection {self.names!r}>"


class Book:
    """以模板開啟的渲染會話。

    生命週期唯一歸屬本物件：載入於建構時發生，落地只在
    :meth:`save` 發生。``with`` 區塊結束**不會**自動存檔。

    Examples:
        >>> with Book("template.xlsx") as bk:            # doctest: +SKIP
        ...     bk.render({"oper_name": "王小明"})
        ...     bk.sheets[0].ws.merge_cells("A9:C9")     # openpyxl 原生
        ...     bk.save("output.xlsx")
    """

    __slots__ = (
        "_template_path", "_workbook", "_sheets", "_closed",
        "_kept_drawings", "_shift_tracking",
    )

    def __init__(self, template: Union[str, os.PathLike]):
        """開啟模板並載入至記憶體（不落地任何檔案）。

        Args:
            template: 模板檔路徑（``.xlsx`` 或 ``.xlsm``）。載入前會檢查 zip
                資源占用（解壓總量/壓縮比上限），防解壓炸彈。

        Raises:
            TemplateNotFoundError: 模板檔不存在。
            FileFormatError: 不支援的檔案格式或無效的 zip 結構。
            TemplateResourceError: 模板 zip 資源占用超過上限。
            RenderError: 載入過程發生未預期錯誤。
        """
        template = str(template)
        self._template_path = template
        self._workbook: Optional[Workbook] = None
        self._sheets = SheetCollection(self)
        self._closed = False
        # openpyxl 不模型化 <xdr:sp> 等圖形，載入即遺失；在此原樣留存，
        # 於 save() 寫出後回填（見 core.drawing_keeper）。
        self._kept_drawings = {}
        self._shift_tracking = {}

        try:
            if not os.path.exists(template):
                raise TemplateNotFoundError(template)
            if not _is_supported_format(template):
                raise FileFormatError(template)
            _check_template_resources(template)
            self._workbook = load_workbook(template)
            # 立即轉存為 Worksheet 物件鍵（design 決策 7）：擷取當下以模板
            # 工作表名稱為鍵，但存檔時工作表可能已經過逃生門改名——物件
            # 身分不受改名影響，才能在存檔時仍正確找到對應的保留內容。
            self._kept_drawings = {
                self._workbook[name]: item
                for name, item in capture_unsupported_drawings(template).items()
            }
        except TemplateError:
            raise
        except Exception as e:
            raise RenderError(f"載入模板發生未預期錯誤: {str(e)}")

    # ---- 基本屬性 -------------------------------------------------

    @property
    def template_path(self) -> str:
        """來源模板路徑。"""
        return self._template_path

    @property
    def workbook(self) -> Workbook:
        """底層 openpyxl ``Workbook``（逃生門）。"""
        return self._require_open()

    @property
    def sheets(self) -> SheetCollection:
        """工作表存取器；支援 ``bk.sheets[0]`` 與 ``bk.sheets["名稱"]``。"""
        return self._sheets

    def _require_open(self) -> Workbook:
        if self._closed or self._workbook is None:
            raise RenderError("Book 已關閉，無法再操作")
        return self._workbook

    # ---- 渲染與存檔 -----------------------------------------------

    def render(
        self,
        data: Optional[Dict[str, Any]] = None,
        *,
        escape_formulas: bool = True,
        with_report: bool = False,
    ) -> BookRenderResult:
        """以資料渲染整份工作簿（記憶體內，不落地）。

        Args:
            data: 渲染資料；鍵為標籤名稱，值為純量或 pandas DataFrame。
            escape_formulas: 是否中和 ``data`` 中危險前綴字串（防公式注入），
                預設開啟；語意與 ``render()`` 相同。
            with_report: 是否產生渲染報告（除錯/維護用），預設關閉。

        Returns:
            BookRenderResult: 含全簿 ``warnings``，以及（啟用時）``report``。

        Raises:
            RenderError: 渲染過程發生錯誤。
        """
        workbook = self._require_open()
        data = dict(data) if data else {}

        try:
            render_context = RenderContext(
                process_id=str(uuid.uuid4()),
                template_path=self._template_path,
                output_path="",
                data=data,
            )
            containers = _execute_render_pipeline(
                workbook, render_context, escape_formulas
            )
            _merge_shift_tracking(workbook, self._shift_tracking,
                                  render_context.shift_tracking)
            warnings = _collect_unresolved_tag_warnings(workbook)
            report = (
                _build_render_report(
                    containers, render_context, self._template_path, ""
                )
                if with_report
                else None
            )
            return BookRenderResult(warnings=warnings, report=report)
        except TemplateError:
            raise
        except Exception as e:
            raise RenderError(f"渲染過程發生未預期錯誤: {str(e)}")

    def save(self, output: Union[str, os.PathLike]) -> str:
        """將目前的工作簿狀態原子寫入輸出檔。

        寫入為原子操作（同目錄暫存檔 + rename），中途失敗不留下半寫入的檔案。
        存檔前執行表格範圍最終同步（``table.ref`` / ``autoFilter.ref``）。

        Args:
            output: 輸出檔路徑；已存在會被覆寫。

        Returns:
            str: 已寫出的輸出檔路徑。

        Raises:
            RenderError: 寫入過程發生未預期錯誤。
        """
        workbook = self._require_open()
        output = str(output)
        try:
            _final_table_autofilter_sync(workbook)
            _atomic_save(workbook, output, post_process=self._restore_drawings)
            return output
        except TemplateError:
            raise
        except Exception as e:
            raise RenderError(f"寫出輸出檔發生未預期錯誤: {str(e)}")

    def _restore_drawings(self, path: str) -> None:
        """回填模板圖形；於原子寫入的暫存檔上執行，故不破壞原子性。

        只回填仍在 ``workbook.worksheets`` 中的物件（design 決策 7）：
        逃生門移除工作表後，其 ``Worksheet`` 物件仍留在
        ``_kept_drawings``／``_shift_tracking`` 裡（未被主動清除）；若只靠
        「移除後 title 在輸出檔對不到部件」自然略過，一旦另一張工作表
        改名重用同一個 title，字串比對會把它誤判為同一張工作表而回填到
        錯誤位置——故須以物件身分（而非名稱字串）明確排除已移除者。
        """
        if not self._kept_drawings:
            return
        live = set(self._workbook.worksheets)
        captured = {ws: item for ws, item in self._kept_drawings.items() if ws in live}
        if not captured:
            return
        shift_tracking = {
            ws: shift for ws, shift in self._shift_tracking.items() if ws in live
        }
        restore_drawings(path, captured, shift_tracking)

    # ---- 生命週期 -------------------------------------------------

    def close(self) -> None:
        """釋放工作簿參照。不存檔；重複呼叫安全。"""
        self._workbook = None
        self._closed = True

    def __enter__(self) -> "Book":
        return self

    def __exit__(self, exc_type, exc_value, traceback) -> bool:
        # 刻意不自動存檔：落地必須顯式呼叫 save()。
        # 回傳 False（不吞例外），區塊內的例外原樣傳播。
        self.close()
        return False

    def __repr__(self) -> str:
        state = "closed" if self._closed else f"{len(self._workbook.worksheets)} sheets"
        return f"<Book {os.path.basename(self._template_path)!r} ({state})>"


# ======================================================================
# 載入 / 寫出
# ======================================================================


def _is_supported_format(file_path: str) -> bool:
    """
    檢查檔案格式是否支援

    Args:
        file_path: 檔案路徑

    Returns:
        bool: 是否支援
    """
    supported_extensions = ['.xlsx', '.xlsm']
    file_ext = os.path.splitext(file_path)[1].lower()
    return file_ext in supported_extensions


def _check_template_resources(template: str) -> None:
    """載入前檢查模板 zip 的資源占用，拒絕解壓炸彈與無效 zip。

    Raises:
        TemplateResourceError: 解壓後總大小或整體壓縮比超過上限。
        FileFormatError: 不是有效的 zip（xlsx/xlsm）結構。
    """
    try:
        with zipfile.ZipFile(template) as zf:
            total_uncompressed = sum(info.file_size for info in zf.infolist())
            total_compressed = max(
                1, sum(info.compress_size for info in zf.infolist())
            )
    except zipfile.BadZipFile:
        raise FileFormatError(template)

    if total_uncompressed > MAX_DECOMPRESSED_SIZE:
        raise TemplateResourceError(
            template,
            f"解壓後總大小 {total_uncompressed} bytes 超過上限 {MAX_DECOMPRESSED_SIZE}",
        )
    ratio = total_uncompressed / total_compressed
    if ratio > MAX_COMPRESSION_RATIO:
        raise TemplateResourceError(
            template,
            f"整體壓縮比 {ratio:.0f}:1 超過上限 {MAX_COMPRESSION_RATIO}:1",
        )


def _atomic_save(workbook: Workbook, output: str, post_process=None) -> None:
    """原子寫入輸出檔：先寫同目錄暫存檔，成功後以 os.replace 取代目標。

    中途失敗時清理暫存檔，目標路徑維持呼叫前狀態（不存在或保有舊檔）。

    Args:
        post_process: 可選的收尾處理，接受暫存檔路徑；在 ``os.replace``
            之前於暫存檔上執行，使 zip 層改寫（如模板圖形回填）同樣受
            原子性保護。
    """
    out_dir = os.path.dirname(os.path.abspath(output)) or "."
    fd, tmp_path = tempfile.mkstemp(suffix=".xlsx.tmp", dir=out_dir)
    os.close(fd)
    try:
        workbook.save(tmp_path)
        if post_process is not None:
            post_process(tmp_path)
        os.replace(tmp_path, output)
    except BaseException:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass
        raise


# ======================================================================
# warnings 收集
# ======================================================================


def _has_tags(worksheet: Worksheet) -> bool:
    """工作表是否仍存在任何模板標籤（決定 sheet 級渲染是否為 no-op）。"""
    for row in worksheet.iter_rows():
        for cell in row:
            value = cell.value
            # '{{' 是 _TAG_PATTERN 命中的必要條件；先以子字串排除絕大多數
            # 資料儲存格，避免對整張表逐格啟動正則引擎（大表為熱路徑）。
            if isinstance(value, str) and '{{' in value and _TAG_PATTERN.search(value):
                return True
    return False


def _collect_worksheet_warnings(worksheet: Worksheet) -> list:
    """掃描單一工作表，回報任何仍未被解析的標籤（非致命警告）。"""
    warnings: list = []
    for row in worksheet.iter_rows():
        for cell in row:
            value = cell.value
            # 同 _has_tags：'{{' 子字串前置過濾，正則只對候選儲存格執行。
            if isinstance(value, str) and '{{' in value and _TAG_PATTERN.search(value):
                warnings.append(
                    f"未解析的標籤 '{value}'"
                    f"（工作表 '{worksheet.title}' 儲存格 {cell.coordinate}）"
                )
    return warnings


def _collect_unresolved_tag_warnings(workbook: Workbook) -> list:
    """掃描輸出工作簿，回報任何仍未被解析的標籤（非致命警告）。"""
    warnings: list = []
    for worksheet in workbook.worksheets:
        warnings.extend(_collect_worksheet_warnings(worksheet))
    return warnings


# ======================================================================
# 渲染管線
# ======================================================================


def _merge_shift_tracking(workbook: Workbook, target: dict, shift_tracking: dict) -> None:
    """把 ``render_context.shift_tracking``（工作表名稱鍵）併入 ``target``
    （Worksheet 物件鍵，與 ``Book._kept_drawings`` 一致）。

    位移登記沿用渲染當下 ``container.sheet_name`` 的名稱字串（見
    ``core.block_manager``），在此轉存至物件鍵，避免「渲染後、存檔前
    改名」使位移資訊與保留內容的鍵對不上（design 決策 7）。
    """
    for name, shift in shift_tracking.items():
        target[workbook[name]] = shift


def _execute_render_pipeline(
    workbook: Workbook,
    render_context: 'RenderContext',
    escape_formulas: bool = True,
    sheet_names: Optional[List[str]] = None,
) -> list:
    """
    執行渲染管道

    Args:
        workbook: Excel工作簿
        render_context: 渲染上下文
        escape_formulas: 是否中和執行期資料值的公式前綴
        sheet_names: 限縮處理的工作表名稱；``None``（預設）處理全部工作表。
            掃描、映射與渲染三階段一併限縮，其他工作表完全不被讀寫。

    Returns:
        list: 容器清單
    """
    # 使用新的架構：統一的模板掃描與註冊表建立
    container_manager = ContainerManager()

    # 重要修正：先掃描原始標籤，保留所有標籤的完整資訊
    parser = TemplateParser()
    original_tags = parser.parse_template(workbook, sheet_names)

    # 建立容器
    containers = container_manager.create_containers(workbook, sheet_names=sheet_names)

    # 3. 區塊分類
    for container in containers:
        container_manager.classify_blocks(container)

    # 4. 建立標籤到物件的映射關係
    # 使用原始標籤資訊建立映射，確保相同名稱的多個標籤都能正確處理
    _build_tag_object_mapping_with_original_tags(containers, original_tags, render_context)

    # 5. 創建渲染器和區塊管理器
    renderer = TemplateRenderer(escape_formulas=escape_formulas)
    block_manager = BlockManager(escape_formulas=escape_formulas)

    # 6. 按區塊順序執行渲染
    for container in containers:
        _render_container(container, workbook, render_context, renderer, block_manager)

    return containers


def _render_container(
    container: Container,
    workbook: Workbook,
    render_context: RenderContext,
    renderer: 'TemplateRenderer',
    block_manager: 'BlockManager'
) -> None:
    """
    渲染單一容器

    Args:
        container: 容器物件
        workbook: Excel工作簿
        render_context: 渲染上下文
        renderer: 渲染器
        block_manager: 區塊管理器
    """
    logger.debug(f"DEBUG: 使用新的block分區渲染搬移機制處理容器: {container.sheet_name}")
    logger.debug(f"DEBUG: Calling block_manager.process_container_with_block_moving")

    try:
        # 使用新的block搬移機制
        block_manager.process_container_with_block_moving(container, workbook, render_context, renderer)
        logger.debug(f"DEBUG: Successfully called process_container_with_block_moving")
    except Exception as e:
        logger.debug(f"DEBUG: Error in process_container_with_block_moving: {e}")
        logger.debug("例外堆疊", exc_info=True)
        raise

    logger.debug(f"DEBUG: Finished processing container: {container.sheet_name}")


def _build_tag_object_mapping_with_original_tags(containers, original_tags, render_context):
    """
    使用原始標籤資訊建立標籤到物件的映射關係

    確保相同名稱的多個標籤都能正確對應到各自的物件

    Args:
        containers: 容器清單
        original_tags: 原始標籤清單（來自TemplateParser的掃描結果）
        render_context: 渲染上下文
    """
    from .models.base import ObjectType

    logger.debug(f"DEBUG_MAPPING: Building tag-object mapping with {len(original_tags)} original tags")

    # 建立位置到標籤清單的映射（支援同一位置多個標籤）
    position_to_tags = {}
    for tag in original_tags:
        key = (tag.sheet_name, tag.cell_position.row, tag.cell_position.col)
        if key not in position_to_tags:
            position_to_tags[key] = []
        position_to_tags[key].append(tag)
        logger.debug(f"DEBUG_MAPPING: Tag {tag.tag_name} at position {key}")

    # 為每個物件找到對應的標籤
    for container in containers:
        logger.debug(f"DEBUG_MAPPING: Processing container {container.sheet_name} with {len(container.objects)} objects")
        for obj in container.objects:
            # 處理標籤相關的物件（SIMPLE、TABLE和TABLE_OBJ類型）
            # TABLE_OBJ是標籤與表格物件綁定後的類型
            if obj.obj_type in [ObjectType.SIMPLE, ObjectType.TABLE, ObjectType.TABLE_OBJ]:
                key = (obj.sheet_name, obj.cell_position.row, obj.cell_position.col)
                logger.debug(f"DEBUG_MAPPING: Object {obj.obj_id} type={obj.obj_type} display_name={obj.display_name} at position {key}")
                # 對於TABLE_OBJ類型，需要特殊處理（因為綁定改變了位置）
                if obj.obj_type == ObjectType.TABLE_OBJ:
                    # TABLE_OBJ的display_name保留了原始標籤名
                    # 需要在所有標籤中查找匹配的標籤名
                    matched_tag = None
                    for pos_key, tags in position_to_tags.items():
                        if pos_key[0] == obj.sheet_name:  # 同一工作表
                            for tag in tags:
                                if tag.tag_name == obj.display_name:
                                    matched_tag = tag
                                    logger.debug(f"DEBUG_MAPPING: Found tag {tag.tag_name} for TABLE_OBJ {obj.obj_id}")
                                    break
                            if matched_tag:
                                break
                elif key in position_to_tags:
                    # 獲取該位置的所有標籤
                    candidate_tags = position_to_tags[key]

                    # 嘗試根據物件名稱匹配標籤
                    matched_tag = None
                    for tag in candidate_tags:
                        # 檢查物件ID是否包含標籤名稱
                        if tag.tag_name in obj.obj_id:
                            matched_tag = tag
                            break

                    # 如果沒有匹配到，使用第一個標籤作為備選
                    if not matched_tag and candidate_tags:
                        matched_tag = candidate_tags[0]
                else:
                    matched_tag = None

                if matched_tag:
                    render_context.add_tag_mapping(obj.obj_id, matched_tag)
                    logger.debug(f"DEBUG: 映射物件 {obj.obj_id} 到標籤 {matched_tag.tag_name} 在位置 ({matched_tag.cell_position.row}, {matched_tag.cell_position.col})")
                else:
                    logger.debug(f"DEBUG: 警告 - 無法找到物件 {obj.obj_id} 的匹配標籤 at position {key}")


def _final_table_autofilter_sync(workbook: Workbook) -> None:
    """
    在儲存前進行最終的表格範圍同步檢查

    確保所有表格的 ``autoFilter.ref`` 等於應有的篩選範圍（表頭列＋資料列，
    不含合計列）。本函式會掃過**全部**工作表的**全部**表格，與該表格是否
    被渲染無關——故篩選範圍一律以 ``expected_autofilter_ref`` 為準，未渲染
    的含合計列表格才不會在存檔時被推成含合計列。

    Args:
        workbook: Excel工作簿
    """
    logger.debug("DEBUG_FINAL_SYNC: 開始進行最終的表格範圍同步檢查")
    sync_count = 0

    for worksheet in workbook.worksheets:
        logger.debug(f"DEBUG_FINAL_SYNC: 檢查工作表 {worksheet.title}")

        for table_name in worksheet.tables:
            table = worksheet.tables[table_name]
            expected_ref = expected_autofilter_ref(table)

            if hasattr(table, 'autoFilter') and table.autoFilter:
                autofilter_ref = table.autoFilter.ref

                if expected_ref != autofilter_ref:
                    logger.debug(f"DEBUG_FINAL_SYNC: 發現不同步 - 表格 {table_name}")
                    logger.debug(f"  table.ref: {getattr(table, 'ref', '')}")
                    logger.debug(f"  autoFilter.ref: {autofilter_ref}")
                    logger.debug(f"  應有篩選範圍: {expected_ref}")
                    logger.debug(f"  正在同步...")

                    # 同步 autoFilter.ref
                    table.autoFilter.ref = expected_ref
                    sync_count += 1

                    logger.debug(f"  已同步為: {table.autoFilter.ref}")
                else:
                    logger.debug(f"DEBUG_FINAL_SYNC: 表格 {table_name} 範圍已同步: {autofilter_ref}")
            else:
                logger.debug(f"DEBUG_FINAL_SYNC: 表格 {table_name} 沒有 autoFilter")

    logger.debug(f"DEBUG_FINAL_SYNC: 同步檢查完成，共修正 {sync_count} 個表格")


def _build_render_report(
    containers: list,
    render_context: 'RenderContext',
    template_path: str,
    output_path: str,
) -> RenderReport:
    """建立渲染報告（記憶體物件，不寫入磁碟）。

    報告內容直接取自渲染後的容器狀態（單一真相來源），不另行重算物件落點，
    避免「報告與實際輸出漂移」的維護陷阱。欄名以 ``get_column_letter`` 轉換，
    正確支援超過 Z 欄的多字母欄名。

    Args:
        containers: 渲染後的容器清單。
        render_context: 渲染上下文（用於對應原始模板標籤）。
        template_path: 模板文件路徑。
        output_path: 輸出文件路徑。

    Returns:
        RenderReport: 物件清單與摘要。
    """
    worksheets: dict = {}
    total_objects = 0

    for container in containers:
        objects = []
        for obj in container.objects:
            coordinate = f"{get_column_letter(obj.cell_position.col)}{obj.cell_position.row}"
            tag = render_context.get_tag_for_object(obj.obj_id)
            objects.append({
                'obj_id': obj.obj_id,
                'display_name': obj.display_name,
                'obj_type': obj.obj_type.value if hasattr(obj.obj_type, 'value') else str(obj.obj_type),
                'sheet_name': obj.sheet_name,
                'having_header': obj.having_header,
                'is_multi_rows': obj.is_multi_rows,
                'block_id': obj.block_id,
                'template_tag': tag.tag_name if tag else None,
                # 渲染後實際狀態（單一真相來源：渲染流程已更新的容器物件）
                'position': {
                    'row': obj.cell_position.row,
                    'col': obj.cell_position.col,
                    'coordinate': coordinate,
                },
                'data_shape': {
                    'rows': obj.data_shape.rows,
                    'cols': obj.data_shape.cols,
                },
            })
            total_objects += 1

        worksheets[container.sheet_name] = {
            'container_id': container.container_id,
            'sheet_name': container.sheet_name,
            'total_objects': len(container.objects),
            'total_blocks': len(container.blocks),
            'objects': objects,
        }

    summary = {
        'total_worksheets': len(containers),
        'total_objects': total_objects,
    }

    return RenderReport(
        template_path=template_path,
        output_path=output_path,
        worksheets=worksheets,
        summary=summary,
    )
