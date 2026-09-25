"""
數據渲染器
"""
import logging
import re

logger = logging.getLogger(__name__)

from typing import Any, Union, Optional
import pandas as pd
from datetime import datetime, date

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import NamedStyle

from ..models.base import DataType, ObjectType
from ..models.tag import Tag
from ..models.container import Container
from ..models.objects import ObjectInfo
from ..exceptions import InvalidDataTypeError, RenderError
from openpyxl.utils import get_column_letter

from .table_sync import update_table_object_range
from .cell_ops import (
    translate_formula,
    apply_cell_style_dict,
    copy_cell_style_fast,
    copy_cell_style_dict,
    format_cell_ref,
    write_data_value,
)


class TemplateRenderer:
    """
    模板渲染器類別

    負責將數據渲染到Excel模板中
    """

    def __init__(self, escape_formulas: bool = True) -> None:
        self._escape_formulas = escape_formulas
        # 跨呼叫的表格原始位置快取（table_sync 動態起始列計算用）
        self._original_table_positions: dict = {}
    
    def render_simple_tag(
        self, 
        tag: Tag, 
        data: Any, 
        workbook: Workbook, 
        worksheet: Worksheet
    ) -> None:
        """
        渲染簡單變數標籤
        
        Args:
            tag: 標籤物件
            data: 要渲染的數據
            workbook: Excel工作簿
            worksheet: 工作表
            
        Raises:
            InvalidDataTypeError: 不支援的數據類型
            RenderError: 渲染過程錯誤
        """
        try:
            # 檢查是否為DataFrame
            if hasattr(data, 'shape') and hasattr(data, 'columns'):
                # DataFrame處理
                self._render_simple_dataframe(tag, data, worksheet)
            else:
                # 一般數據處理
                self._render_simple_data(tag, data, worksheet)
            
        except Exception as e:
            raise RenderError(
                f"渲染簡單標籤失敗: {str(e)}", 
                tag_name=tag.tag_name, 
                sheet_name=worksheet.title
            )
    
    def _render_simple_data(self, tag: Tag, data: Any, worksheet: Worksheet) -> None:
        """渲染簡單數據到單一儲存格"""
        # 取得目標儲存格
        target_cell = worksheet.cell(row=tag.cell_position.row, column=tag.cell_position.col)
        
        # 檢查是否為合併儲存格，如果是則跳過
        if isinstance(target_cell, MergedCell):
            return
        
        # 獲取原始值
        original_value = target_cell.value
        
        # 處理標籤替換
        if original_value and isinstance(original_value, str):
            # 使用正則表達式替換，忽略標籤中的空白
            # 構建支援任意空白的標籤模式：{{ tag_name }} 或 {{tag_name}}
            tag_pattern = rf"\{{\{{\s*{re.escape(tag.tag_name)}\s*\}}\}}"
            
            # 準備替換的數據值
            replacement_value = self._prepare_data_for_replacement(data, tag)

            # 執行標籤替換（使用正則表達式）
            new_value = re.sub(tag_pattern, str(replacement_value), original_value)
            if original_value.startswith('='):
                # 模板作者的公式（含標籤），保留公式語意
                target_cell.value = new_value
            else:
                # 純文字標籤替換結果視為資料值，套用公式中和
                write_data_value(target_cell, new_value, self._escape_formulas)
        else:
            # 如果儲存格沒有文本，直接設定數據值
            replacement_value = self._prepare_data_for_replacement(data, tag)
            write_data_value(
                target_cell, replacement_value, self._escape_formulas, str_fallback=False
            )
    
    def _render_simple_dataframe(self, tag: Tag, dataframe, worksheet: Worksheet) -> None:
        """渲染簡單DataFrame標籤到多個儲存格"""
        
        start_row = tag.cell_position.row
        start_col = tag.cell_position.col
        
        # Simple標籤邏輯：
        # 1. 有noheader條件：只渲染數據，不渲染header
        # 2. 無noheader條件：渲染header + 數據
        # 注意：simple模板標籤不需要考慮往下推移位置與cell風格樣式，直接替換數據即可
        
        has_noheader = tag.has_condition and tag.condition == "noheader"
        
        # 先清除原始標籤文字
        original_cell = worksheet.cell(row=start_row, column=start_col)
        if original_cell.value and isinstance(original_cell.value, str):
            # 清除包含標籤的文字
            tag_pattern = rf"\{{\{{\s*{re.escape(tag.tag_name)}\s*(\|\s*[^}}]*?)?\s*\}}\}}"
            if re.search(tag_pattern, original_cell.value):
                original_cell.value = None
        
        # 檢查是否為合併儲存格的一部分，如果是則需要處理合併範圍內的渲染
        merge_range = self._find_merged_range_for_cell(worksheet, start_row, start_col)
        if merge_range:
            # 如果標籤在合併儲存格中，使用合併儲存格的左上角作為開始位置
            start_row = merge_range.min_row
            start_col = merge_range.min_col
            logger.debug(f"DEBUG: 標籤 {tag.tag_name} 在合併儲存格中，調整起始位置為 ({start_row}, {start_col})")
        
        # 標準Simple標籤處理：根據是否有noheader決定是否渲染header
        logger.debug(f"DEBUG_RENDER: Calling _render_simple_full_dataframe for {tag.tag_name}, has_noheader={has_noheader}")
        self._render_simple_full_dataframe(dataframe, worksheet, start_row, start_col, has_noheader)
        
        # 設定標籤數據類型
        tag.data_type = DataType.DATAFRAME
    
    def _render_simple_full_dataframe(
        self,
        dataframe: pd.DataFrame,
        worksheet: Worksheet,
        start_row: int,
        start_col: int,
        has_noheader: bool = False
    ) -> None:
        """
        渲染簡單標籤的完整DataFrame（包括表頭和數據）
        
        Args:
            dataframe: DataFrame物件
            worksheet: 工作表
            start_row: 開始行
            start_col: 開始列
            has_noheader: 是否跳過表頭（noheader條件）
        """
        
        logger.debug(f"DEBUG_SIMPLE_FULL: Entered _render_simple_full_dataframe, start=({start_row},{start_col}), noheader={has_noheader}")
        logger.debug(f"DEBUG_SIMPLE_FULL: DataFrame shape={dataframe.shape}, columns={list(dataframe.columns)}")
        
        current_row = start_row
        
        # 渲染表頭（如果需要）
        if not has_noheader:
            for col_idx, col_name in enumerate(dataframe.columns):
                cell = worksheet.cell(row=current_row, column=start_col + col_idx)
                cell.value = str(col_name)
            current_row += 1
        
        # 渲染數據行
        for row_idx, (_, row_data) in enumerate(dataframe.iterrows()):
            logger.debug(f"DEBUG_UNIVERSAL_MERGE: 處理第{row_idx}行數據，行號{current_row + row_idx}")
            
            # 追蹤累積的列偏移量
            cumulative_offset = 0
            
            for col_idx, value in enumerate(row_data):
                # 計算目標列位置（考慮累積偏移）
                target_row = current_row + row_idx
                target_col = start_col + col_idx + cumulative_offset
                
                # 檢查是否是合併儲存格的主要儲存格（左上角）
                merge_range = self._find_merged_range_for_cell(worksheet, target_row, target_col)
                if merge_range and merge_range.min_row == target_row and merge_range.min_col == target_col:
                    # 這是合併範圍的主儲存格，計算跨度
                    merge_span = merge_range.max_col - merge_range.min_col
                    if merge_span > 0:
                        cumulative_offset += merge_span
                        logger.debug(f"DEBUG_UNIVERSAL_MERGE: col_idx={col_idx} 是合併主儲存格，跨度={merge_span}，累積偏移={cumulative_offset}")
                
                # 檢查目標儲存格是否可用，如果被合併占用則找下一個可用列
                elif self._is_cell_merged_and_not_top_left(worksheet, target_row, target_col):
                    old_col = target_col
                    target_col = self._find_next_available_column(worksheet, target_row, target_col)
                    additional_offset = target_col - old_col
                    cumulative_offset += additional_offset
                    logger.debug(f"DEBUG_UNIVERSAL_MERGE: col_idx={col_idx}被合併占用，從{old_col}跳到{target_col}，額外偏移={additional_offset}")
                
                cell = worksheet.cell(row=target_row, column=target_col)

                write_data_value(cell, value, self._escape_formulas)
    
    def _prepare_data_for_replacement(self, data: Any, tag: Tag) -> Any:
        """
        準備用於替換的數據值
        
        Args:
            data: 原始數據
            tag: 標籤物件
            
        Returns:
            Any: 準備好的數據值
        """
        if isinstance(data, str):
            tag.data_type = DataType.STRING
            return data
            
        elif isinstance(data, (int, float)):
            tag.data_type = DataType.NUMBER
            return data
            
        elif isinstance(data, (datetime, date)):
            tag.data_type = DataType.DATE
            return data
            
        else:
            # 嘗試轉換為字串
            tag.data_type = DataType.STRING
            return str(data)
    
    def render_table_tag(
        self, 
        tag: Tag, 
        dataframe: pd.DataFrame, 
        workbook: Workbook, 
        worksheet: Worksheet,
        obj_info: ObjectInfo
    ) -> int:
        """
        渲染表格數據標籤
        
        Args:
            tag: 標籤物件
            dataframe: 要渲染的DataFrame
            workbook: Excel工作簿
            worksheet: 工作表
            obj_info: 物件資訊
            
        Returns:
            int: 實際渲染的行數（用於計算位移）
            
        Raises:
            InvalidDataTypeError: 數據類型錯誤
            RenderError: 渲染過程錯誤
        """
        if not isinstance(dataframe, pd.DataFrame):
            raise InvalidDataTypeError(type(dataframe).__name__, tag.tag_name)
        
        try:
            start_row = tag.cell_position.row
            start_col = tag.cell_position.col
            template_position = (start_row, start_col)
            
            # 檢查是否有noheader條件
            skip_header = tag.has_condition and tag.condition == "noheader"
            
            # *** 移除重複的template rows複製邏輯 ***
            # 新的流程對齊機制中，template rows複製已在BlockManager._process_template_rows_by_render_order中處理
            # 這裡不再重複插入行，避免雙重插入問題
            data_rows_needed = len(dataframe)
            header_rows = 1 if (not skip_header and obj_info.having_header) else 0
            total_rows_needed = data_rows_needed + header_rows
            
            logger.debug(f"DEBUG_TABLE_TAG: render_table_tag for {tag.tag_name}, skip_header={skip_header}, data_rows={data_rows_needed}")
            logger.debug(f"DEBUG_TABLE_TAG: DataFrame columns={list(dataframe.columns)}")
            # 移除：self._copy_template_rows(worksheet, start_row, total_rows_needed - 1)
            
            # 渲染表頭（如果需要）
            current_row = start_row
            if not skip_header and obj_info.having_header:
                self._render_dataframe_header(dataframe, worksheet, current_row, start_col, template_position)
                current_row += 1
            
            # 渲染數據行
            rows_rendered = self._render_dataframe_data(dataframe, worksheet, current_row, start_col, template_position)
            
            # 更新物件資訊
            total_rows = rows_rendered
            if not skip_header and obj_info.having_header:
                total_rows += 1
            
            obj_info.data_shape.rows = total_rows
            obj_info.data_shape.cols = len(dataframe.columns)
            tag.data_type = DataType.DATAFRAME
            
            # 如果這是一個TABLE_OBJ類型，需要更新表格物件的範圍
            logger.debug(f"DEBUG_TABLE_CHECK: obj_info.obj_type = {obj_info.obj_type}, ObjectType.TABLE_OBJ = {ObjectType.TABLE_OBJ}")
            logger.debug(f"DEBUG_TABLE_CHECK: obj_info.obj_type == ObjectType.TABLE_OBJ: {obj_info.obj_type == ObjectType.TABLE_OBJ}")
            logger.debug(f"DEBUG_TABLE_CHECK: obj_info.display_name = {obj_info.display_name}")
            
            if obj_info.obj_type == ObjectType.TABLE_OBJ:
                logger.debug(f"DEBUG_TABLE_OBJ: render_table_tag handling TABLE_OBJ for {obj_info.display_name}")
                update_table_object_range(
                    worksheet, obj_info, dataframe, tag, start_row, start_col, total_rows,
                    self._original_table_positions,
                )
            else:
                logger.debug(f"DEBUG_TABLE_CHECK: 不是 TABLE_OBJ 類型，跳過範圍更新")
            
            return total_rows
            
        except Exception as e:
            raise RenderError(
                f"渲染表格標籤失敗: {str(e)}", 
                tag_name=tag.tag_name, 
                sheet_name=worksheet.title
            )
    
    def _render_dataframe_header(
        self, 
        dataframe: pd.DataFrame, 
        worksheet: Worksheet, 
        start_row: int, 
        start_col: int,
        template_cell_position: Optional[tuple] = None
    ) -> None:
        """
        渲染DataFrame表頭
        
        Args:
            dataframe: DataFrame物件
            worksheet: 工作表
            start_row: 開始行
            start_col: 開始列
            template_cell_position: 模板標籤所在位置 (row, col)，用於複製樣式
        """
        for col_idx, column_name in enumerate(dataframe.columns):
            cell = worksheet.cell(row=start_row, column=start_col + col_idx)
            
            # 檢查是否為合併儲存格，如果是則跳過設定值
            if not isinstance(cell, MergedCell):
                cell.value = column_name
            
            # 複製模板樣式到表頭
            if template_cell_position:
                template_cell = worksheet.cell(row=template_cell_position[0], column=template_cell_position[1])
                copy_cell_style_fast(template_cell, cell)
    
    def _render_dataframe_data(
        self, 
        dataframe: pd.DataFrame, 
        worksheet: Worksheet, 
        start_row: int, 
        start_col: int,
        template_cell_position: Optional[tuple] = None
    ) -> int:
        """
        渲染DataFrame數據
        
        Args:
            dataframe: DataFrame物件
            worksheet: 工作表
            start_row: 開始行
            start_col: 開始列
            template_cell_position: 模板標籤所在位置 (row, col)，用於複製樣式
            
        Returns:
            int: 渲染的行數
        """
        # 首先檢查模板行（如果有）中是否有公式
        formula_columns = {}
        if template_cell_position:
            template_row = template_cell_position[0]
            # 檢查整個模板行的所有相關列是否有公式
            # 需要檢查比DataFrame更多的列，因為可能有計算欄位
            max_check_cols = max(len(dataframe.columns) + 3, 10)  # 至少檢查10列
            for col_idx in range(max_check_cols):
                actual_col = start_col + col_idx
                check_cell = worksheet.cell(row=template_row, column=actual_col)
                if check_cell.value and isinstance(check_cell.value, str) and check_cell.value.startswith('='):
                    formula_columns[col_idx] = check_cell.value
        
        # 先偵測模板的合併模式
        logger.debug(f"DEBUG_MERGE_DETECT: Starting merge detection for worksheet '{worksheet.title}'")
        logger.debug(f"DEBUG_MERGE_DETECT: Template row: {template_cell_position[0] if template_cell_position else 'None'}")
        logger.debug(f"DEBUG_MERGE_DETECT: Start col: {start_col}, DataFrame columns: {len(dataframe.columns)}")
        
        merge_pattern = self._detect_template_merge_pattern(
            worksheet, 
            template_cell_position[0] if template_cell_position else start_row,
            start_col,
            len(dataframe.columns)
        ) if template_cell_position else {}
        
        logger.debug(f"DEBUG_MERGE_DETECT: Detected merge pattern: {merge_pattern}")
        
        for row_idx, (_, row_data) in enumerate(dataframe.iterrows()):
            current_row = start_row + row_idx
            
            # 根據偵測到的模式創建合併儲存格
            if merge_pattern:
                for col_offset, span in merge_pattern.items():
                    merge_start_col = start_col + col_offset
                    merge_end_col = merge_start_col + span - 1
                    merge_range = f"{get_column_letter(merge_start_col)}{current_row}:{get_column_letter(merge_end_col)}{current_row}"
                    try:
                        worksheet.merge_cells(merge_range)
                        logger.debug(f"DEBUG_MERGE_CREATE: Created merge range {merge_range}")
                    except Exception as e:
                        logger.debug(f"DEBUG_MERGE_CREATE: Failed to create merge {merge_range}: {e}")
            
            # 處理DataFrame中的數據列 - 使用智能列偏移
            cumulative_offset = 0
            
            
            for col_idx, value in enumerate(row_data):
                # 計算基礎目標列位置
                base_target_col = start_col + col_idx + cumulative_offset
                target_col = base_target_col
                
                
                # 檢查當前位置是否有合併儲存格
                merge_range = self._find_merged_range_for_cell(worksheet, current_row, target_col)
                
                if merge_range:
                    # 如果是合併範圍的起始儲存格
                    if merge_range.min_row == current_row and merge_range.min_col == target_col:
                        # 計算合併跨度並更新偏移
                        merge_span = merge_range.max_col - merge_range.min_col
                        if merge_span > 0:
                            cumulative_offset += merge_span
                    # 如果不是起始儲存格，需要跳過
                    elif self._is_cell_merged_and_not_top_left(worksheet, current_row, target_col):
                        # 找到下一個可用的列
                        old_col = target_col
                        target_col = self._find_next_available_column(worksheet, current_row, target_col)
                        additional_offset = target_col - old_col
                        cumulative_offset += additional_offset

                cell = worksheet.cell(row=current_row, column=target_col)
                
                # 檢查是否為合併儲存格，如果是則跳過
                if isinstance(cell, MergedCell):
                    continue
                
                
                # 檢查該欄位是否有公式
                if col_idx in formula_columns:
                    # 對於有公式的欄位，複製公式並調整引用
                    original_formula = formula_columns[col_idx]
                    if template_cell_position:
                        adjusted_formula = translate_formula(
                            original_formula,
                            template_cell_position[0],  # 模板行
                            start_col + col_idx,        # 公式列
                            current_row,                # 目標行
                            start_col + col_idx         # 目標列
                        )
                        cell.value = adjusted_formula
                else:
                    write_data_value(cell, value, self._escape_formulas)
                
                # 複製模板樣式到數據單元格
                if template_cell_position:
                    # 找到對應列的模板儲存格來複製樣式
                    template_col = cell.column  # 使用實際儲存格的列位置
                    template_cell = worksheet.cell(row=template_cell_position[0], column=template_col)
                    copy_cell_style_fast(template_cell, cell)
            
            # 處理超出DataFrame範圍的公式列（如計算欄位）
            for col_idx in formula_columns:
                if col_idx >= len(dataframe.columns):  # 超出DataFrame範圍的列
                    cell = worksheet.cell(row=current_row, column=start_col + col_idx)
                    
                    # 檢查是否為合併儲存格，如果是則跳過
                    if isinstance(cell, MergedCell):
                        continue
                    
                    # 複製公式並調整引用
                    original_formula = formula_columns[col_idx]
                    if template_cell_position:
                        adjusted_formula = translate_formula(
                            original_formula,
                            template_cell_position[0],  # 模板行
                            start_col + col_idx,        # 公式列
                            current_row,                # 目標行
                            start_col + col_idx         # 目標列
                        )
                        cell.value = adjusted_formula
                    
                    # 複製模板樣式到數據單元格
                    if template_cell_position:
                        # 找到對應列的模板儲存格來複製樣式
                        template_col = start_col + col_idx
                        template_cell = worksheet.cell(row=template_cell_position[0], column=template_col)
                        copy_cell_style_fast(template_cell, cell)
        
            # 為整個行範圍複製樣式，包括沒有數據或公式的列
            if template_cell_position:
                # 檢查模板行有多少列需要複製樣式（檢查到最多15列）
                max_template_cols = 15
                for col_idx in range(max_template_cols):
                    template_col = start_col + col_idx
                    template_cell = worksheet.cell(row=template_cell_position[0], column=template_col)
                    
                    # 如果模板儲存格有樣式（邊框、填充等），則複製到對應的數據行
                    if (template_cell.border and (template_cell.border.left.style or 
                                                 template_cell.border.right.style or 
                                                 template_cell.border.top.style or 
                                                 template_cell.border.bottom.style)) or \
                       (template_cell.fill and template_cell.fill.fill_type) or \
                       (template_cell.font and template_cell.font.name):
                        
                        target_cell = worksheet.cell(row=current_row, column=template_col)
                        
                        # 檢查是否為合併儲存格，如果是則跳過
                        if not isinstance(target_cell, MergedCell):
                            copy_cell_style_fast(template_cell, target_cell)
        
        return len(dataframe)
    
    def _copy_template_rows(self, worksheet: Worksheet, template_row: int, num_copies: int) -> None:
        """
        複製模板行（實現Excel的選取整行→插入複製行的操作）
        
        這個方法模擬Excel中的操作流程：
        1. 選取模板行的entire row（包含風格樣式、公式、合併儲存格）
        2. 在模板行下方插入指定數量的新行
        3. 將模板行的所有內容（樣式、公式、合併）複製到新插入的行
        4. 調整公式引用，保持相對引用正確
        5. 讓原有的其他內容往下推移
        
        Args:
            worksheet: 工作表
            template_row: 模板行號（包含標籤的行）
            num_copies: 需要複製的行數（通常是 DataFrame.shape[0] - 1）
        """
        if num_copies <= 0:
            return
            
        try:
            logger.debug(f"DEBUG: 開始複製模板行 {template_row}，需要複製 {num_copies} 行")
            
            # Step 1: 記錄模板行的所有資訊（在插入行之前）
            template_data = self._capture_template_row_data(worksheet, template_row)
            
            # Step 2: 在模板行下方插入新行（一次性插入所有需要的行）
            # 這會讓原有的其他內容往下推移
            worksheet.insert_rows(template_row + 1, num_copies)
            logger.debug(f"DEBUG: 已插入 {num_copies} 行在第 {template_row + 1} 行")
            
            # Step 3: 將模板行的內容複製到新插入的每一行
            for copy_index in range(num_copies):
                target_row = template_row + 1 + copy_index
                self._copy_template_row_to_target(worksheet, template_data, template_row, target_row)
                logger.debug(f"DEBUG: 已複製模板行到第 {target_row} 行")
                
        except Exception as e:
            logger.error(f"ERROR: 複製模板行失敗: {str(e)}")
            raise RenderError(f"複製模板行失敗: {str(e)}")
    
    def _capture_template_row_data(self, worksheet: Worksheet, template_row: int) -> dict:
        """
        捕獲模板行的所有資訊（樣式、值、公式、合併儲存格等）
        
        Args:
            worksheet: 工作表
            template_row: 模板行號
            
        Returns:
            dict: 包含模板行所有資訊的字典
        """
        template_data = {
            'cells': {},
            'merged_ranges': [],
            'max_col': worksheet.max_column
        }
        
        # 捕獲每個儲存格的資訊
        for col in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row=template_row, column=col)
            
            # 記錄儲存格資訊
            template_data['cells'][col] = {
                'value': cell.value,
                'data_type': cell.data_type,
                'style': copy_cell_style_dict(cell),
                'is_merged': False,
                'merged_range': None
            }
        
        # 捕獲合併儲存格範圍
        for merged_range in worksheet.merged_cells.ranges:
            if merged_range.min_row <= template_row <= merged_range.max_row:
                template_data['merged_ranges'].append({
                    'min_row': merged_range.min_row,
                    'max_row': merged_range.max_row,
                    'min_col': merged_range.min_col,
                    'max_col': merged_range.max_col,
                    'range_string': str(merged_range)
                })
                
                # 標記合併範圍內的儲存格
                for col in range(merged_range.min_col, merged_range.max_col + 1):
                    if col in template_data['cells']:
                        template_data['cells'][col]['is_merged'] = True
                        template_data['cells'][col]['merged_range'] = merged_range
        
        return template_data
    
    def _copy_template_row_to_target(self, worksheet: Worksheet, template_data: dict, template_row: int, target_row: int) -> None:
        """
        將模板行資料複製到目標行
        
        Args:
            worksheet: 工作表
            template_data: 模板行資料
            template_row: 模板行號
            target_row: 目標行號
        """
        
        # 複製每個儲存格
        for col, cell_data in template_data['cells'].items():
            target_cell = worksheet.cell(row=target_row, column=col)
            
            # 跳過合併儲存格（這些會在後續處理合併範圍時處理）
            if isinstance(target_cell, MergedCell):
                continue
            
            # 複製樣式
            apply_cell_style_dict(target_cell, cell_data['style'])
            
            # 複製值和公式
            if cell_data['value'] is not None:
                # 檢查是否為標籤字串（包含{{和}}）
                is_tag_cell = (isinstance(cell_data['value'], str) and 
                             '{{' in cell_data['value'] and '}}' in cell_data['value'])
                
                if not is_tag_cell:
                    # 處理公式
                    if cell_data['data_type'] == 'f' and isinstance(cell_data['value'], str):
                        # 調整公式中的引用
                        adjusted_formula = translate_formula(
                            cell_data['value'],
                            template_row,  # 原始行
                            col,           # 原始列
                            target_row,    # 目標行
                            col            # 目標列
                        )
                        target_cell.value = adjusted_formula
                    else:
                        # 普通值（但不是標籤）
                        target_cell.value = cell_data['value']
                # 如果是標籤儲存格，只複製樣式，不複製值
        
        # 處理合併儲存格
        self._copy_merged_ranges_to_target_row(worksheet, template_data['merged_ranges'], template_row, target_row)
    
    def _copy_merged_ranges_to_target_row(self, worksheet: Worksheet, merged_ranges: list, template_row: int, target_row: int) -> None:
        """
        將模板行的合併範圍複製到目標行
        只複製確實屬於模板行本身的合併儲存格，不複製其他行的合併儲存格
        
        Args:
            worksheet: 工作表
            merged_ranges: 合併範圍列表
            template_row: 模板行號
            target_row: 目標行號
        """
        # 檢查模板行是否包含標籤
        template_has_tag = False
        for col in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row=template_row, column=col)
            if cell.value and isinstance(cell.value, str):
                if '{{' in cell.value and '}}' in cell.value:
                    template_has_tag = True
                    break
        
        if template_has_tag:
            logger.debug(f"DEBUG: renderer - 模板行 {template_row} 包含標籤，不複製合併儲存格到數據行")
            # 標籤行不應該複製合併儲存格給數據行
            return
        
        for merge_info in merged_ranges:
            # 只處理單行合併範圍（確實屬於模板行本身的）
            if merge_info['min_row'] == merge_info['max_row'] == template_row:
                # 計算行偏移
                row_offset = target_row - template_row
                
                # 創建新的合併範圍
                new_min_row = merge_info['min_row'] + row_offset
                new_max_row = merge_info['max_row'] + row_offset
                new_min_col = merge_info['min_col']
                new_max_col = merge_info['max_col']
                
                # 確保新範圍有效且不與現有合併範圍衝突
                if new_min_row > 0 and new_max_row > 0:
                    new_range_string = f"{get_column_letter(new_min_col)}{new_min_row}:{get_column_letter(new_max_col)}{new_max_row}"
                    
                    try:
                        # 檢查是否已存在重疊的合併範圍
                        existing_merge = self._find_merged_range_for_cell(worksheet, new_min_row, new_min_col)
                        if not existing_merge:
                            worksheet.merge_cells(new_range_string)
                            logger.debug(f"DEBUG: renderer創建合併儲存格範圍: {new_range_string}")
                    except Exception as e:
                        logger.debug(f"DEBUG: 合併儲存格失敗 {new_range_string}: {e}")
                        # 繼續處理，不中斷整個流程
            # 繼續處理，不中斷流程

    def _find_merged_range_for_cell(self, worksheet: Worksheet, row: int, col: int) -> Optional[Any]:
        """
        查找指定儲存格所在的合併範圍
        
        Args:
            worksheet: 工作表
            row: 行號
            col: 列號
            
        Returns:
            合併範圍對象，如果不在合併範圍內則返回None
        """
        for merge_range in worksheet.merged_cells.ranges:
            if (merge_range.min_row <= row <= merge_range.max_row and 
                merge_range.min_col <= col <= merge_range.max_col):
                return merge_range
        return None

    def _detect_template_merge_pattern(self, worksheet: Worksheet, template_row: int, start_col: int, num_cols: int) -> dict:
        """
        偵測模板的合併模式，返回每個欄位應該使用的偏移量
        
        Args:
            worksheet: 工作表
            template_row: 模板行號
            start_col: 開始列號
            num_cols: DataFrame的列數
            
        Returns:
            dict: 包含merge_info的字典，記錄哪些欄位需要合併及跨度
        """
        merge_info = {}
        
        # 檢查模板行上方的標題行是否有合併模式
        for row_offset in [-2, -1, 0]:
            check_row = template_row + row_offset
            if check_row <= 0:
                continue
                
            # 檢查每個可能的合併範圍
            checked_ranges = set()
            for col_idx in range(num_cols + 5):  # 檢查更多列以確保完整性
                actual_col = start_col + col_idx
                merge_range = self._find_merged_range_for_cell(worksheet, check_row, actual_col)
                
                if merge_range and merge_range not in checked_ranges:
                    checked_ranges.add(merge_range)
                    # 如果這是合併範圍的起始列
                    if merge_range.min_col == actual_col:
                        # 記錄合併信息
                        col_offset = actual_col - start_col
                        span = merge_range.max_col - merge_range.min_col + 1
                        if col_offset >= 0 and col_offset < num_cols:
                            merge_info[col_offset] = span
                            logger.debug(f"DEBUG: 偵測到合併模式 - 列{col_offset}跨度{span}列")
        
        return merge_info

    def _is_cell_merged_and_not_top_left(self, worksheet: Worksheet, row: int, col: int) -> bool:
        """
        檢查指定儲存格是否是合併儲存格的非主要部分（即不是合併範圍的左上角）
        
        Args:
            worksheet: 工作表
            row: 行號
            col: 列號
            
        Returns:
            bool: 如果是合併儲存格的非主要部分則返回True
        """
        try:
            cell = worksheet.cell(row=row, column=col)
            
            # 如果儲存格是MergedCell類型，表示它是合併範圍的一部分但不是主要儲存格
            if isinstance(cell, MergedCell):
                return True
                
            # 檢查是否是合併範圍的主要儲存格（左上角）
            for merge_range in worksheet.merged_cells.ranges:
                if (merge_range.min_row == row and merge_range.min_col == col):
                    # 這是合併範圍的主要儲存格（左上角），可以使用
                    return False
                elif (merge_range.min_row <= row <= merge_range.max_row and 
                      merge_range.min_col <= col <= merge_range.max_col):
                    # 這是合併範圍的非主要部分，不能使用
                    return True
            
            return False
        except:
            return False

    def _find_next_available_column(self, worksheet: Worksheet, row: int, start_col: int, max_attempts: int = 20) -> int:
        """
        從指定列開始，找到下一個可用的（未被合併占用的）列
        
        Args:
            worksheet: 工作表
            row: 行號
            start_col: 開始搜索的列號
            max_attempts: 最大嘗試次數 
            
        Returns:
            int: 可用的列號
        """
        for offset in range(max_attempts):
            test_col = start_col + offset
            if not self._is_cell_merged_and_not_top_left(worksheet, row, test_col):
                return test_col
        
        # 如果找不到可用列，返回原始列號（容錯處理）
        return start_col

