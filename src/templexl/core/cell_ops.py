"""
儲存格層級的共用操作（leaf 模組：不得 import 其他 core 模組）。

收納：
- 執行期資料值的統一寫入點（區分模板作者公式與 data 值的公式中和）
- 儲存格樣式複製（快速 _style 索引版與逐屬性安全版）
- 公式參照平移（Excel 複製語意，openpyxl Translator 實作）
"""
import logging
import re
from copy import copy
from datetime import date, datetime
from functools import lru_cache
from typing import Any, Dict

import pandas as pd
from openpyxl.cell.cell import Cell
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# Excel/CSV formula injection 的已知危險前綴。
# openpyxl 僅對 '=' 開頭字串自動標為公式；其餘前綴一併強制文字型別作縱深防禦。
FORMULA_RISK_PREFIXES = ('=', '+', '-', '@', '\t', '\r')


def write_data_value(
    cell: Cell,
    value: Any,
    escape_formulas: bool = True,
    *,
    str_fallback: bool = True,
) -> None:
    """把執行期資料值寫入儲存格，統一型別處理與公式中和。

    中和方式：值原樣寫入後將 ``cell.data_type`` 強制為 ``'s'``（文字），
    使 Excel 開啟時不將其視為公式執行；值本身不加任何前綴、不被改寫。
    僅套用於字串型別——數值（含負數）、日期、布林不受影響。

    模板作者在模板檔內撰寫的公式不經過本函式（見 renderer/block_manager
    的模板列複製路徑），因此不受中和影響。

    Args:
        cell: 目標儲存格（呼叫端負責確保非 MergedCell）。
        value: 執行期資料值（DataFrame 儲存格或純量）。
        escape_formulas: 是否中和危險前綴字串（預設 True）。
        str_fallback: 未知型別（如純數值 DataFrame 經 iterrows 產出的
            numpy 純量）的處理——True 轉為字串（renderer 既有行為）、
            False 原樣交給 openpyxl（block_manager 既有行為）。
            兩者刻意保留以維持黃金基準不變。
    """
    # 分支順序依「熱路徑優先」排列：表格渲染的資料值多為字串與數值，
    # 且已於 table_writer._iter_normalized_rows 逐欄正規化缺失值為 None，
    # 故常見型別皆不需觸及 pd.isna()（大表下該呼叫達數百萬次）。
    # 未正規化的呼叫端（純量標籤等）語意不變：所有缺失值仍收斂為空白儲存格。
    if value is None:
        cell.value = None
    elif isinstance(value, str):
        cell.value = value
        if escape_formulas and value.startswith(FORMULA_RISK_PREFIXES):
            cell.data_type = 's'
    elif isinstance(value, (bool, int, float)):
        # float 的 NaN 是唯一可能缺失的情形（bool/int 不可能）；
        # NaN != NaN 為其定義性質，比 pd.isna() 快一個量級。
        cell.value = None if value != value else value
    elif isinstance(value, (datetime, date)):
        # NaT 為 datetime 子類但語意上是缺失值
        cell.value = None if pd.isna(value) else value
    elif pd.isna(value):
        # pd.NA、np.nan 等非上述型別的缺失值
        cell.value = None
    else:
        cell.value = str(value) if str_fallback else value


def copy_cell_style_info(cell: Cell) -> Dict[str, Any]:
    """
    複製儲存格的樣式資訊（避免深拷貝遞迴問題）

    Args:
        cell: 源儲存格

    Returns:
        Dict: 樣式資訊
    """
    # 使用簡單的屬性複製而不是deepcopy來避免openpyxl proxy遞迴問題
    style_info = {}

    # 安全地複製字體屬性
    if cell.font:
        style_info['font'] = {
            'name': cell.font.name,
            'size': cell.font.size,
            'bold': cell.font.bold,
            'italic': cell.font.italic,
            'underline': cell.font.underline,
            'strike': cell.font.strike,
            'color': str(cell.font.color.rgb) if cell.font.color and hasattr(cell.font.color, 'rgb') else None
        }

    # 安全地複製填充屬性
    if cell.fill and hasattr(cell.fill, 'start_color'):
        style_info['fill'] = {
            'fill_type': cell.fill.fill_type,
            'start_color': str(cell.fill.start_color.rgb) if cell.fill.start_color and hasattr(cell.fill.start_color, 'rgb') else None
        }

    # 安全地複製邊框屬性
    if cell.border:
        style_info['border'] = {
            'left_style': cell.border.left.style if cell.border.left else None,
            'left_color': str(cell.border.left.color.rgb) if cell.border.left and cell.border.left.color and hasattr(cell.border.left.color, 'rgb') else None,
            'right_style': cell.border.right.style if cell.border.right else None,
            'right_color': str(cell.border.right.color.rgb) if cell.border.right and cell.border.right.color and hasattr(cell.border.right.color, 'rgb') else None,
            'top_style': cell.border.top.style if cell.border.top else None,
            'top_color': str(cell.border.top.color.rgb) if cell.border.top and cell.border.top.color and hasattr(cell.border.top.color, 'rgb') else None,
            'bottom_style': cell.border.bottom.style if cell.border.bottom else None,
            'bottom_color': str(cell.border.bottom.color.rgb) if cell.border.bottom and cell.border.bottom.color and hasattr(cell.border.bottom.color, 'rgb') else None
        }

    # 安全地複製對齊屬性
    if cell.alignment:
        style_info['alignment'] = {
            'horizontal': cell.alignment.horizontal,
            'vertical': cell.alignment.vertical,
            'wrap_text': cell.alignment.wrap_text
        }

    return style_info



def copy_cell_style_fast(source_cell, target_cell):
    """
    複製儲存格樣式。

    openpyxl 將字型/框線/填色/數字格式/對齊/保護全部存於單一 StyleArray
    （``cell._style`` 索引）。直接複製該索引可一次涵蓋全部樣式，
    較逐一複製六個 StyleProxy 快數倍，且語意完全等價。

    Args:
        source_cell: 來源儲存格
        target_cell: 目標儲存格
    """
    from copy import copy

    try:
        if getattr(source_cell, "has_style", False):
            target_cell._style = copy(source_cell._style)
    except Exception as e:
        logger.debug(f"DEBUG: 樣式複製失敗: {str(e)}")
        # 後備：至少複製數字格式
        try:
            if source_cell.number_format:
                target_cell.number_format = source_cell.number_format
        except Exception:
            pass


def apply_cell_style_info(cell: Cell, style_info: Dict[str, Any]) -> None:
    """
    應用儲存格樣式資訊

    Args:
        cell: 目標儲存格
        style_info: 樣式資訊字典
    """
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, Protection

    try:
        # 應用字體
        if 'font' in style_info and style_info['font']:
            font_data = style_info['font']
            cell.font = Font(
                name=font_data.get('name'),
                size=font_data.get('size'),
                bold=font_data.get('bold', False),
                italic=font_data.get('italic', False),
                color=font_data.get('color')
            )

        # 應用填充
        if 'fill' in style_info and style_info['fill']:
            fill_data = style_info['fill']
            if fill_data.get('pattern_type'):
                cell.fill = PatternFill(
                    fill_type=fill_data.get('pattern_type'),
                    fgColor=fill_data.get('fg_color'),
                    bgColor=fill_data.get('bg_color')
                )

        # 應用邊框
        if 'border' in style_info and style_info['border']:
            border_data = style_info['border']
            sides = {}
            for side in ['left', 'right', 'top', 'bottom']:
                if side in border_data and border_data[side]:
                    side_data = border_data[side]
                    sides[side] = Side(
                        border_style=side_data.get('style'),
                        color=side_data.get('color')
                    )
            if sides:
                cell.border = Border(**sides)

        # 應用對齊
        if 'alignment' in style_info and style_info['alignment']:
            align_data = style_info['alignment']
            cell.alignment = Alignment(
                horizontal=align_data.get('horizontal'),
                vertical=align_data.get('vertical'),
                wrap_text=align_data.get('wrap_text', False)
            )

        # 應用保護
        if 'protection' in style_info and style_info['protection']:
            prot_data = style_info['protection']
            cell.protection = Protection(
                locked=prot_data.get('locked', True),
                hidden=prot_data.get('hidden', False)
            )
    except Exception as e:
        logger.debug(f"DEBUG: 應用樣式時發生錯誤: {e}")


def copy_cell_style_dict(cell) -> dict:
    """
    將儲存格樣式複製到字典中

    Args:
        cell: 來源儲存格

    Returns:
        dict: 樣式資訊字典
    """
    return {
        'font': cell.font,
        'fill': cell.fill,
        'border': cell.border,
        'alignment': cell.alignment,
        'number_format': cell.number_format,
        'protection': cell.protection
    }


def apply_cell_style_dict(target_cell, style_dict: dict) -> None:
    """
    將樣式字典套用到目標儲存格

    Args:
        target_cell: 目標儲存格
        style_dict: 樣式資訊字典
    """
    try:
        target_cell.font = style_dict['font']
        target_cell.fill = style_dict['fill']
        target_cell.border = style_dict['border']
        target_cell.alignment = style_dict['alignment']
        target_cell.number_format = style_dict['number_format']
        target_cell.protection = style_dict['protection']
    except Exception as e:
        logger.debug(f"DEBUG: 套用樣式失敗: {e}")




def safe_copy_cell_style(source_cell, target_cell) -> None:
    """
    安全地複製單元格樣式，過濾掉問題的顏色值

    註：目前無呼叫者，**刻意保留**作為跨 workbook 複製的後備——
    ``copy_cell_style_fast`` 複製的 ``_style`` 索引是 workbook 範圍的
    StyleArray 參照，跨簿複製時索引不可轉移，必須改用本函式的
    逐屬性重建。同 workbook 情境一律用 fast 版（更快、不丟屬性、
    不逐格新建樣式物件）。

    Args:
        source_cell: 來源單元格
        target_cell: 目標單元格
    """
    try:
        from openpyxl.styles import Font, Border, Side, PatternFill, Alignment
        from openpyxl.styles.colors import Color

        # 安全地複製字體樣式
        if source_cell.font:
            font_color = None
            if source_cell.font.color:
                try:
                    # 驗證顏色值是否有效
                    if hasattr(source_cell.font.color, 'rgb') and source_cell.font.color.rgb:
                        color_value = source_cell.font.color.rgb
                        if color_value and len(color_value) == 8 and all(c in '0123456789ABCDEFabcdef' for c in color_value):
                            font_color = source_cell.font.color
                except:
                    pass  # 忽略顏色錯誤，使用預設

            target_cell.font = Font(
                name=source_cell.font.name,
                size=source_cell.font.size,
                bold=source_cell.font.bold,
                italic=source_cell.font.italic,
                color=font_color
            )

        # 安全地複製邊框樣式
        if source_cell.border:
            def safe_border_color(border_side):
                if border_side and border_side.color:
                    try:
                        if hasattr(border_side.color, 'rgb') and border_side.color.rgb:
                            color_value = border_side.color.rgb
                            if color_value and len(color_value) == 8 and all(c in '0123456789ABCDEFabcdef' for c in color_value):
                                return border_side.color
                    except:
                        pass
                return None

            target_cell.border = Border(
                left=Side(
                    style=source_cell.border.left.style if source_cell.border.left else None,
                    color=safe_border_color(source_cell.border.left)
                ),
                right=Side(
                    style=source_cell.border.right.style if source_cell.border.right else None,
                    color=safe_border_color(source_cell.border.right)
                ),
                top=Side(
                    style=source_cell.border.top.style if source_cell.border.top else None,
                    color=safe_border_color(source_cell.border.top)
                ),
                bottom=Side(
                    style=source_cell.border.bottom.style if source_cell.border.bottom else None,
                    color=safe_border_color(source_cell.border.bottom)
                )
            )

        # 安全地複製填充樣式
        if source_cell.fill and hasattr(source_cell.fill, 'fill_type'):
            safe_start_color = None
            safe_end_color = None

            try:
                if source_cell.fill.start_color and hasattr(source_cell.fill.start_color, 'rgb'):
                    color_value = source_cell.fill.start_color.rgb
                    if color_value and len(color_value) == 8 and all(c in '0123456789ABCDEFabcdef' for c in color_value):
                        safe_start_color = source_cell.fill.start_color
            except:
                pass

            try:
                if source_cell.fill.end_color and hasattr(source_cell.fill.end_color, 'rgb'):
                    color_value = source_cell.fill.end_color.rgb
                    if color_value and len(color_value) == 8 and all(c in '0123456789ABCDEFabcdef' for c in color_value):
                        safe_end_color = source_cell.fill.end_color
            except:
                pass

            if safe_start_color or safe_end_color:
                target_cell.fill = PatternFill(
                    fill_type=source_cell.fill.fill_type,
                    start_color=safe_start_color,
                    end_color=safe_end_color
                )

        # 複製對齊樣式
        if source_cell.alignment:
            target_cell.alignment = Alignment(
                horizontal=source_cell.alignment.horizontal,
                vertical=source_cell.alignment.vertical,
                wrap_text=source_cell.alignment.wrap_text
            )

        # 複製數字格式
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format

    except Exception as e:
        logger.debug(f"DEBUG: 安全樣式複製失敗: {e}")
        # 最後的回退方案：只複製最基本的樣式
        try:
            if source_cell.font:
                target_cell.font = Font(
                    name=source_cell.font.name or 'Calibri',
                    size=source_cell.font.size or 11,
                    bold=source_cell.font.bold or False,
                    italic=source_cell.font.italic or False
                )
            if source_cell.alignment:
                target_cell.alignment = Alignment(
                    horizontal=source_cell.alignment.horizontal,
                    vertical=source_cell.alignment.vertical,
                    wrap_text=source_cell.alignment.wrap_text
                )
        except Exception as final_e:
            logger.debug(f"DEBUG: 最終樣式複製回退也失敗: {final_e}")


def format_cell_ref(row: int, col: int) -> str:
    """以 1-based 列/欄號組出儲存格參考字串（如 (2, 27) -> "AA2"）。"""
    return f"{get_column_letter(col)}{row}"


@lru_cache(maxsize=256)
def _cached_translator(formula: str, origin: str) -> Translator:
    """同一 (公式, 起點) 只 tokenize 一次；Translator 實例可安全重複 translate。"""
    return Translator(formula, origin=origin)


def translate_formula(
    formula: str, src_row: int, src_col: int, tgt_row: int, tgt_col: int
) -> str:
    """依 Excel 複製語意平移公式參照（openpyxl Translator 薄包裝）。

    相對參照隨列/欄位移平移（含跨工作表的相對參照——Excel 複製公式
    時的原生行為）；``$`` 錨定的絕對參照維持不變。以公式 tokenizer
    為基礎，函式名稱（含帶數字者如 ``LOG10``）與字串常數不會被改寫。

    無法解析的公式原樣返回（防禦性降級，不中斷渲染）。

    Args:
        formula: 原始公式（以 ``=`` 開頭；非公式字串原樣返回）。
        src_row: 來源列（模板位置）。
        src_col: 來源欄。
        tgt_row: 目標列。
        tgt_col: 目標欄。

    Returns:
        str: 平移後的公式。
    """
    if not isinstance(formula, str) or not formula.startswith('='):
        return formula
    origin = f"{get_column_letter(src_col)}{src_row}"
    dest = f"{get_column_letter(tgt_col)}{tgt_row}"
    try:
        return _cached_translator(formula, origin).translate_formula(dest)
    except Exception as e:
        logger.debug(f"公式平移失敗，原樣保留（{origin}->{dest}）: {e}")
        return formula
