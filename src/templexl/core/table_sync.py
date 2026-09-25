"""
Excel Table 物件的幾何同步（範圍、篩選範圍、合計列保留）。

表格資料渲染後同步 table.ref／autoFilter.ref／tableColumns，保留表格樣式與
合計列，並依同 sheet 前一個表格的實際渲染結果動態計算起始列（gap 區域）。
``original_table_positions`` 為跨呼叫的表格原始位置快取，由呼叫端
（TemplateRenderer 實例）持有並顯式傳入。

表格幾何的四個不變量集中於本模組決定：

- ``table.ref``       ＝ 表頭列＋資料列＋合計列
- ``autoFilter.ref``  ＝ 表頭列＋資料列（``expected_autofilter_ref``，Excel 慣例）
- ``totalsRowCount`` / ``totalsRowShown`` 保留模板原值（不改寫合計列公式）
- body ≥ 1 列（由 ``table_writer`` 的範圍計算保證）
"""
import logging

import pandas as pd
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from .cell_ops import format_cell_ref

logger = logging.getLogger(__name__)


def expected_autofilter_ref(table) -> str:
    """表格的篩選範圍：``table.ref`` 扣除尾端的合計列。

    Excel 的慣例是篩選範圍只涵蓋「表頭列＋資料列」，合計列在外（模板本身
    即 ``ref=A1:C3`` 對 ``autoFilter=A1:C2``）。無合計列時回傳 ``ref`` 本身，
    行為與既有實作相同。

    篩選範圍的規則只寫在這裡，由渲染後同步、屬性驗證與存檔前最終同步三處
    共用，避免任一處日後又把合計列推進篩選範圍。

    Args:
        table: openpyxl ``Table`` 物件

    Returns:
        str: 應有的 ``autoFilter.ref``；``table.ref`` 為空時原樣回傳。
    """
    ref = getattr(table, 'ref', None)
    if not ref:
        return ref

    totals_row_count = getattr(table, 'totalsRowCount', None) or 0
    if totals_row_count <= 0:
        return ref

    min_col, min_row, max_col, max_row = range_boundaries(ref)
    # 合計列數大於 ref 高度時仍須回傳合法（不倒置）的範圍
    end_row = max(max_row - totals_row_count, min_row)
    return f"{format_cell_ref(min_row, min_col)}:{format_cell_ref(end_row, max_col)}"


def _extend_range_rows(range_str: str, extra_rows: int) -> str:
    """把範圍字串的結束列往下延伸 ``extra_rows`` 列。"""
    if extra_rows <= 0:
        return range_str
    min_col, min_row, max_col, max_row = range_boundaries(range_str)
    return (f"{format_cell_ref(min_row, min_col)}:"
            f"{format_cell_ref(max_row + extra_rows, max_col)}")


def _rebuild_table_columns(table, dataframe, debug_name: str = ""):
    """依 DataFrame 欄名重建 ``tableColumns``，以欄名比對沿用合計列設定。

    Excel 的結構化參照本身就以欄名為鍵，因此欄名相同者沿用模板欄的
    ``totalsRowLabel`` / ``totalsRowFunction`` / ``totalsRowFormula``，
    ``id`` 依新順序重編。欄名改變者不繼承任何合計列設定——其 ``[舊欄名]``
    參照本來就會失效，屬模板作者責任（與 Excel 自身行為一致）。
    """
    from openpyxl.worksheet.table import TableColumn

    template_columns = {
        col.name: col for col in (table.tableColumns or [])
    }

    rebuilt = []
    for index, column_name in enumerate(dataframe.columns, start=1):
        name = str(column_name)
        col = TableColumn(id=index, name=name)
        template_col = template_columns.get(name)
        if template_col is not None:
            col.totalsRowLabel = template_col.totalsRowLabel
            col.totalsRowFunction = template_col.totalsRowFunction
            col.totalsRowFormula = template_col.totalsRowFormula
            logger.debug(
                f"DEBUG_SYNC: 表格 {debug_name} - 欄 {name} 沿用模板合計列設定"
                f"（label={col.totalsRowLabel}, function={col.totalsRowFunction}）"
            )
        rebuilt.append(col)

    table.tableColumns = rebuilt


def update_table_range_sync(table, new_range: str, debug_name: str = "", dataframe=None, include_header: bool = True):
    """
    統一更新表格範圍，確保 table.ref 和 autoFilter.ref 保持同步
    並正確設定所有必要的表格屬性以避免Excel錯誤

    ``new_range`` 的語意為「**表頭列＋資料列**」——呼叫端只知道資料列數，
    合計列的延伸由本函式負責：``table.ref`` 為 ``new_range`` 再往下延伸
    ``totalsRowCount`` 列，``autoFilter.ref`` 則為 ``expected_autofilter_ref``
    的結果（即 ``new_range``）。合計列旗標與合計列公式一律保留模板原值。

    Args:
        table: Excel表格物件
        new_range: 新的範圍字串，語意為「表頭列＋資料列」(例如: "A2:F5")
        debug_name: 用於DEBUG訊息的物件名稱
        dataframe: DataFrame物件（用於更新tableColumns）
        include_header: 是否包含表頭
    """
    old_table_ref = table.ref if hasattr(table, 'ref') else 'Unknown'
    old_autofilter_ref = table.autoFilter.ref if (hasattr(table, 'autoFilter') and table.autoFilter) else 'None'

    # *** 先讀取合計列數（後續所有範圍計算的基礎，不得被覆寫）***
    totals_row_count = getattr(table, 'totalsRowCount', None) or 0

    # *** 保留原始表格樣式設定 ***
    original_style_info = preserve_table_style(table, debug_name)

    # 更新表格範圍（new_range 為表頭＋資料列，ref 需再涵蓋合計列）
    table.ref = _extend_range_rows(new_range, totals_row_count)
    logger.debug(f"DEBUG_SYNC: 表格 {debug_name} - table.ref 更新: {old_table_ref} -> {table.ref} (合計列 {totals_row_count} 列)")

    # *** 關鍵修正：設定表格屬性避免Excel錯誤 ***

    # 1. 處理headerRowCount（修正：noheader模式仍有表頭，只是保留模板表頭）
    if hasattr(table, 'headerRowCount'):
        old_headerRowCount = table.headerRowCount
        # noheader模式實際上仍有表頭（保留的模板表頭），所以headerRowCount應該保持為1
        table.headerRowCount = 1  # 所有表格都有表頭（noheader模式保留模板表頭）
        logger.debug(f"DEBUG_SYNC: 表格 {debug_name} - headerRowCount設為: {table.headerRowCount} (include_header={include_header}, 說明: {'動態表頭' if include_header else '保留模板表頭'})")

    # 1.5. 處理autoFilter（同步更新範圍，不清除；合計列不納入篩選範圍）
    if hasattr(table, 'autoFilter') and table.autoFilter:
        table.autoFilter.ref = expected_autofilter_ref(table)
        logger.debug(f"DEBUG_SYNC: 表格 {debug_name} - autoFilter.ref 同步更新: {old_autofilter_ref} -> {table.autoFilter.ref}")
    else:
        logger.debug(f"DEBUG_SYNC: 表格 {debug_name} - 沒有 autoFilter，跳過同步")

    # 2. 合計列旗標（totalsRowCount／totalsRowShown）一律保留模板原值：
    #    清掉它們會讓 [#Totals] 結構化參照與引用它的具名範圍在 Excel 中失效。
    logger.debug(f"DEBUG_SYNC: 表格 {debug_name} - 保留合計列設定: totalsRowCount={getattr(table, 'totalsRowCount', None)}, totalsRowShown={getattr(table, 'totalsRowShown', None)}")

    # 4. 更新tableColumns（重要修正：根據include_header決定是否更新列定義）
    if (dataframe is not None and hasattr(dataframe, 'columns') and
        hasattr(table, 'tableColumns')):
        try:
            if include_header:
                # 有header模式：依 DataFrame 欄名重建，同名欄沿用合計列設定
                _rebuild_table_columns(table, dataframe, debug_name)
                logger.debug(f"DEBUG_SYNC: 表格 {debug_name} - 更新tableColumns: {len(table.tableColumns)} 列")
                for i, col in enumerate(table.tableColumns):
                    logger.debug(f"  列{i+1}: {col.name}")
            else:
                # noheader模式：完全保留原始tableColumns，不做任何修改
                original_columns_count = len(table.tableColumns) if table.tableColumns else 0
                logger.debug(f"DEBUG_SYNC: 表格 {debug_name} - noheader模式，完全保留原始tableColumns")
                logger.debug(f"  原始列數: {original_columns_count}")

                # 顯示原始的tableColumns
                if table.tableColumns:
                    for i, col in enumerate(table.tableColumns):
                        logger.debug(f"  保留列{i+1}: {col.name}")
                else:
                    logger.debug(f"  警告：原始tableColumns為空")

        except Exception as e:
            logger.debug(f"DEBUG_SYNC: 表格 {debug_name} - tableColumns處理失敗: {str(e)}")

    # 5. 恢復原始表格樣式設定
    restore_table_style(table, original_style_info, debug_name)

    # 6. 最終驗證表格屬性的一致性
    validate_table_properties(table, debug_name)

    logger.debug(f"DEBUG_SYNC: 表格 {debug_name} - 所有屬性更新完成")


def validate_table_properties(table, debug_name: str = ""):
    """
    驗證表格屬性的一致性，確保不會導致Excel錯誤

    Args:
        table: Excel表格物件
        debug_name: 用於DEBUG訊息的物件名稱
    """
    try:
        # 驗證基本屬性
        if not hasattr(table, 'ref') or not table.ref:
            logger.debug(f"DEBUG_VALIDATE: 警告 - 表格 {debug_name} 缺少有效的 ref 屬性")
            return

        # 驗證 autoFilter 與應有篩選範圍（表頭＋資料列，不含合計列）的一致性
        if hasattr(table, 'autoFilter') and table.autoFilter:
            if hasattr(table.autoFilter, 'ref'):
                expected_ref = expected_autofilter_ref(table)
                if table.autoFilter.ref != expected_ref:
                    logger.debug(f"DEBUG_VALIDATE: 警告 - 表格 {debug_name} autoFilter.ref 與應有篩選範圍不一致:")
                    logger.debug(f"  table.ref: {table.ref}")
                    logger.debug(f"  autoFilter.ref: {table.autoFilter.ref}")
                    logger.debug(f"  應有篩選範圍: {expected_ref}")
                    # 強制同步
                    table.autoFilter.ref = expected_ref
                    logger.debug(f"DEBUG_VALIDATE: 已強制同步 autoFilter.ref 到: {expected_ref}")

        # 驗證 totalsRowCount 和 totalsRowShown 的一致性
        if hasattr(table, 'totalsRowCount') and hasattr(table, 'totalsRowShown'):
            totals_count = table.totalsRowCount
            totals_shown = table.totalsRowShown

            if totals_count is None or totals_count == 0:
                if totals_shown is True:
                    logger.debug(f"DEBUG_VALIDATE: 修正 - 表格 {debug_name} totalsRowCount 為 {totals_count} 但 totalsRowShown 為 True")
                    table.totalsRowShown = False
                    logger.debug(f"DEBUG_VALIDATE: 已修正 totalsRowShown 為: False")

        # 驗證 headerRowCount 的合理性
        if hasattr(table, 'headerRowCount'):
            header_count = table.headerRowCount
            if header_count is not None and header_count not in [0, 1]:
                logger.debug(f"DEBUG_VALIDATE: 警告 - 表格 {debug_name} headerRowCount 值異常: {header_count}")

        logger.debug(f"DEBUG_VALIDATE: 表格 {debug_name} 屬性驗證完成")

    except Exception as e:
        logger.debug(f"DEBUG_VALIDATE: 表格 {debug_name} 屬性驗證失敗: {str(e)}")


def update_table_object_range(worksheet, obj_info, dataframe, tag, start_row, start_col, total_rows,
                              original_table_positions: dict):
    """
    更新表格物件的範圍，確保 table.ref 和 autoFilter.ref 保持一致

    Args:
        worksheet: Excel工作表
        obj_info: 物件資訊
        dataframe: DataFrame數據
        tag: 標籤物件
        start_row: 起始行
        start_col: 起始列
        total_rows: 總行數
    """
    try:
        # 找到對應的表格物件
        table_name = obj_info.display_name
        logger.debug(f"DEBUG_TABLE_OBJ: 更新表格物件 {table_name} 的範圍")

        if table_name not in worksheet.tables:
            logger.debug(f"DEBUG_TABLE_OBJ: 直接匹配失敗，嘗試位置匹配: {table_name}")
            # 嘗試根據位置匹配表格
            matched_table = find_table_by_position(worksheet, start_row, start_col)
            if matched_table:
                table_name = matched_table
                logger.debug(f"DEBUG_TABLE_OBJ: 位置匹配成功: {table_name}")
            else:
                logger.debug(f"DEBUG_TABLE_OBJ: 警告 - 找不到表格物件: {table_name}")
                return

        table = worksheet.tables[table_name]

        # *** 新增：計算動態起始位置（gap區域計算）***
        original_start_row = start_row
        actual_start_row = calculate_dynamic_table_start_row(
            worksheet, table_name, original_start_row, dataframe, original_table_positions
        )

        logger.debug(f"DEBUG_TABLE_OBJ: 表格 {table_name} 原始起始行: {original_start_row}, 動態調整後起始行: {actual_start_row}")

        # 如果位置有變化，需要使用調整後的起始行
        if actual_start_row != original_start_row:
            start_row = actual_start_row
            logger.debug(f"DEBUG_TABLE_OBJ: 使用動態調整後的起始行: {start_row}")

        # 取得表格原始範圍
        from openpyxl.utils import range_boundaries
        orig_min_col, orig_min_row, orig_max_col, orig_max_row = range_boundaries(table.ref)
        original_table_cols = orig_max_col - orig_min_col + 1
        dataframe_cols = len(dataframe.columns)

        # 使用DataFrame欄位數和原始表格欄位數中的較大值
        actual_cols = max(dataframe_cols, original_table_cols)

        # 計算新的範圍
        new_end_row = start_row + total_rows - 1
        new_end_col = start_col + actual_cols - 1
        start_cell = format_cell_ref(start_row, start_col)
        end_cell = format_cell_ref(new_end_row, new_end_col)
        new_range = f"{start_cell}:{end_cell}"

        # 檢查是否有noheader條件
        skip_header = tag.has_condition and tag.condition == "noheader"
        include_header = not skip_header

        logger.debug(f"DEBUG_TABLE_OBJ: 表格 {table_name} - skip_header={skip_header}, include_header={include_header}")
        logger.debug(f"DEBUG_TABLE_OBJ: tag.has_condition={tag.has_condition}, tag.condition='{tag.condition if tag.has_condition else 'None'}'")

        # 使用統一方法更新表格範圍，確保 table.ref 和 autoFilter.ref 同步
        update_table_range_sync(table, new_range, table_name, dataframe, include_header)

        # 立即驗證更新是否成功
        if hasattr(table, 'autoFilter') and table.autoFilter:
            current_ref = table.autoFilter.ref
            logger.debug(f"DEBUG_TABLE_OBJ: 驗證 - 當前 autoFilter.ref: {current_ref}")
            if current_ref != new_range:
                logger.debug(f"DEBUG_TABLE_OBJ: 警告！更新後驗證失敗，期望: {new_range}, 實際: {current_ref}")
        else:
            logger.debug(f"DEBUG_TABLE_OBJ: 表格 {table_name} 沒有 autoFilter")

    except Exception as e:
        logger.debug(f"DEBUG_TABLE_OBJ: 更新表格物件範圍失敗: {str(e)}")
        import traceback
        logger.debug("例外堆疊", exc_info=True)


def calculate_dynamic_table_start_row(worksheet: Worksheet,
    current_table_name: str,
    original_start_row: int,
    dataframe: pd.DataFrame,
    original_table_positions: dict,
) -> int:
    """
    計算表格物件的動態起始位置（考慮gap區域）

    當同一sheet中有多個表格物件時，需要根據以下規則動態調整：
    1. 找到上一個順序渲染的表格物件
    2. 計算gap區域：上一個表格的max row到當前表格min row的差距
    3. 根據上一個表格的實際渲染結果重新計算當前表格的起始位置

    Args:
        worksheet: 工作表
        current_table_name: 當前表格名稱
        original_start_row: 原始起始行
        dataframe: 當前表格的數據

    Returns:
        int: 動態調整後的起始行 
    """
    try:
        logger.debug(f"DEBUG_GAP: 開始計算表格 {current_table_name} 的動態起始位置")

        # 獲取當前表格的列位置以判斷是否在同一個垂直列
        from openpyxl.utils import range_boundaries
        current_table = worksheet.tables[current_table_name]
        current_min_col, current_min_row, current_max_col, current_max_row = range_boundaries(current_table.ref)

        logger.debug(f"DEBUG_GAP: 當前表格 {current_table_name} 列範圍: {current_min_col}-{current_max_col}")

        # 收集同一垂直列的表格（只考慮垂直方向推移）
        same_column_tables = []

        for table_name, table in worksheet.tables.items():
            if hasattr(table, 'ref') and table.ref:
                min_col, min_row, max_col, max_row = range_boundaries(table.ref)

                # 檢查是否在同一垂直列（列範圍有重疊）
                has_column_overlap = not (max_col < current_min_col or min_col > current_max_col)

                if has_column_overlap:
                    table_info = {
                        'name': table_name,
                        'min_row': min_row,
                        'max_row': max_row,
                        'min_col': min_col,
                        'max_col': max_col,
                        'table': table
                    }
                    same_column_tables.append(table_info)
                    logger.debug(f"DEBUG_GAP: 同列表格 {table_name} - 行範圍:{min_row}-{max_row}, 列範圍:{min_col}-{max_col}")

                # 保存原始位置信息
                if table_name not in original_table_positions:
                    original_table_positions[table_name] = {
                        'original_min_row': min_row,
                        'original_max_row': max_row
                    }

        # 按起始行排序（只考慮同一垂直列的表格）
        same_column_tables.sort(key=lambda x: x['min_row'])
        logger.debug(f"DEBUG_GAP: 在同一垂直列發現 {len(same_column_tables)} 個表格物件")

        # 找到當前表格在同列表格中的位置
        current_table_index = -1
        for i, table_info in enumerate(same_column_tables):
            if table_info['name'] == current_table_name:
                current_table_index = i
                break

        if current_table_index == -1:
            logger.debug(f"DEBUG_GAP: 找不到當前表格 {current_table_name} 在同列表格中，使用原始位置")
            return original_start_row

        # 如果是同列中的第一個表格，直接使用原始位置
        if current_table_index == 0:
            logger.debug(f"DEBUG_GAP: 當前表格是同列中的第一個表格，使用原始位置 {original_start_row}")
            return original_start_row

        # 計算gap區域 - 只考慮同列中的上一個表格
        previous_table_info = same_column_tables[current_table_index - 1]
        current_table_info = same_column_tables[current_table_index]

        # 獲取上一個表格的原始位置信息
        previous_table_name = previous_table_info['name']
        if previous_table_name in original_table_positions:
            original_previous_max = original_table_positions[previous_table_name]['original_max_row']
        else:
            original_previous_max = previous_table_info['max_row']

        # 獲取當前表格的原始位置信息
        if current_table_name in original_table_positions:
            original_current_min = original_table_positions[current_table_name]['original_min_row']
        else:
            original_current_min = current_table_info['min_row']

        # 計算原始gap大小（基於模板中的設計）
        original_gap_size = original_current_min - original_previous_max

        # 上一個表格的當前實際max row（可能已經被渲染擴展）
        previous_table_actual_max_row = previous_table_info['max_row']

        # 新的起始位置 = 上一個表格的實際結束位置 + 原始gap區域大小
        new_start_row = previous_table_actual_max_row + original_gap_size

        logger.debug(f"DEBUG_GAP: 同列上一個表格 {previous_table_name}:")
        logger.debug(f"DEBUG_GAP:   原始max_row: {original_previous_max}")
        logger.debug(f"DEBUG_GAP:   實際max_row: {previous_table_actual_max_row}")
        logger.debug(f"DEBUG_GAP: 當前表格 {current_table_name}:")
        logger.debug(f"DEBUG_GAP:   原始min_row: {original_current_min}")
        logger.debug(f"DEBUG_GAP:   原始gap大小: {original_gap_size}")
        logger.debug(f"DEBUG_GAP: 動態調整後起始行: {original_start_row} -> {new_start_row}")

        return max(new_start_row, original_start_row)  # 確保不會往上移動

    except Exception as e:
        logger.debug(f"DEBUG_GAP: 計算動態起始位置失敗: {str(e)}")
        import traceback
        logger.debug("例外堆疊", exc_info=True)
        return original_start_row


def preserve_table_style(table, debug_name: str = "") -> dict:
    """
    保留表格的原始樣式設定

    Args:
        table: Excel表格物件
        debug_name: 用於DEBUG訊息的物件名稱

    Returns:
        dict: 保存的樣式信息
    """
    style_info = {}

    try:
        # 保存表格樣式名稱
        if hasattr(table, 'tableStyleInfo') and table.tableStyleInfo:
            style_info['table_style_name'] = table.tableStyleInfo.name if hasattr(table.tableStyleInfo, 'name') else None
            style_info['show_first_column'] = table.tableStyleInfo.showFirstColumn if hasattr(table.tableStyleInfo, 'showFirstColumn') else False
            style_info['show_last_column'] = table.tableStyleInfo.showLastColumn if hasattr(table.tableStyleInfo, 'showLastColumn') else False
            style_info['show_row_stripes'] = table.tableStyleInfo.showRowStripes if hasattr(table.tableStyleInfo, 'showRowStripes') else True
            style_info['show_column_stripes'] = table.tableStyleInfo.showColumnStripes if hasattr(table.tableStyleInfo, 'showColumnStripes') else False

            logger.debug(f"DEBUG_STYLE: 保留表格 {debug_name} 樣式: {style_info['table_style_name']}")
            logger.debug(f"  - showRowStripes: {style_info['show_row_stripes']}")
            logger.debug(f"  - showFirstColumn: {style_info['show_first_column']}")
            logger.debug(f"  - showLastColumn: {style_info['show_last_column']}")
            logger.debug(f"  - showColumnStripes: {style_info['show_column_stripes']}")
        else:
            logger.debug(f"DEBUG_STYLE: 表格 {debug_name} 沒有 tableStyleInfo，使用預設樣式")
            style_info['table_style_name'] = None
            style_info['show_first_column'] = False
            style_info['show_last_column'] = False
            style_info['show_row_stripes'] = True
            style_info['show_column_stripes'] = False

    except Exception as e:
        logger.debug(f"DEBUG_STYLE: 保留表格樣式失敗: {str(e)}")
        # 設定預設值
        style_info = {
            'table_style_name': None,
            'show_first_column': False,
            'show_last_column': False,
            'show_row_stripes': True,
            'show_column_stripes': False
        }

    return style_info


def restore_table_style(table, style_info: dict, debug_name: str = ""):
    """
    恢復表格的樣式設定

    Args:
        table: Excel表格物件
        style_info: 保存的樣式信息
        debug_name: 用於DEBUG訊息的物件名稱
    """
    try:
        from openpyxl.worksheet.table import TableStyleInfo

        # 確保有樣式信息可恢復
        if not style_info:
            logger.debug(f"DEBUG_STYLE: 表格 {debug_name} 沒有樣式信息可恢復")
            return

        # 創建或更新 tableStyleInfo
        if not hasattr(table, 'tableStyleInfo') or not table.tableStyleInfo:
            table.tableStyleInfo = TableStyleInfo()
            logger.debug(f"DEBUG_STYLE: 為表格 {debug_name} 創建新的 tableStyleInfo")

        # 恢復樣式屬性
        if style_info.get('table_style_name'):
            table.tableStyleInfo.name = style_info['table_style_name']
            logger.debug(f"DEBUG_STYLE: 恢復表格 {debug_name} 樣式名: {style_info['table_style_name']}")

        table.tableStyleInfo.showFirstColumn = style_info.get('show_first_column', False)
        table.tableStyleInfo.showLastColumn = style_info.get('show_last_column', False)
        table.tableStyleInfo.showRowStripes = style_info.get('show_row_stripes', True)
        table.tableStyleInfo.showColumnStripes = style_info.get('show_column_stripes', False)

        logger.debug(f"DEBUG_STYLE: 恢復表格 {debug_name} 樣式設定完成")
        logger.debug(f"  - showRowStripes: {table.tableStyleInfo.showRowStripes}")
        logger.debug(f"  - showFirstColumn: {table.tableStyleInfo.showFirstColumn}")
        logger.debug(f"  - showLastColumn: {table.tableStyleInfo.showLastColumn}")
        logger.debug(f"  - showColumnStripes: {table.tableStyleInfo.showColumnStripes}")

    except Exception as e:
        logger.debug(f"DEBUG_STYLE: 保留表格樣式設定失敗: {str(e)}")
        import traceback
        logger.debug("例外堆疊", exc_info=True)


def find_table_by_position(worksheet, target_row, target_col):
    """
    根據位置查找表格物件

    Args:
        worksheet: 工作表
        target_row: 目標行
        target_col: 目標列

    Returns:
        str: 匹配的表格名稱，如果沒有匹配則返回None
    """
    try:
        from openpyxl.utils import range_boundaries

        logger.debug(f"DEBUG_TABLE_MATCH: 尋找位置 ({target_row}, {target_col}) 對應的表格")

        best_match = None
        best_distance = float('inf')

        for table_name in worksheet.tables:
            table = worksheet.tables[table_name]
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)

            logger.debug(f"DEBUG_TABLE_MATCH: 檢查表格 {table_name} - 範圍: 行{min_row}-{max_row}, 列{min_col}-{max_col}")

            # 計算距離用於除錯
            row_distance = abs(target_row - min_row)
            col_distance = abs(target_col - min_col)
            total_distance = row_distance + col_distance
            logger.debug(f"DEBUG_TABLE_MATCH: 表格 {table_name} 距離 - 行距:{row_distance}, 列距:{col_distance}, 總距:{total_distance}")

            # 首先檢查是否目標位置在表格範圍內（優先級最高）
            if (min_row <= target_row <= max_row and
                min_col <= target_col <= max_col):
                logger.debug(f"DEBUG_TABLE_MATCH: 目標位置在表格範圍內: {table_name}")
                return table_name

            # 檢查是否在容錯範圍內
            row_tolerance = 8  # 增加容錯範圍到8行
            col_tolerance = 3  # 增加容錯範圍到3列

            if (row_distance <= row_tolerance and col_distance <= col_tolerance):
                # 同列匹配給予獎勵（減少距離）
                adjusted_distance = total_distance
                if target_col == min_col:
                    adjusted_distance = total_distance * 0.7  # 同列匹配30%獎勵
                    logger.debug(f"DEBUG_TABLE_MATCH: 同列獎勵 {table_name} - 調整距離: {adjusted_distance}")

                # 檢查是否為最佳匹配
                if adjusted_distance < best_distance:
                    best_match = table_name
                    best_distance = adjusted_distance
                    logger.debug(f"DEBUG_TABLE_MATCH: 更新最佳匹配: {table_name} (距離: {adjusted_distance})")

        if best_match:
            logger.debug(f"DEBUG_TABLE_MATCH: 找到最佳匹配表格: {best_match}")
            return best_match

        logger.debug(f"DEBUG_TABLE_MATCH: 未找到匹配的表格")
        return None

    except Exception as e:
        logger.debug(f"DEBUG_TABLE_MATCH: 位置匹配發生錯誤: {str(e)}")
        return None
