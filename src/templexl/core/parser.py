"""
模板解析器
"""
import re
from typing import Iterable, List, Optional, Tuple

from openpyxl import Workbook

from ..exceptions import InvalidTagSyntaxError
from ..models.base import TagType, DataType, RenderDirection, CellPosition
from ..models.tag import Tag


def resolve_sheet_names(
    workbook: Workbook, sheet_names: Optional[Iterable[str]] = None
) -> List[str]:
    """解析目標工作表名稱清單（維持工作簿的實際順序）。

    Args:
        workbook: Excel工作簿物件。
        sheet_names: 欲限縮的工作表名稱；``None`` 表示全部工作表。
            不存在於工作簿的名稱會被忽略。

    Returns:
        List[str]: 依工作簿順序排列的目標工作表名稱。
    """
    if sheet_names is None:
        return list(workbook.sheetnames)
    wanted = set(sheet_names)
    return [name for name in workbook.sheetnames if name in wanted]


class TemplateParser:
    """
    模板解析器類別
    
    負責識別和解析Excel模板中的標籤
    """
    
    # 正則表達式模式 - 支援標籤內的空白字符
    SIMPLE_TAG_PATTERN = r'\{\{\s*([^{}#|]+?)\s*(\|\s*[^{}]*?)?\s*\}\}'  # 排除 # 字符，支援空白
    TABLE_TAG_PATTERN = r'#\{\{\s*([^{}|]+?)\s*(\|\s*[^{}]*?)?\s*\}\}'  # 支援空白
    
    @staticmethod
    def parse_template(
        workbook: Workbook, sheet_names: Optional[Iterable[str]] = None
    ) -> List[Tag]:
        """
        解析模板中的所有標籤

        Args:
            workbook: Excel工作簿物件
            sheet_names: 限縮掃描的工作表名稱；``None``（預設）掃描全部工作表

        Returns:
            List[Tag]: 解析出的標籤清單

        Raises:
            InvalidTagSyntaxError: 標籤語法錯誤
        """
        tags = []

        for sheet_name in resolve_sheet_names(workbook, sheet_names):
            worksheet = workbook[sheet_name]
            sheet_tags = TemplateParser._parse_sheet_tags(worksheet, sheet_name)
            tags.extend(sheet_tags)
            
        return tags
    
    @staticmethod
    def _parse_sheet_tags(worksheet, sheet_name: str) -> List[Tag]:
        """
        解析單一工作表中的標籤
        
        Args:
            worksheet: 工作表物件
            sheet_name: 工作表名稱
            
        Returns:
            List[Tag]: 標籤清單
        """
        tags = []
        
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    cell_tags = TemplateParser._extract_tags_from_cell(
                        cell.value, cell.row, cell.column, sheet_name
                    )
                    tags.extend(cell_tags)
                    
        return tags
    
    @staticmethod
    def _extract_tags_from_cell(cell_value: str, row: int, col: int, sheet_name: str) -> List[Tag]:
        """
        從單一儲存格中提取標籤
        
        Args:
            cell_value: 儲存格值
            row: 行號
            col: 列號  
            sheet_name: 工作表名稱
            
        Returns:
            List[Tag]: 標籤清單
        """
        tags = []
        
        # 檢查表格標籤（#{{...}}）
        table_matches = re.finditer(TemplateParser.TABLE_TAG_PATTERN, cell_value)
        for match in table_matches:
            tag = TemplateParser._create_tag_from_match(
                match, TagType.TABLE, row, col, sheet_name
            )
            tags.append(tag)
        
        # 檢查簡單標籤（{{...}}）- 排除已被表格標籤匹配的部分
        # 先移除表格標籤再匹配簡單標籤
        simple_text = re.sub(TemplateParser.TABLE_TAG_PATTERN, '', cell_value)
        simple_matches = re.finditer(TemplateParser.SIMPLE_TAG_PATTERN, simple_text)
        for match in simple_matches:
            tag = TemplateParser._create_tag_from_match(
                match, TagType.SIMPLE, row, col, sheet_name
            )
            tags.append(tag)
            
        return tags
    
    @staticmethod
    def _create_tag_from_match(match, tag_type: TagType, row: int, col: int, sheet_name: str) -> Tag:
        """
        從正則匹配結果建立標籤物件
        
        Args:
            match: 正則匹配物件
            tag_type: 標籤類型
            row: 行號
            col: 列號
            sheet_name: 工作表名稱
            
        Returns:
            Tag: 標籤物件
        """
        tag_name = match.group(1).strip()
        condition_part = match.group(2)
        
        has_condition = condition_part is not None
        condition = None
        if has_condition:
            condition = condition_part.strip().lstrip('|').strip()
        
        # 驗證標籤語法
        if not TemplateParser.validate_tag_name(tag_name):
            position = f"{sheet_name}!{row}:{col}"
            raise InvalidTagSyntaxError(match.group(0), position)
        
        return Tag(
            tag_name=tag_name,
            tag_type=tag_type,
            has_condition=has_condition,
            condition=condition,
            cell_position=CellPosition(row=row, col=col),
            data_type=DataType.UNKNOWN,  # 初始設為未知，後續根據實際數據確定
            render_direction=RenderDirection.VERTICAL if tag_type == TagType.TABLE else RenderDirection.HORIZONTAL,
            sheet_name=sheet_name
        )
    
    @staticmethod
    def validate_syntax(tag_string: str) -> bool:
        """
        驗證標籤語法是否有效
        
        Args:
            tag_string: 標籤字串
            
        Returns:
            bool: 語法是否有效
        """
        # 檢查是否為表格標籤（完整匹配）
        if re.fullmatch(TemplateParser.TABLE_TAG_PATTERN, tag_string):
            return True
            
        # 檢查是否為簡單標籤（完整匹配）
        if re.fullmatch(TemplateParser.SIMPLE_TAG_PATTERN, tag_string):
            return True
            
        return False
    
    @staticmethod
    def validate_tag_name(tag_name: str) -> bool:
        """
        驗證標籤名稱是否有效
        
        Args:
            tag_name: 標籤名稱
            
        Returns:
            bool: 名稱是否有效
        """
        if not tag_name or not tag_name.strip():
            return False
            
        # 標籤名稱不能包含特殊字元
        invalid_chars = ['{{', '}}', '|', '#']
        for char in invalid_chars:
            if char in tag_name:
                return False
                
        return True
    
    @staticmethod
    def extract_condition(tag_string: str) -> Optional[str]:
        """
        提取條件標籤
        
        Args:
            tag_string: 標籤字串
            
        Returns:
            Optional[str]: 條件字串，如果沒有條件則返回None
        """
        # 尋找管道符號
        pipe_pos = tag_string.find('|')
        if pipe_pos == -1:
            return None
            
        # 提取條件部分
        condition_part = tag_string[pipe_pos + 1:]
        # 移除結尾的 }}
        condition_part = condition_part.rstrip('}').strip()
        
        return condition_part if condition_part else None
    
    @staticmethod
    def is_table_tag(tag_string: str) -> bool:
        """
        判斷是否為表格標籤
        
        Args:
            tag_string: 標籤字串 
            
        Returns:
            bool: 是否為表格標籤
        """
        return tag_string.startswith('#{{') and tag_string.endswith('}}')
    
    @staticmethod
    def extract_tag_info(tag_string: str) -> Tuple[str, Optional[str], TagType]:
        """
        提取標籤資訊
        
        Args:
            tag_string: 標籤字串
            
        Returns:
            Tuple[str, Optional[str], TagType]: (標籤名稱, 條件, 標籤類型)
            
        Raises:
            InvalidTagSyntaxError: 標籤語法錯誤
        """
        if not TemplateParser.validate_syntax(tag_string):
            raise InvalidTagSyntaxError(tag_string)
        
        tag_type = TagType.TABLE if TemplateParser.is_table_tag(tag_string) else TagType.SIMPLE
        condition = TemplateParser.extract_condition(tag_string)
        
        # 提取標籤名稱
        if tag_type == TagType.TABLE:
            # 移除 #{{ 和 }}
            content = tag_string[3:-2]
        else:
            # 移除 {{ 和 }}
            content = tag_string[2:-2]
        
        # 如果有條件，移除條件部分
        if condition:
            pipe_pos = content.find('|')
            tag_name = content[:pipe_pos].strip()
        else:
            tag_name = content.strip()
        
        return tag_name, condition, tag_type
