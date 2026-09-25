"""工作表結構操作（leaf 模組：不得 import 其他 core 模組）。

收納 openpyxl「做不到或做不對」的結構編輯——目前為 Excel 等效的刪整欄。

openpyxl 的 ``Worksheet.delete_cols`` 只搬動儲存格（值與樣式），**不會**調整
合併範圍、欄寬與列印範圍，且以 ``translate=False`` 呼叫內部搬移（公式參照
完全不改寫）。本模組補齊前三者；對無法保證正確性的特徵一律 fail fast。

另補償合併範圍錨點格的樣式：openpyxl 載入時，合併範圍內的非錨點格會被換成
無真實樣式的 ``MergedCell`` 佔位格，``unmerge_cells`` 又會直接刪除這些佔位格；
而 ``delete_cols`` 搬移前會把右側範圍補實為預設空白格再左移覆蓋。三者疊加使
「錨點欄恰為被刪除欄」的合併，其存活儲存格淪為無樣式空白格（Excel 的
``EntireColumn.Delete`` 則保留格式）。故刪欄前快照錨點樣式、刪欄後套回。

為什麼公式一律拒絕而非嘗試平移：``cell_ops.translate_formula`` 實作的是
Excel **複製語意**（公式自己移動，相對參照隨之平移、``$`` 絕對參照不動），
而刪欄需要的是 **參照改寫語意**（公式不動也要改：指向刪除欄右側的參照
——含絕對參照——左移一欄，指向被刪欄者變 ``#REF!``，左側不動）。兩者僅在
窄情境碰巧等價，硬套會靜默算錯。 
"""
from __future__ import annotations

import logging
from copy import copy
from typing import Any, Iterable, List, Optional, Tuple

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string, range_boundaries
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from ..exceptions import StructuralEditError

logger = logging.getLogger(__name__)

_OPERATION = "刪除整欄"

# 刪欄時會一併搬移的欄寬屬性（ColumnDimension 的視覺設定）。
# 不含 min/max（openpyxl 寫出時依 index 重算）與 customWidth（唯讀衍生屬性，
# 由 width 是否設定決定）。
_COLUMN_DIMENSION_ATTRS = (
    "width",
    "hidden",
    "bestFit",
    "outlineLevel",
    "collapsed",
)

# 合併範圍調整後的座標：(min_row, min_col, max_row, max_col)
_MergeCoords = Tuple[int, int, int, int]

# 錨點樣式復原指令：(row, col, 錨點格的 StyleArray 快照)
_StyleRestore = Tuple[int, int, Any]


def delete_column(worksheet: Worksheet, idx: int) -> None:
    """刪除第 ``idx`` 欄（1-based），結果與 Excel「刪除整欄」等效。

    涵蓋範圍：儲存格值與樣式左移（委派 openpyxl）、合併範圍修正、
    錨點樣式復原、欄寬左移、列印範圍收縮。

    多欄刪除請由呼叫端**由大到小**逐次呼叫，使座標全程以原始欄位計算、
    不受先前刪除的位移影響。

    Args:
        worksheet: 目標工作表。
        idx: 欲刪除的欄索引（1-based，A 欄為 1）。

    Raises:
        ValueError: ``idx`` 小於 1。
        StructuralEditError: 工作表存在無法保證正確調整的特徵。
    """
    if idx < 1:
        raise ValueError(f"欄索引須為 1-based 正整數，收到 {idx}")

    _reject_unsupported_features(worksheet)

    # 順序關鍵：先卸下受影響的合併範圍（避免 MergedCell 干擾儲存格搬移，
    # 並讓刪欄後可用調整過的座標重建），再刪欄，接著把錨點樣式補回存活位置，
    # 最後補回欄寬/列印範圍/合併。樣式復原必須早於重建合併——merge_cells 會
    # 依新錨點格的框線重繪範圍外框，錨點得先拿回樣式。
    pending_merges, style_restores = _detach_affected_merges(worksheet, idx)
    worksheet.delete_cols(idx)
    _restore_anchor_styles(worksheet, style_restores)
    _shift_column_dimensions(worksheet, idx)
    _shrink_print_area(worksheet, idx)
    _reattach_merges(worksheet, pending_merges)

    logger.debug("已刪除工作表 '%s' 第 %d 欄", worksheet.title, idx)


# ======================================================================
# fail fast 特徵掃描
# ======================================================================


def _reject_unsupported_features(worksheet: Worksheet) -> None:
    """ 工作表存在無法正確調整的特徵時拒絕刪欄（訊息含種類與位置）。 """

    formula_cells = _find_formula_cells(worksheet, limit=5)
    if formula_cells:
        raise StructuralEditError(
            _OPERATION,
            f"工作表含公式，openpyxl 刪欄不會改寫公式參照（範例位置: "
            f"{', '.join(formula_cells)}）",
            worksheet.title,
        )

    images = getattr(worksheet, "_images", None)
    if images:
        raise StructuralEditError(
            _OPERATION, f"工作表含 {len(images)} 張圖片，欄向錨點無法正確位移",
            worksheet.title,
        )

    charts = getattr(worksheet, "_charts", None)
    if charts:
        raise StructuralEditError(
            _OPERATION, f"工作表含 {len(charts)} 個圖表，資料參照無法正確位移",
            worksheet.title,
        )

    if worksheet.tables:
        raise StructuralEditError(
            _OPERATION,
            f"工作表含 Excel 表格 {sorted(worksheet.tables)}，表格範圍無法正確調整",
            worksheet.title,
        )

    conditional = list(worksheet.conditional_formatting)
    if conditional:
        ranges = [str(cf.sqref) for cf in conditional[:5]]
        raise StructuralEditError(
            _OPERATION, f"工作表含條件格式（範圍: {', '.join(ranges)}）",
            worksheet.title,
        )

    validations = worksheet.data_validations.dataValidation
    if validations:
        ranges = [str(dv.sqref) for dv in validations[:5]]
        raise StructuralEditError(
            _OPERATION, f"工作表含資料驗證（範圍: {', '.join(ranges)}）",
            worksheet.title,
        )


def _find_formula_cells(worksheet: Worksheet, limit: int = 5) -> List[str]:
    """回傳含公式的儲存格座標（最多 ``limit`` 筆，供錯誤訊息定位）。"""
    found: List[str] = []
    for row in worksheet.iter_rows():
        for cell in row:
            value = cell.value
            if cell.data_type == "f" or (
                isinstance(value, str) and value.startswith("=")
            ):
                found.append(cell.coordinate)
                if len(found) >= limit:
                    return found
    return found


# ======================================================================
# 合併範圍
# ======================================================================


def _detach_affected_merges(
    worksheet: Worksheet, idx: int
) -> Tuple[List[_MergeCoords], List[_StyleRestore]]:
    """卸下所有受刪欄影響的合併範圍。

    回傳 ``(調整後應重建的座標, 應復原的錨點樣式)``。座標調整規則：

    - 完全位於刪除欄左側：不受影響，保留原樣。
    - 完全落在刪除欄之內（單欄合併）：整段隨欄消失，丟棄。
    - 跨越刪除欄：寬度縮 1。
    - 完全位於右側：整體左移 1。
    - 調整後塌成單一儲存格：丟棄（不留 1×1 合併）。

    錨點欄恰為刪除欄且範圍延伸至其右側者，另快照錨點樣式待刪欄後復原
    （見模組 docstring 的三段機制）。單欄合併整段隨欄消失，其位置由右側
    儲存格遞補，不在此列。
    """
    pending: List[_MergeCoords] = []
    restores: List[_StyleRestore] = []

    for merged in list(worksheet.merged_cells.ranges):
        cell_range = CellRange(str(merged))
        if cell_range.max_col < idx:
            continue  # 完全在左側，不受影響

        if cell_range.min_col == idx and cell_range.max_col > idx:
            # 快照須早於 unmerge：此刻錨點格仍持有載入時 MergedCellRange.format()
            # 依合併範圍補上的完整外框。刪欄後存活位置即 (min_row, idx)。
            anchor = worksheet.cell(row=cell_range.min_row, column=idx)
            restores.append((cell_range.min_row, idx, copy(anchor._style)))

        worksheet.unmerge_cells(str(merged))

        if cell_range.min_col == idx and cell_range.max_col == idx:
            continue  # 整段就在被刪的欄上，隨欄消失

        min_col = cell_range.min_col - 1 if cell_range.min_col > idx else cell_range.min_col
        max_col = cell_range.max_col - 1

        if min_col == max_col and cell_range.min_row == cell_range.max_row:
            continue  # 塌成 1×1，不再是合併

        pending.append((cell_range.min_row, min_col, cell_range.max_row, max_col))

    return pending, restores


def _restore_anchor_styles(
    worksheet: Worksheet, restores: Iterable[_StyleRestore]
) -> None:
    """把刪欄前快照的錨點樣式套回存活位置的儲存格。

    直接指派 ``_style``（儲存格樣式在各樣式表中的完整索引紀錄）而非逐一複製
    font/border/fill/alignment/number_format/protection：後者項目零散易漏，
    而索引在同一 workbook 內恆有效。

    僅復原樣式、不回存值——值隨錨點欄消失即 Excel 語意，回寫由呼叫端負責。
    """
    for row, col, style in restores:
        worksheet.cell(row=row, column=col)._style = style


def _reattach_merges(worksheet: Worksheet, pending: Iterable[_MergeCoords]) -> None:
    """以調整後的座標重建合併範圍。"""
    for min_row, min_col, max_row, max_col in pending:
        worksheet.merge_cells(
            start_row=min_row,
            start_column=min_col,
            end_row=max_row,
            end_column=max_col,
        )


# ======================================================================
# 欄寬
# ======================================================================


def _shift_column_dimensions(worksheet: Worksheet, idx: int) -> None:
    """欄寬設定自刪除欄起左移一位。

    ``column_dimensions`` 是稀疏字典（可能有 A–O 與 Q 而缺 P）：右鄰欄
    無紀錄者一律清除該欄紀錄回歸預設寬度，不得殘留刪除前的舊寬度。
    """
    dimensions = worksheet.column_dimensions
    snapshot = {}
    for letter in list(dimensions.keys()):
        try:
            col = column_index_from_string(letter)
        except ValueError:
            continue  # 非欄名鍵（理論上不存在），保守略過
        snapshot[col] = {
            attr: getattr(dimensions[letter], attr, None)
            for attr in _COLUMN_DIMENSION_ATTRS
        }

    if not snapshot:
        return

    for col in range(idx, max(snapshot) + 1):
        letter = get_column_letter(col)
        source = snapshot.get(col + 1)
        if source is None:
            if letter in dimensions:
                del dimensions[letter]
            continue
        target = dimensions[letter]
        for attr, value in source.items():
            setattr(target, attr, value)


# ======================================================================
# 列印範圍
# ======================================================================


def _shrink_print_area(worksheet: Worksheet, idx: int) -> None:
    """列印範圍隨刪欄收縮/左移；範圍整段被刪則移除該段。"""
    raw = worksheet.print_area
    if not raw:
        return

    adjusted: List[str] = []
    for part in str(raw).split(","):
        part = part.strip()
        if not part:
            continue
        # 去除 "'工作表名'!" 前綴；setter 會依工作表標題重新加上
        ref = part.rsplit("!", 1)[-1]
        shifted = _shift_range_ref(ref, idx)
        if shifted:
            adjusted.append(shifted)

    worksheet.print_area = adjusted or None


def _shift_range_ref(ref: str, idx: int) -> Optional[str]:
    """調整單一 ``$B$1:$AC$10`` 形式的範圍字串；整段被刪則回傳 ``None``。"""
    try:
        min_col, min_row, max_col, max_row = range_boundaries(ref)
    except ValueError:
        logger.debug("列印範圍 '%s' 無法解析，原樣保留", ref)
        return ref

    if None in (min_col, max_col):
        return ref
    if max_col < idx:
        return ref  # 完全在左側

    new_min_col = min_col - 1 if min_col > idx else min_col
    new_max_col = max_col - 1
    if new_max_col < new_min_col:
        return None  # 整段隨欄消失

    return (
        f"${get_column_letter(new_min_col)}${min_row}"
        f":${get_column_letter(new_max_col)}${max_row}"
    )
