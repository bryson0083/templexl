"""
不受信任輸入處理測試（對應 specs/input-safety 與 public-api 原子寫入）。

驗證公式中和（escape_formulas）、模板 zip 資源檢查、原子寫入。 
"""

from __future__ import annotations

import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from templexl import (
    FileFormatError,
    RenderError,
    TemplateResourceError,
    render,
)

EVIL_FORMULA = '=HYPERLINK("http://evil.example","x")'


@pytest.fixture
def simple_template(tmp_path):
    """含簡單標籤、模板公式與表格標籤的最小模板。"""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Hello {{name}}"
    ws["A2"] = "{{scalar}}"
    ws["B1"] = "=SUM(C1:C2)"  # 模板作者的公式
    ws["A5"] = "#{{df}}"
    path = tmp_path / "template.xlsx"
    wb.save(path)
    return path


def _render_data() -> dict:
    return {
        "name": "=2+5",
        "scalar": EVIL_FORMULA,
        "df": pd.DataFrame({
            "備註": [EVIL_FORMULA, "正常文字", "@提及"],
            "金額": [-42, 3.14, 100],
        }),
    }


def _formula_cells(path) -> list:
    wb = load_workbook(path)
    return [
        (ws.title, cell.coordinate, cell.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]


def test_escape_formulas_default_neutralizes_data(simple_template, tmp_path):
    """預設中和：資料帶入的公式前綴字串以純文字寫入，值不被改寫。"""
    out = tmp_path / "out.xlsx"
    render(str(simple_template), str(out), data=_render_data())

    formulas = _formula_cells(out)
    # 唯一的公式是模板作者寫的 B1
    assert [(t, c) for t, c, _ in formulas] == [("Sheet", "B1")]
    assert formulas[0][2] == "=SUM(C1:C2)"

    # 值原樣保留（無前綴單引號等改寫）
    ws = load_workbook(out).active
    values = {c.value for row in ws.iter_rows() for c in row}
    assert EVIL_FORMULA in values
    assert "Hello =2+5" in values


def test_escape_formulas_off_preserves_legacy_behavior(simple_template, tmp_path):
    """顯式停用：= 開頭的資料值以公式形式寫入（第一階段行為）。"""
    out = tmp_path / "out.xlsx"
    render(str(simple_template), str(out), data=_render_data(), escape_formulas=False)

    formula_coords = {(t, c) for t, c, _ in _formula_cells(out)}
    assert len(formula_coords) > 1  # 除模板 B1 外，資料公式也被寫成公式


def test_numeric_values_unaffected(simple_template, tmp_path):
    """數值（含負數）不受中和影響，維持數值型別。"""
    out = tmp_path / "out.xlsx"
    render(str(simple_template), str(out), data=_render_data())
    ws = load_workbook(out).active
    numeric = {
        c.value
        for row in ws.iter_rows()
        for c in row
        if c.data_type == "n" and c.value is not None
    }
    assert -42 in numeric
    assert 3.14 in numeric


def test_zip_bomb_template_rejected(tmp_path):
    """壓縮比超限的模板在載入前被拒絕。"""
    bomb = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/huge.bin", b"\x00" * (64 * 1024 * 1024))  # 64MB 零 → 壓縮比極高
    with pytest.raises(TemplateResourceError):
        render(str(bomb), str(tmp_path / "out.xlsx"), data={})


def test_invalid_zip_rejected(tmp_path):
    """副檔名正確但非 zip 結構的檔案被拒絕。"""
    fake = tmp_path / "fake.xlsx"
    fake.write_text("not a zip")
    with pytest.raises(FileFormatError):
        render(str(fake), str(tmp_path / "out.xlsx"), data={})


def test_failed_save_leaves_no_partial_output(simple_template, tmp_path, monkeypatch):
    """存檔中途失敗：目標路徑維持舊檔內容、無暫存檔殘留。"""
    out = tmp_path / "out.xlsx"
    out.write_bytes(b"OLD_CONTENT")

    from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook

    def boom(self, filename):
        raise OSError("disk full")

    monkeypatch.setattr(OpenpyxlWorkbook, "save", boom)
    with pytest.raises(RenderError):
        render(str(simple_template), str(out), data=_render_data())

    assert out.read_bytes() == b"OLD_CONTENT"
    assert not list(tmp_path.glob("*.tmp"))
