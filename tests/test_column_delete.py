"""Excel 等效刪整欄契約測試（對應 specs/column-delete）。

模板形狀取自實務案例：12 欄、標題橫跨 A1:L1、
表頭 K7:L7 群組合併、稀疏欄寬設定。 
"""
from __future__ import annotations

import pytest
from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Border, Font, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from templexl import Book, StructuralEditError

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RED_BOLD = Font(color="FFFF0000", bold=True)

# 欄寬刻意稀疏：有 A~O 與 Q，缺 P
COLUMN_WIDTHS = {
    "A": 21.5, "B": 27.5, "C": 39.0, "D": 12.5, "E": 13.0, "F": 16.5,
    "G": 16.0, "H": 13.5, "I": 12.5, "J": 13.5, "K": 12.5, "L": 16.25,
    "M": 9.0, "N": 11.5, "O": 10.5, "Q": 9.0,
}


def _build_template(path, *, with_print_area=False):
    """建立形似案例報表的 12 欄模板。"""
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws.title = "補貼款報表"

    ws["A1"] = "Templexl Corporation"
    ws.merge_cells("A1:L1")                 # 跨越刪除欄 → 應縮為 A1:K1

    headers = {
        "A7": "管轄單位", "B7": "運具", "C7": "業者名稱",
        "I7": "中央補助金額", "K7": "地方補助金額",
    }
    for coord, value in headers.items():
        ws[coord] = value

    # 表頭樣式取自實務模板：「地方補助金額」群組為紅字粗體、四邊框線。
    # 合併錨點的樣式是刪欄後唯一可承接的來源（成員格於 openpyxl 載入時
    # 已被換成無樣式的 MergedCell 佔位格），故須在合併前設於錨點上。
    ws["A1"].font = Font(bold=True, size=16)
    ws["I7"].font = Font(bold=True)
    ws["I7"].border = BOX
    ws["K7"].font = RED_BOLD
    ws["K7"].border = BOX

    ws.merge_cells("A7:A8")                 # 左側，不受影響
    ws.merge_cells("I7:J7")                 # 左側，不受影響
    ws.merge_cells("K7:L7")                 # 恰為刪除欄 + 右鄰 → 塌陷移除

    sub_headers = ["比例", "筆數", "金額", "金額", "金額", "比例", "金額", "比例", "金額"]
    for offset, text in enumerate(sub_headers):
        ws.cell(row=8, column=4 + offset, value=text)

    # 資料列：K 欄為比例（百分比格式）、L 欄為金額（千分位格式）
    for row in (9, 10):
        for col in range(1, 13):
            cell = ws.cell(row=row, column=col, value=f"r{row}c{col}")
            cell.border = BOX
        ws.cell(row=row, column=11).value = 0.5
        ws.cell(row=row, column=11).number_format = "0.00%"
        ws.cell(row=row, column=12).value = 12345
        ws.cell(row=row, column=12).number_format = "#,##0"

    ws.merge_cells("A9:C9")                 # 完全在左側 → 原樣保留

    for letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[letter].width = width

    if with_print_area:
        ws.print_area = "$B$1:$AC$10"

    wb.save(path)
    return path


@pytest.fixture
def template(tmp_path):
    return _build_template(tmp_path / "template.xlsx")


def _saved(book, tmp_path, name="out.xlsx"):
    out = tmp_path / name
    book.save(out)
    return load_workbook(out).active


# ===================================================================
# 值與格式左移
# ===================================================================


def test_values_and_formats_shift_left(template, tmp_path):
    """新 K 欄承接原 L 欄的值、數字格式與框線。"""
    with Book(template) as bk:
        bk.sheets[0].delete_columns(11)
        ws = _saved(bk, tmp_path)

    cell = ws.cell(row=9, column=11)
    assert cell.value == 12345
    assert cell.number_format == "#,##0"
    assert [getattr(cell.border, s).style for s in ("left", "right", "top", "bottom")] == [
        "thin", "thin", "thin", "thin"
    ]
    # 原 L 欄（第 12 欄）已無內容
    assert ws.cell(row=9, column=12).value is None


def test_multi_column_delete_high_to_low(template, tmp_path):
    """由大到小逐次刪欄：兩欄各自正確刪除，第二次不受第一次位移影響。"""
    with Book(template) as bk:
        sheet = bk.sheets[0]
        for col in sorted([12, 4], reverse=True):
            sheet.delete_columns(col)
        ws = _saved(bk, tmp_path)

    # 原 12 欄剩 10 欄；D（原第 4 欄，值 r9c4）與 L（原第 12 欄，值 12345）都應消失
    # 註：B9/C9 為 None 是模板的 A9:C9 合併所致（合併只保留左上值），與刪欄無關
    row9 = [ws.cell(row=9, column=c).value for c in range(1, 11)]
    assert "r9c4" not in row9
    assert 12345 not in row9
    assert row9[0] == "r9c1"
    assert row9[3] == "r9c5"   # D 承接原 E
    assert row9[9] == 0.5      # J 承接原 K（比例欄）


# ===================================================================
# 合併範圍
# ===================================================================


def test_spanning_merge_shrinks(template, tmp_path):
    with Book(template) as bk:
        bk.sheets[0].delete_columns(11)
        ws = _saved(bk, tmp_path)
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "A1:K1" in merged
    assert "A1:L1" not in merged


def test_collapsed_merge_removed_and_writable(template, tmp_path):
    """K7:L7 塌成單格：合併消失，新 K7 為普通儲存格（可直接寫值）。"""
    with Book(template) as bk:
        sheet = bk.sheets[0]
        sheet.delete_columns(11)
        merged_in_memory = {str(r) for r in sheet.ws.merged_cells.ranges}
        assert not any(m.startswith("K7") for m in merged_in_memory)
        # 表頭回寫：塌陷合併已移除，故 K7 可直接賦值（MergedCell 為唯讀）
        sheet.ws.cell(row=7, column=11).value = "地方補助金額"
        ws = _saved(bk, tmp_path)

    assert ws.cell(row=7, column=11).value == "地方補助金額"
    assert ws.cell(row=8, column=11).value == "金額"  # 原 L8 左移


def _borders(cell):
    return [getattr(cell.border, s).style for s in ("left", "right", "top", "bottom")]


def test_collapsed_merge_anchor_style_inherited(template, tmp_path):
    """K7:L7 塌成單格：新 K7 承接原錨點的紅字/粗體/框線，值不回存。"""
    with Book(template) as bk:
        bk.sheets[0].delete_columns(11)
        ws = _saved(bk, tmp_path)

    cell = ws.cell(row=7, column=11)
    assert cell.font.color.rgb == "FFFF0000"
    assert cell.font.bold is True
    assert _borders(cell) == ["thin", "thin", "thin", "thin"]
    # 值隨錨點欄消失（Excel 語意），回寫由呼叫端負責
    assert cell.value is None


def test_anchor_column_deleted_but_merge_survives(tmp_path):
    """K7:M7 錨點欄被刪、合併縮為 K7:L7：新錨點仍承接原錨點樣式。

    只斷言 left/top 框線——``MergedCellRange._get_borders`` 於重建合併時會把
    範圍右下角儲存格的 right/bottom 併入錨點，那兩側不純粹來自快照。
    """
    path = tmp_path / "surviving_merge.xlsx"
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws["K7"] = "地方補助金額"
    ws["K7"].font = RED_BOLD
    ws["K7"].border = BOX
    ws.merge_cells("K7:M7")
    wb.save(path)

    with Book(path) as bk:
        bk.sheets[0].delete_columns(11)
        out = _saved(bk, tmp_path)

    assert "K7:L7" in {str(r) for r in out.merged_cells.ranges}
    anchor = out.cell(row=7, column=11)
    assert anchor.font.color.rgb == "FFFF0000"
    assert anchor.font.bold is True
    assert [anchor.border.left.style, anchor.border.top.style] == ["thin", "thin"]
    assert anchor.value is None


def test_anchors_outside_deleted_column_unaffected(template, tmp_path):
    """錨點不在刪除欄者（I7:J7 在左側、A1:L1 錨點在左），樣式與值皆不變。"""
    with Book(template) as bk:
        bk.sheets[0].delete_columns(11)
        ws = _saved(bk, tmp_path)

    title = ws["A1"]
    assert title.value == "Templexl Corporation"
    assert (title.font.bold, title.font.size) == (True, 16)

    central = ws["I7"]
    assert central.value == "中央補助金額"
    assert central.font.bold is True
    assert _borders(central) == ["thin", "thin", "thin", "thin"]


def test_left_side_merges_untouched(template, tmp_path):
    with Book(template) as bk:
        bk.sheets[0].delete_columns(11)
        ws = _saved(bk, tmp_path)
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "A9:C9" in merged
    assert "A7:A8" in merged
    assert "I7:J7" in merged


def test_merge_entirely_within_deleted_column_dropped(tmp_path):
    """整段落在被刪欄上的單欄合併（K9:K10）隨欄消失，不留殘範圍。"""
    path = tmp_path / "single_col_merge.xlsx"
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws["K9"] = "整欄合併"
    ws.merge_cells("K9:K10")
    ws["M9"] = "右側"
    wb.save(path)

    with Book(path) as bk:
        bk.sheets[0].delete_columns(11)
        out = _saved(bk, tmp_path)

    assert [str(r) for r in out.merged_cells.ranges] == []
    assert out["L9"].value == "右側"


# ===================================================================
# 欄寬
# ===================================================================


def test_column_width_inherited_from_right_neighbour(template, tmp_path):
    with Book(template) as bk:
        bk.sheets[0].delete_columns(11)
        ws = _saved(bk, tmp_path)
    assert ws.column_dimensions["K"].width == pytest.approx(COLUMN_WIDTHS["L"])
    assert ws.column_dimensions["L"].width == pytest.approx(COLUMN_WIDTHS["M"])


def test_sparse_column_widths_do_not_smear(template, tmp_path):
    """P 欄原本無紀錄：左移後 O 應回歸預設（無紀錄），Q 的寬度落到 P。"""
    with Book(template) as bk:
        bk.sheets[0].delete_columns(11)
        ws = _saved(bk, tmp_path)

    recorded = {k: v.width for k, v in ws.column_dimensions.items() if v.width}
    assert "O" not in recorded
    assert recorded["N"] == pytest.approx(COLUMN_WIDTHS["O"])
    assert recorded["P"] == pytest.approx(COLUMN_WIDTHS["Q"])
    assert "Q" not in recorded


# ===================================================================
# 列印範圍
# ===================================================================


def test_print_area_shrinks(tmp_path):
    path = _build_template(tmp_path / "with_print_area.xlsx", with_print_area=True)
    with Book(path) as bk:
        bk.sheets[0].delete_columns(13)
        ws = _saved(bk, tmp_path)
    assert ws.print_area.endswith("$B$1:$AB$10")


def test_print_area_left_of_deleted_column_untouched(tmp_path):
    path = tmp_path / "left_print_area.xlsx"
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws["A1"] = "x"
    ws.print_area = "$A$1:$C$5"
    wb.save(path)

    with Book(path) as bk:
        bk.sheets[0].delete_columns(11)
        out = _saved(bk, tmp_path)
    assert out.print_area.endswith("$A$1:$C$5")


# ===================================================================
# fail fast
# ===================================================================


def _sheet_snapshot(path):
    ws = load_workbook(path).active
    return (
        [(c.coordinate, c.value) for row in ws.iter_rows() for c in row],
        sorted(str(r) for r in ws.merged_cells.ranges),
        {k: v.width for k, v in ws.column_dimensions.items()},
    )


def test_formula_rejects_delete_with_coordinate(tmp_path):
    path = tmp_path / "with_formula.xlsx"
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws["A1"] = 1
    ws["M5"] = "=SUM(A1:A3)"
    wb.save(path)
    before = _sheet_snapshot(path)

    with Book(path) as bk:
        with pytest.raises(StructuralEditError) as excinfo:
            bk.sheets[0].delete_columns(11)
        assert "M5" in str(excinfo.value)
        assert "公式" in str(excinfo.value)

    assert _sheet_snapshot(path) == before  # 來源檔未被改動


def test_rejects_images(tmp_path):
    pytest.importorskip("PIL", reason="openpyxl 需要 Pillow 才能嵌入圖片")
    png = tmp_path / "dot.png"
    png.write_bytes(
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01"
        b"\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    path = tmp_path / "with_image.xlsx"
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws.add_image(OpenpyxlImage(str(png)), "B2")
    wb.save(path)

    with Book(path) as bk:
        with pytest.raises(StructuralEditError, match="圖片"):
            bk.sheets[0].delete_columns(11)


def test_rejects_excel_tables(tmp_path):
    path = tmp_path / "with_table.xlsx"
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws.append(["a", "b"])
    ws.append([1, 2])
    ws.add_table(Table(displayName="T1", ref="A1:B2"))
    wb.save(path)

    with Book(path) as bk:
        with pytest.raises(StructuralEditError, match="表格"):
            bk.sheets[0].delete_columns(11)


def test_rejects_conditional_formatting(tmp_path):
    path = tmp_path / "with_cf.xlsx"
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws["A1"] = 5
    ws.conditional_formatting.add(
        "A1:A10", CellIsRule(operator="greaterThan", formula=["3"])
    )
    wb.save(path)

    with Book(path) as bk:
        with pytest.raises(StructuralEditError, match="條件格式"):
            bk.sheets[0].delete_columns(11)


def test_rejects_data_validation(tmp_path):
    path = tmp_path / "with_dv.xlsx"
    wb = OpenpyxlWorkbook()
    ws = wb.active
    dv = DataValidation(type="list", formula1='"甲,乙"')
    ws.add_data_validation(dv)
    dv.add("A1:A10")
    wb.save(path)

    with Book(path) as bk:
        with pytest.raises(StructuralEditError, match="資料驗證"):
            bk.sheets[0].delete_columns(11)


# ===================================================================
# 模板保留內容（preserve-template-drawings）：形狀／openpyxl 會丟棄的圖片
# 拒絕刪欄；頁首頁尾圖片不阻擋刪欄。
# ===================================================================


def _extract_header_footer_image_from_real_template():
    """取出 `template_non_table.xlsx` 工作表12 的頁首圖片 VML／rels／媒體。

    供「僅含頁首頁尾圖片時刪欄不受阻擋」測試注入乾淨模板，不依賴黃金基準。
    """
    import zipfile

    from tests.fixtures import TEMPLATE_NON_TABLE
    from tests.golden_compare import (
        _zip_rels_map,
        _zip_rels_path,
        _zip_resolve,
        _zip_sheet_part_map,
    )

    with zipfile.ZipFile(TEMPLATE_NON_TABLE) as z:
        names = set(z.namelist())
        sheet_part = _zip_sheet_part_map(z, names)["工作表12"]
        vml_part = next(
            _zip_resolve(sheet_part, rel.get("Target"))
            for rel in _zip_rels_map(z, names, sheet_part).values()
            if rel.get("Type", "").endswith("/vmlDrawing")
        )
        vml_bytes = z.read(vml_part)
        vml_rels_bytes = z.read(_zip_rels_path(vml_part))
        media = {}
        for rel in _zip_rels_map(z, names, vml_part).values():
            target = _zip_resolve(vml_part, rel.get("Target"))
            media[target.rsplit("/", 1)[-1]] = z.read(target)
    return vml_bytes, vml_rels_bytes, media


def _inject_header_footer_image(xlsx_path, vml_bytes, vml_rels_xml, media: dict) -> None:
    """把頁首頁尾 VML（含其 rels 與媒體）以 zip 手術塞進既有 xlsx（比照 legacyDrawingHF）。"""
    import shutil
    import zipfile

    src = str(xlsx_path) + ".orig"
    shutil.move(str(xlsx_path), src)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(xlsx_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = payload.decode("utf-8")
                if "xmlns:r=" not in text:
                    text = text.replace(
                        "<worksheet ",
                        '<worksheet xmlns:r="http://schemas.openxmlformats.org/'
                        'officeDocument/2006/relationships" ',
                        1,
                    )
                text = text.replace(
                    "</worksheet>",
                    "<headerFooter><oddHeader>&amp;C&amp;G</oddHeader></headerFooter>"
                    '<legacyDrawingHF r:id="rIdVml1"/></worksheet>',
                )
                payload = text.encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                text = payload.decode("utf-8")
                extras = ""
                if 'Extension="vml"' not in text:
                    extras += (
                        '<Default Extension="vml" ContentType='
                        '"application/vnd.openxmlformats-officedocument.vmlDrawing"/>'
                    )
                if 'Extension="png"' not in text and any(n.endswith(".png") for n in media):
                    extras += '<Default Extension="png" ContentType="image/png"/>'
                payload = text.replace("</Types>", extras + "</Types>").encode("utf-8")
            zout.writestr(item, payload)
        zout.writestr("xl/drawings/vmlDrawing1.vml", vml_bytes)
        zout.writestr("xl/drawings/_rels/vmlDrawing1.vml.rels", vml_rels_xml)
        for name, data in media.items():
            zout.writestr(f"xl/media/{name}", data)
        zout.writestr(
            "xl/worksheets/_rels/sheet1.xml.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rIdVml1"'
            ' Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/vmlDrawing"'
            ' Target="../drawings/vmlDrawing1.vml"/></Relationships>',
        )


def test_rejects_template_shapes(tmp_path):
    """僅含文字方塊（openpyxl 無法表示）時拒絕刪欄，訊息含「形狀」與錨點座標。"""
    from tests.test_drawing_preservation import _anchor_rows, _inject_shape

    path = tmp_path / "with_shape.xlsx"
    wb = OpenpyxlWorkbook()
    wb.active["A1"] = "no tag"
    wb.save(path)
    _inject_shape(path, 2, 6)  # 錨點左上角：col=1、row=2（0-based）→ B3

    out = tmp_path / "out_shape.xlsx"
    with Book(path) as bk:
        with pytest.raises(StructuralEditError) as excinfo:
            bk.sheets[0].delete_columns(2)
        assert "形狀" in str(excinfo.value)
        assert "B3" in str(excinfo.value)
        bk.save(out)  # 拒絕後未做任何修改，仍可正常存檔

    assert _anchor_rows(out) == [(2, 6)]  # 形狀仍位於模板原始位置


def test_rejects_openpyxl_dropped_images(tmp_path):
    """僅含 openpyxl 會丟棄的圖片（EMF）時拒絕刪欄，訊息含「圖片」與錨點座標。"""
    from tests.test_drawing_preservation import (
        _drawing_rels_xml,
        _inject_drawing,
        _minimal_emf_bytes,
        _pic_anchor_xml,
        _wsdr_xml,
    )

    path = tmp_path / "with_emf.xlsx"
    wb = OpenpyxlWorkbook()
    wb.active["A1"] = "no tag"
    wb.save(path)
    _inject_drawing(
        path,
        _wsdr_xml(_pic_anchor_xml("rId1", 2, "EmfPic")),  # 錨點左上角：col=1、row=2 → B3
        _drawing_rels_xml({"rId1": "image1.emf"}),
        {"image1.emf": _minimal_emf_bytes()},
    )

    with Book(path) as bk:
        with pytest.raises(StructuralEditError) as excinfo:
            bk.sheets[0].delete_columns(2)
        assert "圖片" in str(excinfo.value)
        assert "B3" in str(excinfo.value)


def test_header_footer_image_does_not_block_delete(tmp_path):
    """僅含頁首頁尾圖片時刪欄成功、不拋例外，存檔後頁首頁尾圖片保留。"""
    from tests.test_drawing_preservation import _inspect_sheet

    path = tmp_path / "with_header_image.xlsx"
    wb = OpenpyxlWorkbook()
    wb.active["A1"] = "no tag"
    wb.save(path)
    vml_bytes, vml_rels_xml, media = _extract_header_footer_image_from_real_template()
    _inject_header_footer_image(path, vml_bytes, vml_rels_xml, media)

    before = _inspect_sheet(path, "Sheet")
    out = tmp_path / "out_header_image.xlsx"
    with Book(path) as bk:
        bk.sheets[0].delete_columns(2)  # 不應拋例外
        bk.save(out)

    after = _inspect_sheet(out, "Sheet")
    assert len(after["header_images"]) == 1
    assert after["header_images"][0]["media_hash"] == before["header_images"][0]["media_hash"]


def test_clean_sheet_deletes_without_error(template, tmp_path):
    with Book(template) as bk:
        bk.sheets[0].delete_columns(11)  # 不應拋例外
        ws = _saved(bk, tmp_path)
    assert ws.cell(row=9, column=11).value == 12345


def test_invalid_index_raises_value_error(template):
    with Book(template) as bk:
        with pytest.raises(ValueError):
            bk.sheets[0].delete_columns(0)


# ===================================================================
# 與渲染組合
# ===================================================================


def test_delete_after_render(tmp_path):
    """渲染 → 刪欄 → 存檔：單次 load/save 完成完整後處理流程。"""
    import pandas as pd

    path = tmp_path / "render_then_delete.xlsx"
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws["A1"] = "{{title}}"
    ws.merge_cells("A1:C1")
    ws["A3"] = "#{{df | noheader}}"
    wb.save(path)

    out = tmp_path / "out.xlsx"
    with Book(path) as bk:
        bk.render({"title": "報表", "df": pd.DataFrame([[1, 2, 3], [4, 5, 6]])})
        bk.sheets[0].delete_columns(2)
        bk.save(out)

    ws = load_workbook(out).active
    assert ws["A1"].value == "報表"
    assert "A1:B1" in {str(r) for r in ws.merged_cells.ranges}
    assert [ws.cell(row=3, column=c).value for c in (1, 2)] == [1, 3]
