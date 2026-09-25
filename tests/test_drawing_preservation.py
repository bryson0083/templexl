"""模板圖形（AutoShape）保留契約測試。

openpyxl 只模型化 pic 與 chart，``<xdr:sp>`` 這類圖形在 load/save round-trip
會被無聲丟棄（openpyxl/reader/drawings.py 的 find_images 只回傳 charts 與
images）。實務上大小章方框、註記文字方塊都是這種圖形，掉了不會報錯、
只會產出缺框的報表。

本檔驗證 templexl 會把這些圖形原樣帶到輸出，並在表格插列時比照 Excel
的插入語意平移錨點。
"""

from __future__ import annotations

import hashlib
import posixpath
import shutil
import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook as OpenpyxlWorkbook

from templexl import render
from templexl.book import Book
from templexl.exceptions import RenderError
from tests.fixtures import (
    TEMPLATE_NON_TABLE,
    TEMPLATE_TABLE,
    non_table_data,
    table_data,
)

DRAWING_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"\
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><xdr:twoCellAnchor>\
<xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{from_row}</xdr:row>\
<xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>3</xdr:col><xdr:colOff>0</xdr:colOff>\
<xdr:row>{to_row}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>\
<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="2" name="矩形 1"/><xdr:cNvSpPr/>\
</xdr:nvSpPr><xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1573530" cy="1693545"/></a:xfrm>\
<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:schemeClr val="bg1"/></a:solidFill>\
<a:ln w="12700"><a:solidFill><a:schemeClr val="tx1"/></a:solidFill></a:ln></xdr:spPr>\
<xdr:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="zh-TW"/></a:p></xdr:txBody></xdr:sp>\
<xdr:clientData/></xdr:twoCellAnchor></xdr:wsDr>"""

SHEET_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">\
<Relationship Id="rIdShape1"\
 Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"\
 Target="../drawings/drawing1.xml"/></Relationships>"""

DRAWING_CT = ('<Override PartName="/xl/drawings/drawing1.xml"'
              ' ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>')


def _inject_shape(xlsx_path, from_row: int, to_row: int) -> None:
    """把一個矩形圖形以 zip 手術塞進既有 xlsx（openpyxl 無法建立圖形）。

    from_row / to_row 為 0-based 錨點列號，與 OOXML 相同。
    """
    src = str(xlsx_path) + ".orig"
    shutil.move(str(xlsx_path), src)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(xlsx_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = payload.decode("utf-8")
                if "xmlns:r=" not in text:
                    text = text.replace(
                        "<worksheet ",
                        '<worksheet xmlns:r="http://schemas.openxmlformats.org/'
                        'officeDocument/2006/relationships" ',
                        1,
                    )
                text = text.replace("</worksheet>", '<drawing r:id="rIdShape1"/></worksheet>')
                payload = text.encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                text = payload.decode("utf-8")
                payload = text.replace("</Types>", DRAWING_CT + "</Types>").encode("utf-8")
            zout.writestr(item, payload)
        zout.writestr("xl/drawings/drawing1.xml",
                      DRAWING_XML.format(from_row=from_row, to_row=to_row))
        zout.writestr("xl/worksheets/_rels/sheet1.xml.rels", SHEET_RELS)


NS_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"


def _anchor_rows(xlsx_path):
    """讀出輸出檔中所有「圖形」錨點的 (from_row, to_row)，0-based。"""
    from lxml import etree

    rows = []
    with zipfile.ZipFile(xlsx_path) as z:
        for name in z.namelist():
            if not (name.startswith("xl/drawings/") and name.endswith(".xml")):
                continue
            root = etree.fromstring(z.read(name))
            for anchor in root:
                if anchor.find(f".//{{{NS_XDR}}}sp") is None:
                    continue
                found = [int(el.text) for el in anchor.iter(f"{{{NS_XDR}}}row")]
                rows.append((found[0], found[1]))
    return rows


def _text_row(xlsx_path, needle, sheet_name=None):
    """回傳 A 欄等於 needle 的列號（1-based）；``sheet_name`` 省略時用作用中工作表。"""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path)
    ws = wb[sheet_name] if sheet_name else wb.active
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == needle:
            return r
    return None


@pytest.fixture
def template_with_shape(tmp_path):
    """A3 有表格標籤、第 9 列起有一個矩形圖形（模擬大小章方框）。"""
    def _make(shape_from_row: int, shape_to_row: int, with_tag: bool = True):
        wb = OpenpyxlWorkbook()
        ws = wb.active
        ws["A1"] = "報表 {{title}}"
        if with_tag:
            ws["A3"] = "#{{df | noheader}}"
        ws["A9"] = "客運公司大小章："
        path = tmp_path / f"tpl_{shape_from_row}_{int(with_tag)}.xlsx"
        wb.save(path)
        _inject_shape(path, shape_from_row, shape_to_row)
        return path
    return _make


def _df(n: int) -> pd.DataFrame:
    return pd.DataFrame({"a": [f"r{i}" for i in range(n)], "b": list(range(n))})


def test_template_fixture_really_contains_shape(template_with_shape):
    """先確認 fixture 本身有效：模板裡真的有圖形。"""
    tpl = template_with_shape(8, 16)
    assert _anchor_rows(tpl) == [(8, 16)]


def test_shape_survives_render_without_table_expansion(template_with_shape, tmp_path):
    """無表格展開時，圖形應原樣保留、錨點不動。"""
    tpl = template_with_shape(8, 16, with_tag=False)
    out = tmp_path / "out_no_tag.xlsx"
    render(str(tpl), str(out), data={"title": "X"})
    assert _anchor_rows(out) == [(8, 16)], "圖形在渲染後消失或位移了"


def test_shape_below_table_shifts_with_inserted_rows(template_with_shape, tmp_path):
    """圖形在表格標籤下方時，錨點應隨插列平移（Excel 插入語意）。"""
    tpl = template_with_shape(8, 16)
    out = tmp_path / "out_shift.xlsx"
    render(str(tpl), str(out), data={"title": "X", "df": _df(5)})
    # 標籤在第 3 列（1-based），5 列資料（noheader）→ 插入 4 列
    assert _anchor_rows(out) == [(12, 20)]
    # 圖形位移量必須與標籤下方文字的實際位移量一致（第 9 列 → 第 13 列）
    assert _text_row(out, "客運公司大小章：") == 13


def test_shape_above_table_not_shifted(template_with_shape, tmp_path):
    """圖形在表格標籤上方時不應被平移。"""
    tpl = template_with_shape(0, 1)
    out = tmp_path / "out_above.xlsx"
    render(str(tpl), str(out), data={"title": "X", "df": _df(5)})
    assert _anchor_rows(out) == [(0, 1)]


def test_book_save_also_preserves_shape(template_with_shape, tmp_path):
    """Book 會話路徑與 render() 行為一致。"""
    tpl = template_with_shape(8, 16)
    out = tmp_path / "out_book.xlsx"
    with Book(str(tpl)) as bk:
        bk.render({"title": "X", "df": _df(5)})
        bk.save(str(out))
    assert _anchor_rows(out) == [(12, 20)]


NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_V = "urn:schemas-microsoft-com:vml"
NS_O = "urn:schemas-microsoft-com:office:office"


def assert_package_integrity(xlsx_path):
    """檢查 OPC 套件自洽：每個 r:id 都解析得到、目標 part 存在、宣告了內容型別。

    openpyxl 讀檔非常寬鬆，缺 rels 也照樣載入；Excel 則會直接拒絕開啟。
    因此「渲染後 openpyxl 還讀得到圖形」不足以證明檔案有效，必須驗套件本身。
    """
    import posixpath
    from lxml import etree

    with zipfile.ZipFile(xlsx_path) as z:
        names = set(z.namelist())
        content_types = etree.fromstring(z.read("[Content_Types].xml"))
        overrides = {el.get("PartName")
                     for el in content_types.iter(f"{{{NS_CT}}}Override")}
        defaults = {el.get("Extension", "").lower()
                    for el in content_types.iter(f"{{{NS_CT}}}Default")}

        for part in sorted(n for n in names
                           if n.startswith("xl/worksheets/") and n.endswith(".xml")):
            root = etree.fromstring(z.read(part))
            rels_part = posixpath.join(posixpath.dirname(part), "_rels",
                                       posixpath.basename(part) + ".rels")
            rels = {}
            if rels_part in names:
                for rel in etree.fromstring(z.read(rels_part)):
                    rels[rel.get("Id")] = rel.get("Target")

            used = {el.get(f"{{{NS_REL}}}id")
                    for el in root.iter()
                    if el.get(f"{{{NS_REL}}}id")}
            missing = used - set(rels)
            assert not missing, (
                f"{part} 參照了不存在的關聯 {sorted(missing)}；"
                f"rels={rels_part if rels_part in names else '(缺檔)'}"
            )

            for rid in used:
                target = rels[rid]
                if target.startswith(("http://", "https://")):
                    continue
                resolved = _resolve(part, target)
                assert resolved in names, f"{part} 的 {rid} 指向不存在的 part {resolved}"
                ext = posixpath.splitext(resolved)[1].lstrip(".").lower()
                assert (f"/{resolved}" in overrides or ext in defaults), (
                    f"part {resolved} 未在 [Content_Types].xml 宣告內容型別")

        # VML 部件（頁首頁尾圖片）：o:relid 皆須解析得到關聯，媒體 part 存在，
        # 且 vml 副檔名須在 [Content_Types].xml 宣告——openpyxl 原生輸出不含
        # VML，缺這個 Default 會讓 Excel 判定檔案損毀。
        for part in sorted(n for n in names
                           if n.startswith("xl/drawings/") and n.endswith(".vml")):
            ext = posixpath.splitext(part)[1].lstrip(".").lower()
            assert (f"/{part}" in overrides or ext in defaults), (
                f"part {part} 未在 [Content_Types].xml 宣告內容型別")

            root = etree.fromstring(z.read(part))
            rels_part = posixpath.join(posixpath.dirname(part), "_rels",
                                       posixpath.basename(part) + ".rels")
            rels = {}
            if rels_part in names:
                for rel in etree.fromstring(z.read(rels_part)):
                    rels[rel.get("Id")] = rel.get("Target")

            used = {el.get(f"{{{NS_O}}}relid")
                    for el in root.iter(f"{{{NS_V}}}imagedata")
                    if el.get(f"{{{NS_O}}}relid")}
            missing = used - set(rels)
            assert not missing, (
                f"{part} 參照了不存在的關聯 {sorted(missing)}；"
                f"rels={rels_part if rels_part in names else '(缺檔)'}"
            )
            for rid in used:
                resolved = _resolve(part, rels[rid])
                assert resolved in names, f"{part} 的 {rid} 指向不存在的 part {resolved}"
                media_ext = posixpath.splitext(resolved)[1].lstrip(".").lower()
                assert (f"/{resolved}" in overrides or media_ext in defaults), (
                    f"part {resolved} 未在 [Content_Types].xml 宣告內容型別")


def test_output_package_is_valid_for_excel(template_with_shape, tmp_path):
    """回填後的檔案必須是自洽的 OPC 套件（Excel 才打得開）。

    回歸測試：曾因新建的 sheet rels 沒被寫進 zip，導致 <drawing r:id> 懸空，
    openpyxl 照樣讀得到圖形、Excel 卻拒絕開啟。
    """
    tpl = template_with_shape(8, 16)
    out = tmp_path / "out_valid.xlsx"
    render(str(tpl), str(out), data={"title": "X", "df": _df(5)})
    assert_package_integrity(out)


def test_output_sheet_rels_written_for_new_drawing(template_with_shape, tmp_path):
    """新建 drawing 時，工作表的 _rels 必須實際落進 zip。"""
    tpl = template_with_shape(8, 16)
    out = tmp_path / "out_rels.xlsx"
    render(str(tpl), str(out), data={"title": "X", "df": _df(5)})
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        sheet_rels = [n for n in names if n.startswith("xl/worksheets/_rels/")]
    assert sheet_rels, "工作表 rels 檔沒有被寫入，<drawing r:id> 會懸空"



# ======================================================================
# 對照黃金模板（tests/test_templates/）的正確性測試
#
# 與上方以合成模板驗證「機制」不同，這裡直接對真實模板（含形狀、
# EMF 無關但含頁首頁尾圖片與文字浮水印）驗證「結果與模板一致」——
# 黃金基準只能防回歸，正確與否要靠這裡直接比對模板本身。
# ======================================================================


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _resolve(base_part, target):
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_part), target))


def _rels_map(z, names, part):
    from lxml import etree

    rels_part = posixpath.join(posixpath.dirname(part), "_rels",
                               posixpath.basename(part) + ".rels")
    if rels_part not in names:
        return {}
    return {rel.get("Id"): rel for rel in etree.fromstring(z.read(rels_part))}


def _sheet_part_map(z, names):
    from lxml import etree

    wb = etree.fromstring(z.read("xl/workbook.xml"))
    wb_rels = _rels_map(z, names, "xl/workbook.xml")
    mapping = {}
    for sheet in wb.iter(f"{{{NS_MAIN}}}sheet"):
        rel = wb_rels.get(sheet.get(f"{{{NS_REL}}}id"))
        if rel is not None:
            mapping[sheet.get("name")] = _resolve("xl/workbook.xml", rel.get("Target"))
    return mapping


def _sheet_drawing_and_vml_parts(z, names, sheet_part):
    drawing = vml = None
    for rel in _rels_map(z, names, sheet_part).values():
        rel_type = rel.get("Type", "")
        if rel_type.endswith("/drawing"):
            drawing = _resolve(sheet_part, rel.get("Target"))
        elif rel_type.endswith("/vmlDrawing"):
            vml = _resolve(sheet_part, rel.get("Target"))
    return drawing, vml


def _inspect_sheet(xlsx_path, sheet_name):
    """回傳工作表的形狀／圖片／頁首頁尾圖片摘要，供直接比對模板。

    Returns:
        dict: ``{"shapes": [...], "pics": [...], "header_text": str|None,
        "header_images": [...]}``；``shapes``／``pics`` 內元素含 ``from``
        （0-based 列欄）與內容雜湊／文字／旋轉角度。
    """
    from lxml import etree

    with zipfile.ZipFile(xlsx_path) as z:
        names = set(z.namelist())
        sheet_part = _sheet_part_map(z, names).get(sheet_name)
        assert sheet_part, f"找不到工作表 {sheet_name!r}"
        drawing_part, vml_part = _sheet_drawing_and_vml_parts(z, names, sheet_part)

        shapes, pics = [], []
        if drawing_part and drawing_part in names:
            root = etree.fromstring(z.read(drawing_part))
            rels = _rels_map(z, names, drawing_part)
            for anchor in root:
                frm_el = anchor.find(f"{{{NS_XDR}}}from")
                frm = None
                if frm_el is not None:
                    frm = (int(frm_el.find(f"{{{NS_XDR}}}row").text),
                           int(frm_el.find(f"{{{NS_XDR}}}col").text))
                blip = anchor.find(f".//{{{NS_A}}}blip")
                media_hash = None
                if blip is not None:
                    rel = rels.get(blip.get(f"{{{NS_REL}}}embed"))
                    if rel is not None:
                        target = _resolve(drawing_part, rel.get("Target"))
                        if target in names:
                            media_hash = _sha256(z.read(target))

                if anchor.find(f".//{{{NS_XDR}}}sp") is not None:
                    xfrm = anchor.find(f".//{{{NS_A}}}xfrm")
                    rotation = int(xfrm.get("rot", "0")) if xfrm is not None else 0
                    text = "".join(t.text or "" for t in anchor.iter(f"{{{NS_A}}}t"))
                    shapes.append({
                        "from": frm, "rotation": rotation, "text": text,
                        "media_hash": media_hash,
                    })
                elif anchor.find(f".//{{{NS_XDR}}}pic") is not None:
                    pics.append({"from": frm, "media_hash": media_hash})

        header_text = None
        sheet_xml = etree.fromstring(z.read(sheet_part))
        odd_header = sheet_xml.find(f"{{{NS_MAIN}}}headerFooter/{{{NS_MAIN}}}oddHeader")
        if odd_header is not None and odd_header.text is not None:
            # openpyxl 寫出時把換行編碼為 _x000a_（Excel 讀回等同字面換行，
            # 見 design.md context：「換行寫成 _x000a_，Excel 讀回正確」）；
            # 正規化後才能與模板原始 XML 的字面換行比較語意是否相同。
            header_text = odd_header.text.replace("_x000a_", "\n")

        header_images = []
        if vml_part and vml_part in names:
            vml_root = etree.fromstring(z.read(vml_part))
            vml_rels = _rels_map(z, names, vml_part)
            for shape in vml_root.iter(f"{{{NS_V}}}shape"):
                imagedata = shape.find(f"{{{NS_V}}}imagedata")
                if imagedata is None:
                    continue
                rel = vml_rels.get(imagedata.get(f"{{{NS_O}}}relid"))
                media_hash = None
                if rel is not None:
                    target = _resolve(vml_part, rel.get("Target"))
                    if target in names:
                        media_hash = _sha256(z.read(target))
                header_images.append({"position": shape.get("id"), "media_hash": media_hash})

    return {
        "shapes": shapes,
        "pics": pics,
        "header_text": header_text,
        "header_images": header_images,
    }


def test_sheet_without_table_tags_shape_coexists_with_images(tmp_path):
    """non_table 工作表11：無表格標籤，形狀與兩張圖片共存、錨點不動（design 決策 3）。"""
    out = tmp_path / "non_table_ws11.xlsx"
    render(str(TEMPLATE_NON_TABLE), str(out), data=non_table_data())

    before = _inspect_sheet(TEMPLATE_NON_TABLE, "工作表11")
    after = _inspect_sheet(out, "工作表11")

    assert len(after["shapes"]) == 1
    assert len(after["pics"]) == 2
    assert after["shapes"][0]["from"] == before["shapes"][0]["from"]
    assert after["shapes"][0]["rotation"] == before["shapes"][0]["rotation"] != 0
    assert after["shapes"][0]["text"] == before["shapes"][0]["text"] == "Watermark Demo"


def test_shape_between_table_tags_shifts_like_report_footer(tmp_path):
    """table 工作表8：形狀夾在兩個表格標籤之間，下移量與 Report Footer、Picture 2 相同。

    對照 spec「保留內容的錨點位移／形狀夾在兩個表格標籤之間」：夾在中間的形狀
    下移「整表插入列數總和」，而非只受上方那個表格標籤影響。
    """
    out = tmp_path / "table_ws8.xlsx"
    render(str(TEMPLATE_TABLE), str(out), data=table_data())

    before = _inspect_sheet(TEMPLATE_TABLE, "工作表8")
    after = _inspect_sheet(out, "工作表8")

    assert len(after["shapes"]) == 1
    assert len(after["pics"]) == 2
    assert after["shapes"][0]["text"] == before["shapes"][0]["text"] == "Watermark Demo"
    assert after["shapes"][0]["rotation"] == before["shapes"][0]["rotation"] != 0

    footer_before = _text_row(TEMPLATE_TABLE, "Report Footer", "工作表8")
    footer_after = _text_row(out, "Report Footer", "工作表8")
    shift = footer_after - footer_before
    assert shift == 8, "目前 table_data() 資料量下 Report Footer 的實際下移量"

    assert after["shapes"][0]["from"][0] - before["shapes"][0]["from"][0] == shift

    # Picture 2：模板中列 10（0-based）起的較大圖片，以媒體雜湊在輸出中定位。
    before_pic2 = next(p for p in before["pics"] if p["from"][0] == 10)
    after_pic2 = next(
        p for p in after["pics"] if p["media_hash"] == before_pic2["media_hash"]
    )
    assert after_pic2["from"][0] - before_pic2["from"][0] == shift


def test_header_footer_image_center_matches_template(tmp_path):
    """table 工作表9：置中頁首圖片的媒體雜湊與頁首文字與模板相同。"""
    out = tmp_path / "table_ws9.xlsx"
    render(str(TEMPLATE_TABLE), str(out), data=table_data())

    before = _inspect_sheet(TEMPLATE_TABLE, "工作表9")
    after = _inspect_sheet(out, "工作表9")

    assert len(after["header_images"]) == 1
    assert after["header_images"][0]["position"] == "CH"
    assert after["header_images"][0]["media_hash"] == before["header_images"][0]["media_hash"]
    assert after["header_text"] == before["header_text"]


def test_header_footer_image_left_matches_template(tmp_path):
    """non_table 工作表12：靠左頁首圖片的媒體雜湊與頁首文字與模板相同。"""
    out = tmp_path / "non_table_ws12.xlsx"
    render(str(TEMPLATE_NON_TABLE), str(out), data=non_table_data())

    before = _inspect_sheet(TEMPLATE_NON_TABLE, "工作表12")
    after = _inspect_sheet(out, "工作表12")

    assert len(after["header_images"]) == 1
    assert after["header_images"][0]["position"] == "LH"
    assert after["header_images"][0]["media_hash"] == before["header_images"][0]["media_hash"]
    assert after["header_text"] == before["header_text"]


def test_header_footer_image_signature_unaffected_by_cell_comment(tmp_path):
    """工作表同時有頁首圖片與儲存格註解時，golden_compare 的
    header_footer_images 簽章仍須正確指向頁首頁尾 VML，而非誤抓 openpyxl
    為註解寫出的 VML——兩者的關聯 Type 相同（皆為 vmlDrawing），只取
    「第一個符合類型的關聯」會抓錯。
    """
    from openpyxl.comments import Comment

    from tests.golden_compare import workbook_signature

    out = tmp_path / "table_ws9_comment.xlsx"
    with Book(str(TEMPLATE_TABLE)) as bk:
        bk.render(table_data())
        ws9 = bk.workbook["工作表9"]
        ws9["A1"].comment = Comment("測試註解", "tester")  # 逃生門新增註解
        bk.save(str(out))

    before = _inspect_sheet(TEMPLATE_TABLE, "工作表9")
    sig = workbook_signature(out)["sheets"]["工作表9"]["header_footer_images"]

    assert len(sig) == 1
    assert sig[0]["position"] == "CH"
    assert sig[0]["media_hash"] == before["header_images"][0]["media_hash"]


def test_golden_templates_output_pass_package_integrity(tmp_path):
    """兩份黃金模板的渲染輸出皆須是自洽的 OPC 套件（含 VML 檢查）。"""
    out_table = tmp_path / "table_pkg.xlsx"
    out_non_table = tmp_path / "non_table_pkg.xlsx"
    render(str(TEMPLATE_TABLE), str(out_table), data=table_data())
    render(str(TEMPLATE_NON_TABLE), str(out_non_table), data=non_table_data())
    assert_package_integrity(out_table)
    assert_package_integrity(out_non_table)



# ======================================================================
# openpyxl 會丟棄的圖片（EMF／WMF）與媒體命名衝突
#
# openpyxl 對「Pillow 開不了、或開得了但 format 是 WMF」的內嵌圖片一律丟棄
# （reader/drawings.py 的 find_images()）。EMF 對 Pillow 而言就是這種圖片：
# _accept() 只認 record type 1 + offset 40 的 " EMF" 簽章，開啟後一律回報
# format == "WMF"（EMF 是 WMF 的後繼格式，Pillow 共用同一個 plugin 識別）。
# 以下用最小合法標頭建構 EMF，不需要真的畫出圖形內容。
# ======================================================================


def _minimal_emf_bytes() -> bytes:
    """最小 ENHMETAHEADER：record type 1、offset 40 為 b" EMF"、bounds/frame 非零。"""
    import struct

    header = struct.pack(
        "<II4i4i4sIIIHHIII2i2i",
        1, 88,                      # iType=EMR_HEADER, nSize
        0, 0, 100, 100,             # rclBounds（裝置單位，非零）
        0, 0, 2540, 2540,           # rclFrame（.01mm 單位，非零）
        b" EMF", 0x00010000,        # dSignature, nVersion
        88, 1,                      # nBytes, nRecords
        0, 0,                       # nHandles, sReserved
        0, 0,                       # nDescription, offDescription
        0,                          # nPalEntries
        100, 100,                   # szlDevice
        26, 26,                     # szlMillimeters
    )
    assert len(header) == 88
    assert header[40:44] == b" EMF"
    return header


def _tiny_png_bytes(color) -> bytes:
    import io

    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (2, 2), color).save(buf, format="PNG")
    return buf.getvalue()


def _make_base_workbook(path) -> None:
    """無任何標籤、無任何圖形的最小工作簿（供 zip 手術注入繪圖）。"""
    wb = OpenpyxlWorkbook()
    wb.active["A1"] = "no tag"
    wb.save(path)


def _pic_anchor_xml(rid, row, name="Pic") -> str:
    return (
        f'<xdr:twoCellAnchor><xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff>'
        f'<xdr:row>{row}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
        f'<xdr:to><xdr:col>3</xdr:col><xdr:colOff>0</xdr:colOff>'
        f'<xdr:row>{row + 4}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
        f'<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="{row + 2}" name="{name}"/>'
        f'<xdr:cNvPicPr/></xdr:nvPicPr><xdr:blipFill>'
        f'<a:blip xmlns:r="{NS_REL}" r:embed="{rid}"/>'
        f'<a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
        f'<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="200000" cy="200000"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:pic>'
        f'<xdr:clientData/></xdr:twoCellAnchor>'
    )


def _filled_shape_anchor_xml(rid, row, name="FilledShape") -> str:
    """以圖片填滿（``<a:blipFill>`` 內 ``r:embed``）的矩形形狀。"""
    return (
        f'<xdr:oneCellAnchor><xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff>'
        f'<xdr:row>{row}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
        f'<xdr:ext cx="300000" cy="300000"/>'
        f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{row + 2}" name="{name}"/>'
        f'<xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr><a:xfrm><a:off x="0" y="0"/>'
        f'<a:ext cx="300000" cy="300000"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f'<a:blipFill><a:blip xmlns:r="{NS_REL}" r:embed="{rid}"/>'
        f'<a:stretch><a:fillRect/></a:stretch></a:blipFill></xdr:spPr>'
        f'<xdr:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="zh-TW"/></a:p></xdr:txBody>'
        f'</xdr:sp><xdr:clientData/></xdr:oneCellAnchor>'
    )


def _wsdr_xml(anchors_xml: str) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<xdr:wsDr xmlns:xdr="{NS_XDR}" xmlns:a="{NS_A}">{anchors_xml}</xdr:wsDr>'
    )


def _drawing_rels_xml(entries: dict) -> str:
    body = "".join(
        f'<Relationship Id="{rid}" Type="{NS_REL}/image" Target="../media/{name}"/>'
        for rid, name in entries.items()
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Relationships xmlns="{NS_PKG_REL}">{body}</Relationships>'
    )


_EXTRA_MEDIA_CONTENT_TYPES = {"png": "image/png", "emf": "image/x-emf"}


def _inject_drawing(xlsx_path, drawing_xml: str, rels_xml: str, media: dict) -> None:
    """把一份完整 drawing1.xml（含其 rels 與媒體）以 zip 手術塞進既有 xlsx。

    與 :func:`_inject_shape` 的差異：呼叫端自行組好完整的錨點清單（可混合
    ``pic``／``sp``、任意數量），供 EMF、媒體命名衝突等測試共用。
    """
    src = str(xlsx_path) + ".orig"
    shutil.move(str(xlsx_path), src)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(xlsx_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text_ = payload.decode("utf-8")
                if "xmlns:r=" not in text_:
                    text_ = text_.replace(
                        "<worksheet ",
                        '<worksheet xmlns:r="http://schemas.openxmlformats.org/'
                        'officeDocument/2006/relationships" ',
                        1,
                    )
                text_ = text_.replace("</worksheet>", '<drawing r:id="rIdDrawing1"/></worksheet>')
                payload = text_.encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                text_ = payload.decode("utf-8")
                extras = DRAWING_CT
                for ext, ct in _EXTRA_MEDIA_CONTENT_TYPES.items():
                    if any(name.endswith(f".{ext}") for name in media):
                        extras += f'<Default Extension="{ext}" ContentType="{ct}"/>'
                payload = text_.replace("</Types>", extras + "</Types>").encode("utf-8")
            zout.writestr(item, payload)
        zout.writestr("xl/drawings/drawing1.xml", drawing_xml)
        zout.writestr("xl/drawings/_rels/drawing1.xml.rels", rels_xml)
        for name, data in media.items():
            zout.writestr(f"xl/media/{name}", data)
        zout.writestr(
            "xl/worksheets/_rels/sheet1.xml.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Relationships xmlns="{NS_PKG_REL}">'
            '<Relationship Id="rIdDrawing1"'
            f' Type="{NS_REL}/drawing"'
            ' Target="../drawings/drawing1.xml"/></Relationships>',
        )


def test_emf_fixture_recognized_as_wmf_by_pillow():
    """先確認 fixture 本身有效：Pillow 對這個最小 EMF 標頭回報 format == WMF。

    這正是 openpyxl 丟棄它的原因（find_images() 對 format == "WMF" 的圖片
    一律略過），也是本套件需要接手保留它的理由。
    """
    import io

    from PIL import Image

    img = Image.open(io.BytesIO(_minimal_emf_bytes()))
    assert img.format == "WMF"


def test_single_emf_image_preserved(tmp_path):
    """單張 EMF 圖片：openpyxl 完全丟棄，本套件仍需原樣保留，媒體雜湊與模板相同。"""
    tpl = tmp_path / "tpl_emf.xlsx"
    _make_base_workbook(tpl)
    emf = _minimal_emf_bytes()
    _inject_drawing(
        tpl,
        _wsdr_xml(_pic_anchor_xml("rId1", 2, "EmfPic")),
        _drawing_rels_xml({"rId1": "image1.emf"}),
        {"image1.emf": emf},
    )

    out = tmp_path / "out_emf.xlsx"
    render(str(tpl), str(out), data={})

    info = _inspect_sheet(out, "Sheet")
    assert len(info["pics"]) == 1
    assert info["pics"][0]["media_hash"] == _sha256(emf)


def test_png_and_emf_images_coexist_without_duplication(tmp_path):
    """PNG 與 EMF 並存：openpyxl 保留 PNG、本套件補回 EMF，輸出恰各一張。"""
    tpl = tmp_path / "tpl_png_emf.xlsx"
    _make_base_workbook(tpl)
    png = _tiny_png_bytes((10, 200, 10))
    emf = _minimal_emf_bytes()
    _inject_drawing(
        tpl,
        _wsdr_xml(_pic_anchor_xml("rId1", 2, "PngPic") + _pic_anchor_xml("rId2", 10, "EmfPic")),
        _drawing_rels_xml({"rId1": "image1.png", "rId2": "image2.emf"}),
        {"image1.png": png, "image2.emf": emf},
    )

    out = tmp_path / "out_png_emf.xlsx"
    render(str(tpl), str(out), data={})

    info = _inspect_sheet(out, "Sheet")
    assert len(info["pics"]) == 2
    assert {p["media_hash"] for p in info["pics"]} == {_sha256(png), _sha256(emf)}


def _linked_pic_anchor_xml(rid, row, name="LinkedPic") -> str:
    """僅有 ``r:link``（外部連結，非內嵌）的 pic 錨點；openpyxl 的
    ``_blip_rels`` 只認 ``r:embed``，故完全不模型化這種圖片。
    """
    return (
        f'<xdr:twoCellAnchor><xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff>'
        f'<xdr:row>{row}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
        f'<xdr:to><xdr:col>3</xdr:col><xdr:colOff>0</xdr:colOff>'
        f'<xdr:row>{row + 4}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
        f'<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="{row + 2}" name="{name}"/>'
        f'<xdr:cNvPicPr/></xdr:nvPicPr><xdr:blipFill>'
        f'<a:blip xmlns:r="{NS_REL}" r:link="{rid}"/>'
        f'<a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
        f'<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="200000" cy="200000"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:pic>'
        f'<xdr:clientData/></xdr:twoCellAnchor>'
    )


def _external_image_rels_xml(rid, target) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Relationships xmlns="{NS_PKG_REL}">'
        f'<Relationship Id="{rid}" Type="{NS_REL}/image"'
        f' Target="{target}" TargetMode="External"/></Relationships>'
    )


def test_externally_linked_image_kept_without_crashing_save(tmp_path):
    """僅有 r:link（外部連結，TargetMode=External）的圖片：openpyxl 完全不
    模型化（``_blip_rels`` 只認 ``r:embed``），本套件依 design 決策 2 判定為
    需擷取；回填時外部關聯須原樣複製 Target／TargetMode，不得誤走
    ``_place_media`` 到 ``item.media`` 找內部媒體——該媒體不在 zip 內，
    原本會 ``KeyError`` 導致整份存檔失敗（HEAD 版只是靜默少一張圖，本套件
    不該從「靜默丟圖」惡化成「整份報表產不出來」）。連結圖片本身的媒體
    取得仍屬 Non-Goal，這裡只驗證錨點與原外部連結被保留、存檔不拋例外。
    """
    tpl = tmp_path / "tpl_linked.xlsx"
    _make_base_workbook(tpl)
    external_target = "file:///C:/nonexistent/logo.png"
    _inject_drawing(
        tpl,
        _wsdr_xml(_linked_pic_anchor_xml("rId1", 2, "LinkedPic")),
        _external_image_rels_xml("rId1", external_target),
        {},
    )

    out = tmp_path / "out_linked.xlsx"
    render(str(tpl), str(out), data={})

    info = _inspect_sheet(out, "Sheet")
    assert len(info["pics"]) == 1

    with zipfile.ZipFile(out) as z:
        names = set(z.namelist())
        sheet_part = _sheet_part_map(z, names).get("Sheet")
        drawing_part, _ = _sheet_drawing_and_vml_parts(z, names, sheet_part)
        rels = _rels_map(z, names, drawing_part)
        linked = [r for r in rels.values() if r.get("TargetMode") == "External"]
        assert len(linked) == 1
        assert linked[0].get("Target") == external_target
        assert linked[0].get("Type") == f"{NS_REL}/image"

    assert_package_integrity(out)


def test_filled_shape_and_separate_png_media_do_not_collide(tmp_path):
    """以圖片填滿的形狀與另一張 PNG 同表：兩者媒體皆須與模板相同，不互相覆蓋或遺失。

    回歸測試：舊回填邏輯沿用模板檔名寫媒體，當 openpyxl 重新編號後的檔名與
    模板保留媒體的原始檔名相同時，會略過保留媒體（關聯指向錯圖）。
    """
    tpl = tmp_path / "tpl_fill_png.xlsx"
    _make_base_workbook(tpl)
    red = _tiny_png_bytes((200, 30, 30))
    blue = _tiny_png_bytes((30, 30, 200))
    _inject_drawing(
        tpl,
        _wsdr_xml(_filled_shape_anchor_xml("rId1", 2, "FilledShape")
                  + _pic_anchor_xml("rId2", 10, "BluePic")),
        _drawing_rels_xml({"rId1": "image1.png", "rId2": "image2.png"}),
        {"image1.png": red, "image2.png": blue},
    )

    out = tmp_path / "out_fill_png.xlsx"
    render(str(tpl), str(out), data={})

    info = _inspect_sheet(out, "Sheet")
    assert len(info["shapes"]) == 1
    assert len(info["pics"]) == 1
    assert info["shapes"][0]["media_hash"] == _sha256(red)
    assert info["pics"][0]["media_hash"] == _sha256(blue)



# ======================================================================
# 失敗與生命週期（design 決策 6／7）
# ======================================================================


class _BoomEtree:
    """讓 templexl.core.drawing_keeper 內部任何 ``etree.*`` 呼叫立即拋例外。

    只替換 drawing_keeper 模組自己的 ``etree`` 名稱綁定（``from lxml import
    etree``），不觸碰 lxml.etree 本身，故不會波及 golden_compare、openpyxl
    等其他也用到 lxml 的模組。這樣才是真正打破 capture/restore **內部**的
    解析步驟，而不是整個替換 capture_unsupported_drawings／restore_drawings
    ——後者只驗得到 Book 的外層包裝，鎖不住 design 決策 6「移除
    drawing_keeper 內部 try/except 吞例外」：任務 4.3 之前這幾個測試須為
    RED（內部吞例外，Book 不拋錯），4.3 移除吞例外後才轉 GREEN。
    """

    def __getattr__(self, name):
        raise RuntimeError("boom: drawing_keeper internal etree call failed")


def _break_drawing_keeper_parsing(monkeypatch) -> None:
    import templexl.core.drawing_keeper as drawing_keeper

    monkeypatch.setattr(drawing_keeper, "etree", _BoomEtree())


def test_save_restore_failure_raises_render_error_no_leftover_tmp(
    template_with_shape, tmp_path, monkeypatch
):
    """回填拋例外時 Book.save() 須拋 RenderError，輸出路徑不存在、無殘留暫存檔。"""
    tpl = template_with_shape(8, 16)
    out = tmp_path / "out_fail.xlsx"

    with Book(str(tpl)) as bk:
        bk.render({"title": "X", "df": _df(5)})
        _break_drawing_keeper_parsing(monkeypatch)  # 只打破存檔時的回填，不影響已完成的擷取
        with pytest.raises(RenderError):
            bk.save(str(out))

    assert not out.exists()
    assert list(tmp_path.glob("*.xlsx.tmp")) == []


def test_save_restore_failure_keeps_existing_output_unchanged(
    template_with_shape, tmp_path, monkeypatch
):
    """輸出路徑已有舊檔時，回填失敗不得覆寫，舊檔位元組維持不變。"""
    tpl = template_with_shape(8, 16)
    out = tmp_path / "out_existing.xlsx"
    original_bytes = b"not a real xlsx, just a marker"
    out.write_bytes(original_bytes)

    with Book(str(tpl)) as bk:
        bk.render({"title": "X", "df": _df(5)})
        _break_drawing_keeper_parsing(monkeypatch)
        with pytest.raises(RenderError):
            bk.save(str(out))

    assert out.read_bytes() == original_bytes


def test_capture_failure_raises_render_error(template_with_shape, monkeypatch):
    """擷取拋例外時 Book(template) 須拋 RenderError（不得吞例外、靜默略過保留）。"""
    tpl = template_with_shape(8, 16)

    _break_drawing_keeper_parsing(monkeypatch)

    with pytest.raises(RenderError):
        Book(str(tpl))


def test_shape_survives_sheet_rename_before_save(template_with_shape, tmp_path):
    """存檔前以逃生門改名工作表，形狀仍保留於改名後的工作表（design 決策 7）。"""
    tpl = template_with_shape(8, 16)
    out = tmp_path / "out_renamed.xlsx"
    with Book(str(tpl)) as bk:
        bk.render({"title": "X", "df": _df(5)})
        bk.sheets[0].ws.title = "Renamed"
        bk.save(str(out))

    assert _anchor_rows(out) == [(12, 20)]
    from openpyxl import load_workbook

    assert load_workbook(out).sheetnames == ["Renamed"]


def test_shape_skipped_when_sheet_removed_before_save(tmp_path):
    """存檔前以逃生門移除含形狀的工作表，存檔成功且不拋例外（design 決策 7）。"""
    wb = OpenpyxlWorkbook()
    wb.active.title = "WithShape"
    wb.active["A1"] = "no tag"
    wb.create_sheet("Other")
    tpl = tmp_path / "tpl_two_sheets.xlsx"
    wb.save(tpl)
    _inject_shape(tpl, 2, 6)

    out = tmp_path / "out_removed.xlsx"
    with Book(str(tpl)) as bk:
        bk.render({})
        bk.workbook.remove(bk.workbook["WithShape"])
        bk.save(str(out))

    from openpyxl import load_workbook

    sheetnames = load_workbook(out).sheetnames
    assert "WithShape" not in sheetnames
    assert "Other" in sheetnames


def test_shape_not_restored_into_unrelated_sheet_reusing_removed_title(tmp_path):
    """移除含形狀的工作表後，另一張工作表改名為同一個名稱，形狀不得誤回填。

    design 決策 7：「物件已不在 workbook.worksheets 中（被移除）則略過」，
    這須是明確的物件身分檢查——只靠「移除後 title 在輸出檔對不到部件」
    自然略過並不足夠：若移除後另一張工作表改名為同一個 title，靠字串比對
    的 restore_drawings 會把形狀誤回填到這張無關的工作表，且不拋任何錯誤。
    """
    wb = OpenpyxlWorkbook()
    wb.active.title = "WithShape"
    wb.active["A1"] = "no tag"
    wb.create_sheet("Other")
    tpl = tmp_path / "tpl_reuse_title.xlsx"
    wb.save(tpl)
    _inject_shape(tpl, 2, 6)

    out = tmp_path / "out_reuse_title.xlsx"
    with Book(str(tpl)) as bk:
        bk.render({})
        bk.workbook.remove(bk.workbook["WithShape"])
        bk.workbook["Other"].title = "WithShape"  # 重用被移除工作表的舊名稱
        bk.save(str(out))

    assert _anchor_rows(out) == [], "形狀不應誤回填到重用同名的無關工作表"


def test_book_and_sheet_render_produce_same_drawing_signatures(tmp_path):
    """Book.render() 與逐一 Sheet.render() 兩條路徑，輸出的繪圖／頁首頁尾簽章須相同。"""
    from tests.golden_compare import workbook_signature

    out_book = tmp_path / "out_book_render.xlsx"
    with Book(str(TEMPLATE_TABLE)) as bk:
        bk.render(table_data())
        bk.save(str(out_book))

    out_sheet = tmp_path / "out_sheet_render.xlsx"
    with Book(str(TEMPLATE_TABLE)) as bk:
        data = table_data()
        for sht in bk.sheets:
            sht.render(data)
        bk.save(str(out_sheet))

    sig_book = workbook_signature(out_book)
    sig_sheet = workbook_signature(out_sheet)
    assert sig_book["sheets"].keys() == sig_sheet["sheets"].keys()
    for name in sig_book["sheets"]:
        assert (sig_book["sheets"][name]["drawings"]
                == sig_sheet["sheets"][name]["drawings"]), name
        assert (sig_book["sheets"][name]["header_footer_images"]
                == sig_sheet["sheets"][name]["header_footer_images"]), name
