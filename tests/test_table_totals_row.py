"""Excel 表格幾何契約測試（對應 specs/excel-table-geometry）。

表格標籤綁定 Excel Table 物件時，渲染後與存檔時的幾何四不變量：

- ``ref``      ＝ 表頭列＋資料列＋合計列
- ``autoFilter.ref`` ＝ 表頭列＋資料列（**不含**合計列，Excel 慣例）
- ``totalsRowCount`` / ``totalsRowShown`` 保留模板原值
- body ≥ 1 列（Excel 表格至少需一列資料列）

模板一律以 openpyxl 程式化建構，才能精準控制「含合計列」的各種形狀
（noheader／header／空 DataFrame／疊放／標籤落在合計列／未渲染表格）。

已知陷阱：表格 ``displayName`` **不可長得像儲存格參照**（例如 ``T1``、``C3``），
否則 Excel 會判定檔案損毀而拒絕開啟。本檔一律以 ``Tbl*`` 命名；又因
``table_writer`` 的表格比對優先以「標籤名稱是否為表格名稱的子字串」進行，
表格名刻意寫成 ``Tbl_<標籤名>``，使綁定結果與表格數量無關而穩定。
"""

from __future__ import annotations

import logging

import pandas as pd
import pytest
from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn

from templexl import Book, render

from tests.golden_compare import render_quietly

COLUMNS = ("TABLE_NAME", "TABLE_ROWS", "TABLE_SIZE_MB")
TOTALS_LABEL = "合計"


# ======================================================================
# 模板建構 helper
# ======================================================================


def _add_table(
    worksheet,
    *,
    name: str,
    header_row: int,
    tag: str | None,
    columns: tuple = COLUMNS,
    with_totals: bool = True,
    totals_tag: str | None = None,
):
    """在 ``worksheet`` 建立「表頭列＋模板列（＋合計列）」的 Excel Table。

    Args:
        name: 表格 displayName（務必 ``Tbl`` 開頭，見模組 docstring）。
        header_row: 表頭列列號；模板列為其下一列，合計列再下一列。
        tag: 寫在模板列 A 欄的標籤字串；``None`` 表示不放標籤（未渲染表格）。
        with_totals: 是否建立合計列（``totalsRowCount=1``）。
        totals_tag: 寫在合計列 A 欄的標籤字串（測「標籤落在合計列」用）。

    Returns:
        openpyxl ``Table``：已加入工作表的表格物件。
    """
    tpl_row = header_row + 1
    last_col_letter = get_column_letter(len(columns))

    for idx, column_name in enumerate(columns, start=1):
        worksheet.cell(row=header_row, column=idx, value=column_name)
    if tag is not None:
        worksheet.cell(row=tpl_row, column=1, value=tag)

    table_columns = [
        TableColumn(id=idx, name=column_name)
        for idx, column_name in enumerate(columns, start=1)
    ]

    last_row = tpl_row
    totals_row_count = None
    if with_totals:
        totals_row = tpl_row + 1
        worksheet.cell(
            row=totals_row, column=1, value=totals_tag or TOTALS_LABEL
        )
        worksheet.cell(
            row=totals_row,
            column=len(columns),
            value=f"=SUBTOTAL(109,{name}[{columns[-1]}])",
        )
        table_columns[0].totalsRowLabel = TOTALS_LABEL
        table_columns[-1].totalsRowFunction = "sum"
        last_row = totals_row
        totals_row_count = 1

    table = Table(
        displayName=name,
        name=name,
        ref=f"A{header_row}:{last_col_letter}{last_row}",
        headerRowCount=1,
        totalsRowCount=totals_row_count,
        tableColumns=table_columns,
        autoFilter=AutoFilter(ref=f"A{header_row}:{last_col_letter}{tpl_row}"),
    )
    worksheet.add_table(table)
    return table


def _template(tmp_path, filename: str = "template.xlsx", **kwargs):
    """單表格模板：表頭列＋模板列（＋合計列），表格名為 ``Tbl_<標籤名>``。"""
    workbook = OpenpyxlWorkbook()
    worksheet = workbook.active
    _add_table(worksheet, header_row=1, **kwargs)
    path = tmp_path / filename
    workbook.save(path)
    return path


def _data(rows: int, columns: tuple = COLUMNS) -> pd.DataFrame:
    """``rows`` 筆確定性資料；``TABLE_SIZE_MB`` 之和為 ``rows * 10``。"""
    return pd.DataFrame(
        {
            columns[0]: [f"TB_{i:02d}" for i in range(1, rows + 1)],
            columns[1]: [i * 100 for i in range(1, rows + 1)],
            columns[2]: [10] * rows,
        }
    )


def _rendered(template_path, tmp_path, data: dict, out_name: str = "out.xlsx"):
    """渲染並回傳 ``(worksheet, table)``（單表格模板用）。"""
    out = tmp_path / out_name
    render_quietly(lambda: render(str(template_path), str(out), data=data))
    worksheet = load_workbook(out).active
    return worksheet, next(iter(worksheet.tables.values()))


# ======================================================================
# Requirement: 合計列隨資料列展開而保留
# ======================================================================


def test_noheader_totals_row_survives_expansion(tmp_path):
    """noheader 展開 N 筆：合計列緊接最後一筆資料、公式原文不變、範圍涵蓋之。"""
    template = _template(
        tmp_path, name="Tbl_rep_df", tag="#{{rep_df | noheader}}"
    )
    worksheet, table = _rendered(template, tmp_path, {"rep_df": _data(3)})

    # 資料佔 2..4 列，合計列被推到第 5 列（模板第 3 列 + 插入的 2 列）
    assert worksheet["A2"].value == "TB_01"
    assert worksheet["A4"].value == "TB_03"
    assert worksheet["A5"].value == TOTALS_LABEL
    assert worksheet["C5"].value == "=SUBTOTAL(109,Tbl_rep_df[TABLE_SIZE_MB])"

    assert table.ref == "A1:C5"
    assert table.autoFilter.ref == "A1:C4"
    assert table.totalsRowCount == 1


def test_noheader_totals_flags_preserved(tmp_path):
    """合計列旗標保留模板原值，不被同步流程清成 None/False。"""
    template = _template(
        tmp_path, name="Tbl_rep_df", tag="#{{rep_df | noheader}}"
    )
    original = load_workbook(template).active.tables["Tbl_rep_df"]
    _, table = _rendered(template, tmp_path, {"rep_df": _data(3)})

    assert table.totalsRowCount == original.totalsRowCount == 1
    assert table.totalsRowShown == original.totalsRowShown
    assert table.headerRowCount == 1


def test_noheader_totals_row_column_definitions_preserved(tmp_path):
    """noheader 模式完全保留模板欄定義（含合計列設定）。"""
    template = _template(
        tmp_path, name="Tbl_rep_df", tag="#{{rep_df | noheader}}"
    )
    _, table = _rendered(template, tmp_path, {"rep_df": _data(3)})

    assert [col.name for col in table.tableColumns] == list(COLUMNS)
    assert table.tableColumns[0].totalsRowLabel == TOTALS_LABEL
    assert table.tableColumns[-1].totalsRowFunction == "sum"


def test_header_mode_reuses_totals_settings_by_column_name(tmp_path):
    """header 模式重建欄定義時，以欄名比對沿用模板欄的合計列設定。"""
    template = _template(tmp_path, name="Tbl_rep_df", tag="#{{rep_df}}")
    # 中間欄改名（TABLE_ROWS -> ROW_COUNT），首尾欄名與模板相同
    columns = ("TABLE_NAME", "ROW_COUNT", "TABLE_SIZE_MB")
    worksheet, table = _rendered(
        template, tmp_path, {"rep_df": _data(3, columns)}
    )

    assert worksheet["A1"].value == "TABLE_NAME"
    assert worksheet["B1"].value == "ROW_COUNT"
    assert worksheet["A5"].value == TOTALS_LABEL
    assert worksheet["C5"].value == "=SUBTOTAL(109,Tbl_rep_df[TABLE_SIZE_MB])"
    assert table.ref == "A1:C5"
    assert table.autoFilter.ref == "A1:C4"
    assert table.totalsRowCount == 1

    assert [col.id for col in table.tableColumns] == [1, 2, 3]
    assert [col.name for col in table.tableColumns] == list(columns)
    # 同名欄沿用；改名欄無合計列設定
    assert table.tableColumns[0].totalsRowLabel == TOTALS_LABEL
    assert table.tableColumns[1].totalsRowLabel is None
    assert table.tableColumns[1].totalsRowFunction is None
    assert table.tableColumns[2].totalsRowFunction == "sum"


def test_content_below_totals_row_shifts_by_data_rows_minus_one(tmp_path):
    """合計列下方的內容依實際展開列數位移 N−1 列，不被覆蓋。"""
    workbook = OpenpyxlWorkbook()
    worksheet = workbook.active
    _add_table(
        worksheet, name="Tbl_rep_df", header_row=1, tag="#{{rep_df | noheader}}"
    )
    worksheet["A5"] = "footer"
    worksheet["B5"] = "=B4*2"
    template = tmp_path / "with_footer.xlsx"
    workbook.save(template)

    worksheet, table = _rendered(template, tmp_path, {"rep_df": _data(3)})

    assert table.ref == "A1:C5"
    assert worksheet["A5"].value == TOTALS_LABEL  # 合計列佔住第 5 列
    assert worksheet["A7"].value == "footer"  # 5 + (3-1)
    assert worksheet["B7"].value == "=B4*2"  # 插入點下方公式不平移（現行語意）


def test_stacked_tables_upper_has_totals_row(tmp_path):
    """含合計列的表格上方疊放另一表格：兩表範圍不重疊、間距維持。"""
    workbook = OpenpyxlWorkbook()
    worksheet = workbook.active
    # 上表：1-3 列（含合計列）；空 1 列；下表：5-6 列（無合計列）
    _add_table(
        worksheet, name="Tbl_top_df", header_row=1, tag="#{{top_df | noheader}}"
    )
    _add_table(
        worksheet,
        name="Tbl_btm_df",
        header_row=5,
        tag="#{{btm_df | noheader}}",
        with_totals=False,
    )
    template = tmp_path / "stacked.xlsx"
    workbook.save(template)

    out = tmp_path / "stacked_out.xlsx"
    render_quietly(
        lambda: render(
            str(template),
            str(out),
            data={"top_df": _data(3), "btm_df": _data(2)},
        )
    )
    worksheet = load_workbook(out).active
    top = worksheet.tables["Tbl_top_df"]
    btm = worksheet.tables["Tbl_btm_df"]

    # 上表：表頭 1、資料 2-4、合計 5
    assert top.ref == "A1:C5"
    assert top.autoFilter.ref == "A1:C4"
    assert worksheet["A5"].value == TOTALS_LABEL
    # 下表：模板 5 -> 位移 2 列後表頭 7、資料 8-9（模板間距 1 列維持）
    assert btm.ref == "A7:C9"
    assert btm.autoFilter.ref == "A7:C9"
    assert worksheet["A7"].value == "TABLE_NAME"
    assert worksheet["A9"].value == "TB_02"


# ======================================================================
# Requirement: 表格至少保留一列資料列
# ======================================================================


def test_empty_dataframe_with_totals_row_keeps_one_body_row(tmp_path):
    """空 DataFrame 且含合計列：表頭＋一列空白資料列＋合計列共三列。"""
    template = _template(
        tmp_path, name="Tbl_rep_df", tag="#{{rep_df | noheader}}"
    )
    worksheet, table = _rendered(
        template, tmp_path, {"rep_df": _data(0)}
    )

    assert table.ref == "A1:C3"
    assert table.autoFilter.ref == "A1:C2"
    assert table.totalsRowCount == 1
    assert worksheet["A2"].value is None  # 空白資料列
    assert worksheet["A3"].value == TOTALS_LABEL
    assert worksheet["C3"].value == "=SUBTOTAL(109,Tbl_rep_df[TABLE_SIZE_MB])"


def test_empty_dataframe_without_totals_row_keeps_one_body_row(tmp_path):
    """空 DataFrame 且無合計列：表頭＋一列空白資料列共兩列，篩選範圍相同。"""
    template = _template(
        tmp_path,
        name="Tbl_rep_df",
        tag="#{{rep_df | noheader}}",
        with_totals=False,
    )
    worksheet, table = _rendered(template, tmp_path, {"rep_df": _data(0)})

    assert table.ref == "A1:C2"
    assert table.autoFilter.ref == "A1:C2"
    assert worksheet["A2"].value is None


# ======================================================================
# Requirement: 篩選範圍不含合計列
# ======================================================================


def test_unrendered_totals_table_survives_book_save(tmp_path):
    """未渲染的含合計列表格：存檔後範圍與篩選範圍皆維持模板原值。"""
    template = _template(tmp_path, name="Tbl_idle", tag=None)
    out = tmp_path / "idle_out.xlsx"

    with Book(template) as book:
        render_quietly(lambda: book.render({}))
        book.save(out)

    table = load_workbook(out).active.tables["Tbl_idle"]
    assert table.ref == "A1:C3"
    assert table.autoFilter.ref == "A1:C2"
    assert table.totalsRowCount == 1


def test_table_without_totals_autofilter_equals_ref(tmp_path):
    """無合計列的表格：篩選範圍等於表格範圍（既有行為不變）。"""
    template = _template(
        tmp_path,
        name="Tbl_rep_df",
        tag="#{{rep_df | noheader}}",
        with_totals=False,
    )
    _, table = _rendered(template, tmp_path, {"rep_df": _data(4)})

    assert table.ref == "A1:C5"
    assert table.autoFilter.ref == table.ref


# ======================================================================
# Requirement: 標籤位於合計列時不綁定表格物件
# ======================================================================


def test_tag_on_totals_row_is_not_bound_and_warns(tmp_path, caplog):
    """標籤落在合計列：發出含表格名稱與座標的 warning，表格範圍不變。

    此案例的表格名刻意**不含**標籤名（``TblArgus`` 對 ``rep_df``，即 Excel
    自動命名如 ``表格1`` 的實務形狀），走的是位置備援比對路徑。
    ``table_writer`` 另有「標籤名為表格名子字串」的名稱比對策略，那條路徑
    與標籤所在位置無關，不在本能力的範圍內。
    """
    template = _template(
        tmp_path,
        name="TblArgus",
        tag=None,
        totals_tag="#{{rep_df | noheader}}",
    )
    out = tmp_path / "totals_tag_out.xlsx"

    with caplog.at_level(logging.WARNING, logger="templexl"):
        render_quietly(
            lambda: render(str(template), str(out), data={"rep_df": _data(3)})
        )

    messages = [rec.getMessage() for rec in caplog.records]
    assert any(
        "TblArgus" in msg and "A3" in msg for msg in messages
    ), f"未見合計列標籤 warning，實際記錄: {messages}"

    table = load_workbook(out).active.tables["TblArgus"]
    assert table.ref == "A1:C3"
    assert table.autoFilter.ref == "A1:C2"


# ======================================================================
# helper: expected_autofilter_ref
# ======================================================================


@pytest.mark.parametrize(
    "ref,totals_row_count,expected",
    [
        ("A1:C3", 1, "A1:C2"),
        ("A1:C3", None, "A1:C3"),
        ("A1:C3", 0, "A1:C3"),
        ("A1:C5", 2, "A1:C3"),
        ("B2:D10", 1, "B2:D9"),
    ],
)
def test_expected_autofilter_ref(ref, totals_row_count, expected):
    """篩選範圍規則：``ref`` 扣除尾端 ``totalsRowCount`` 列。"""
    from templexl.core.table_sync import expected_autofilter_ref

    table = Table(displayName="TblX", ref=ref, totalsRowCount=totals_row_count)
    assert expected_autofilter_ref(table) == expected


def test_expected_autofilter_ref_never_collapses_below_header(tmp_path):
    """合計列數大於等於 ``ref`` 高度時仍回傳合法範圍（不產生倒置範圍）。"""
    from templexl.core.table_sync import expected_autofilter_ref

    table = Table(displayName="TblX", ref="A1:C2", totalsRowCount=2)
    assert expected_autofilter_ref(table) == "A1:C1"
