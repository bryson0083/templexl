"""
黃金檔案比對工具。

將 .xlsx 序列化為穩定的「簽章」結構（值、數字格式、關鍵樣式、合併範圍、
Excel Table 範圍、圖片錨點），並比對兩份簽章、產出可讀的差異清單。 

這是重構安全網的核心：相同模板 + 相同資料 + 相同程式 → 相同簽章。
任何差異都代表渲染輸出改變，須逐項判定為「改對」或「改壞」。
"""

from __future__ import annotations

import hashlib
import os
import posixpath
import zipfile
from contextlib import redirect_stderr, redirect_stdout

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

NS_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_V = "urn:schemas-microsoft-com:vml"
NS_O = "urn:schemas-microsoft-com:office:office"

_ANCHOR_CHILD_KINDS = ("sp", "pic", "graphicFrame", "cxnSp", "grpSp")


def _color_repr(color) -> str | None:
    """安全擷取顏色的穩定表示。"""
    if color is None:
        return None
    rgb = getattr(color, "rgb", None)
    if isinstance(rgb, str):
        return rgb
    theme = getattr(color, "theme", None)
    if theme is not None:
        return f"theme:{theme}:{getattr(color, 'tint', 0)}"
    return None


def _font_sig(font) -> tuple:
    if font is None:
        return ()
    return (
        font.name,
        font.size,
        bool(font.bold),
        bool(font.italic),
        _color_repr(font.color),
    )


def _fill_sig(fill) -> tuple:
    if fill is None:
        return ()
    pattern = getattr(fill, "patternType", None)
    if not pattern:
        return ()
    return (pattern, _color_repr(getattr(fill, "fgColor", None)))


def _align_sig(alignment) -> tuple:
    if alignment is None:
        return ()
    return (
        alignment.horizontal,
        alignment.vertical,
        bool(alignment.wrap_text),
    )


def _border_sig(border) -> tuple:
    if border is None:
        return ()
    return tuple(
        getattr(getattr(border, side), "style", None)
        for side in ("left", "right", "top", "bottom")
    )


def _cell_sig(cell) -> dict:
    sig = {
        "value": cell.value,
        "data_type": cell.data_type,
        "number_format": cell.number_format,
    }
    if cell.has_style:
        sig["font"] = _font_sig(cell.font)
        sig["fill"] = _fill_sig(cell.fill)
        sig["alignment"] = _align_sig(cell.alignment)
        sig["border"] = _border_sig(cell.border)
    return sig


def _image_sigs(ws) -> list:
    sigs = []
    for img in getattr(ws, "_images", []):
        anchor = getattr(img, "anchor", None)
        entry = {"anchor_type": type(anchor).__name__ if anchor else None}
        frm = getattr(anchor, "_from", None)
        if frm is not None:
            entry["from"] = (getattr(frm, "row", None), getattr(frm, "col", None))
        to = getattr(anchor, "to", None)
        if to is not None:
            entry["to"] = (getattr(to, "row", None), getattr(to, "col", None))
        sigs.append(entry)
    # 以錨點位置排序，避免順序造成假性差異
    sigs.sort(key=lambda e: (str(e.get("from")), str(e.get("to"))))
    return sigs


def _sheet_sig(ws, drawings: list, header_footer_images: list) -> dict:
    cells = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None or cell.has_style:
                cells[cell.coordinate] = _cell_sig(cell)

    tables = {}
    for name in ws.tables:
        table = ws.tables[name]
        af = getattr(table, "autoFilter", None)
        tables[name] = {
            "ref": getattr(table, "ref", None),
            "autofilter_ref": getattr(af, "ref", None) if af else None,
            # 合計列／表頭列旗標：鎖住「Excel 是否認得這張表的合計列」，
            # 光看 ref 與 autofilter_ref 是看不出來的。
            "totals_row_count": getattr(table, "totalsRowCount", None),
            "totals_row_shown": getattr(table, "totalsRowShown", None),
            "header_row_count": getattr(table, "headerRowCount", None),
            "columns": _table_column_sigs(table),
        }

    return {
        "dimensions": ws.dimensions,
        "cells": cells,
        "merged_cells": sorted(str(r) for r in ws.merged_cells.ranges),
        "tables": tables,
        "images": _image_sigs(ws),
        "drawings": drawings,
        "header_footer_images": header_footer_images,
    }


# ----------------------------------------------------------------------
# zip 層繪圖簽章（openpyxl 讀不到的內容：形狀、頁首頁尾圖片）
# ----------------------------------------------------------------------


def _zip_resolve(base_part: str, target: str) -> str:
    """把 rels 的相對 Target 正規化為 zip 內的完整路徑。"""
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_part), target))


def _zip_rels_path(part: str) -> str:
    return posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")


def _zip_rels_map(zf, names, part: str) -> dict:
    """回傳 ``{rId: <Relationship 元素>}``；part 沒有 rels 時回傳空字典。"""
    rels_part = _zip_rels_path(part)
    if rels_part not in names:
        return {}
    return {rel.get("Id"): rel for rel in etree.fromstring(zf.read(rels_part))}


def _zip_sheet_part_map(zf, names) -> dict:
    """回傳 ``{工作表名稱: zip 內的 sheet part 路徑}``。"""
    if "xl/workbook.xml" not in names:
        return {}
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    wb_rels = _zip_rels_map(zf, names, "xl/workbook.xml")
    mapping = {}
    for sheet in wb.iter(f"{{{NS_MAIN}}}sheet"):
        rid = sheet.get(f"{{{NS_REL}}}id")
        rel = wb_rels.get(rid)
        if rel is not None:
            mapping[sheet.get("name")] = _zip_resolve("xl/workbook.xml", rel.get("Target"))
    return mapping


def _media_hash(zf, names, media_path: str) -> str | None:
    if media_path not in names:
        return None
    return hashlib.sha256(zf.read(media_path)).hexdigest()


def _anchor_child_kind(anchor) -> str | None:
    """錨點的實際內容種類（sp／pic／graphicFrame／cxnSp／grpSp）。"""
    for child in anchor:
        local = etree.QName(child.tag).localname
        if local in _ANCHOR_CHILD_KINDS:
            return local
    return None


def _anchor_from_to(anchor):
    def _pos(tag):
        el = anchor.find(f"{{{NS_XDR}}}{tag}")
        if el is None:
            return None
        col = el.find(f"{{{NS_XDR}}}col")
        row = el.find(f"{{{NS_XDR}}}row")
        return (
            int(row.text) if row is not None else None,
            int(col.text) if col is not None else None,
        )

    return _pos("from"), _pos("to")


def _anchor_rotation(anchor) -> int:
    for xfrm in anchor.iter(f"{{{NS_A}}}xfrm"):
        rot = xfrm.get("rot")
        if rot is not None:
            return int(rot)
    return 0


def _anchor_text(anchor) -> str:
    return "".join(t.text or "" for t in anchor.iter(f"{{{NS_A}}}t"))


def _anchor_media_hash(zf, names, drawing_part: str, anchor, rels: dict) -> str | None:
    """錨點（pic 或以圖片填滿的形狀）引用的媒體內容雜湊；沒有則 ``None``。"""
    for blip in anchor.iter(f"{{{NS_A}}}blip"):
        rid = blip.get(f"{{{NS_REL}}}embed") or blip.get(f"{{{NS_REL}}}link")
        rel = rels.get(rid) if rid else None
        if rel is not None:
            target = _zip_resolve(drawing_part, rel.get("Target"))
            return _media_hash(zf, names, target)
    return None


def _drawing_sigs(zf, names, drawing_part: str | None) -> list:
    """該工作表繪圖部件中所有錨點的排序清單（不含部件名稱與 rId）。"""
    if not drawing_part or drawing_part not in names:
        return []
    root = etree.fromstring(zf.read(drawing_part))
    rels = _zip_rels_map(zf, names, drawing_part)
    sigs = []
    for anchor in root:
        kind = _anchor_child_kind(anchor)
        if kind is None:
            continue
        frm, to = _anchor_from_to(anchor)
        sigs.append({
            "anchor_type": etree.QName(anchor.tag).localname,
            "child_kind": kind,
            "from": frm,
            "to": to,
            "rotation": _anchor_rotation(anchor),
            "text": _anchor_text(anchor),
            "media_hash": _anchor_media_hash(zf, names, drawing_part, anchor, rels),
        })
    sigs.sort(key=lambda e: (str(e["from"]), str(e["to"]), e["child_kind"], e["text"]))
    return sigs


def _header_footer_image_sigs(zf, names, sheet_part: str | None) -> list:
    """該工作表頁首頁尾 VML 中每個圖片 shape 的位置代碼與媒體雜湊排序清單。

    以工作表 XML 的 ``<legacyDrawingHF r:id>`` 精確定位頁首頁尾 VML 的
    關聯 Id，不能只取「第一個 Type 結尾為 /vmlDrawing 的關聯」——openpyxl
    寫出的儲存格註解 VML 使用同一個關聯類型，同一張工作表若同時有註解與
    頁首頁尾圖片，「取第一個」會誤抓註解 VML。
    """
    if not sheet_part or sheet_part not in names:
        return []
    sheet_root = etree.fromstring(zf.read(sheet_part))
    legacy = sheet_root.find(f"{{{NS_MAIN}}}legacyDrawingHF")
    if legacy is None:
        return []
    rid = legacy.get(f"{{{NS_REL}}}id")
    rel = _zip_rels_map(zf, names, sheet_part).get(rid)
    if rel is None:
        return []
    vml_part = _zip_resolve(sheet_part, rel.get("Target"))
    if vml_part not in names:
        return []
    vml_root = etree.fromstring(zf.read(vml_part))
    vml_rels = _zip_rels_map(zf, names, vml_part)
    sigs = []
    for shape in vml_root.iter(f"{{{NS_V}}}shape"):
        imagedata = shape.find(f"{{{NS_V}}}imagedata")
        if imagedata is None:
            continue
        rel = vml_rels.get(imagedata.get(f"{{{NS_O}}}relid"))
        media_hash = None
        if rel is not None:
            target = _zip_resolve(vml_part, rel.get("Target"))
            media_hash = _media_hash(zf, names, target)
        sigs.append({"position": shape.get("id"), "media_hash": media_hash})
    sigs.sort(key=lambda e: str(e["position"]))
    return sigs


def _table_column_sigs(table) -> list:
    """表格欄定義簽章：欄名與合計列設定（結構化參照以欄名為鍵）。"""
    return [
        {
            "name": col.name,
            "totals_row_function": getattr(col, "totalsRowFunction", None),
            "totals_row_label": getattr(col, "totalsRowLabel", None),
            "totals_row_formula": getattr(col, "totalsRowFormula", None),
        }
        for col in (table.tableColumns or [])
    ]


def _defined_name_sigs(wb) -> dict:
    """具名範圍簽章：全簿層以名稱為鍵，工作表層以 ``<工作表>!<名稱>`` 為鍵。

    引用 ``表格[[#Totals],[欄]]`` 的具名範圍是合計列能否被下游公式取用的
    關鍵證據，故納入簽章。
    """
    sigs = {name: dn.value for name, dn in wb.defined_names.items()}
    for ws in wb.worksheets:
        for name, dn in getattr(ws, "defined_names", {}).items():
            sigs[f"{ws.title}!{name}"] = dn.value
    return sigs


def workbook_signature(path) -> dict:
    """載入 .xlsx 並回傳穩定簽章結構。

    結構為 ``{"sheets": {工作表名: 工作表簽章}, "defined_names": {...}}``——
    具名範圍屬全簿層級，故簽章多包一層而非與工作表平放。

    工作表簽章另包含 zip 層才看得到的 ``drawings``（形狀等 openpyxl 無法
    表示的錨點）與 ``header_footer_images``（頁首頁尾 VML 圖片），使黃金
    比對能守護 openpyxl 模型之外的繪圖內容。
    """
    path = str(path)
    wb = load_workbook(path)
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        sheet_parts = _zip_sheet_part_map(zf, names)
        sheets = {}
        for title in wb.sheetnames:
            sheet_part = sheet_parts.get(title)
            drawing_part = None
            if sheet_part:
                for rel in _zip_rels_map(zf, names, sheet_part).values():
                    if rel.get("Type", "").endswith("/drawing"):
                        drawing_part = _zip_resolve(sheet_part, rel.get("Target"))
                        break
            sheets[title] = _sheet_sig(
                wb[title],
                _drawing_sigs(zf, names, drawing_part),
                _header_footer_image_sigs(zf, names, sheet_part),
            )
    return {
        "sheets": sheets,
        "defined_names": _defined_name_sigs(wb),
    }


def _diff(a, b, path, out, limit):
    if len(out) >= limit:
        return
    if isinstance(a, dict) and isinstance(b, dict):
        for key in sorted(set(a) | set(b)):
            if key not in a:
                out.append(f"{path}.{key}: 僅出現於實際輸出 = {b[key]!r}")
            elif key not in b:
                out.append(f"{path}.{key}: 僅出現於黃金基準 = {a[key]!r}")
            else:
                _diff(a[key], b[key], f"{path}.{key}", out, limit)
            if len(out) >= limit:
                return
    elif a != b:
        out.append(f"{path}: 黃金={a!r} vs 實際={b!r}")


def compare_signatures(golden: dict, actual: dict, limit: int = 50) -> list:
    """比對黃金 vs 實際簽章，回傳人類可讀的差異字串清單（最多 limit 筆）。"""
    out: list = []
    _diff(golden, actual, "", out, limit)
    return out


def render_quietly(render_callable):
    """執行渲染並抑制現況程式的大量 print 除錯輸出（stdout/stderr）。"""
    with open(os.devnull, "w") as devnull:
        with redirect_stdout(devnull), redirect_stderr(devnull):
            return render_callable()
