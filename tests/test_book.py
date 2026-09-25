"""``Book`` 會話 API 契約測試（對應 specs/book-session-api）。 

驗證生命週期（載入檢查、顯式存檔、context manager 語意）、sheet 存取、
逃生門，以及最關鍵的 parity：``Book`` 全簿渲染 ≡ ``render()``。 
"""

from __future__ import annotations

import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl import load_workbook

from templexl import (
    FileFormatError,
    RenderError,
    TemplateNotFoundError,
    TemplateResourceError,
)
from templexl.book import Book

from tests.fixtures import SCENARIOS
from tests.golden_compare import (
    compare_signatures,
    render_quietly,
    workbook_signature,
)
from tests.render_adapter import do_render


@pytest.fixture
def simple_template(tmp_path):
    """含純量標籤與表格標籤的最小模板。"""
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws["A1"] = "Hello {{name}}"
    ws["A3"] = "#{{df}}"
    path = tmp_path / "template.xlsx"
    wb.save(path)
    return path


def _simple_data() -> dict:
    return {"name": "王小明", "df": pd.DataFrame({"欄": ["a", "b"]})}


# ===================================================================
# parity：Book 全簿渲染 ≡ render()
# ===================================================================


@pytest.mark.parametrize(
    "name,template,data_fn",
    SCENARIOS,
    ids=[s[0] for s in SCENARIOS],
)
def test_book_render_matches_render(name, template, data_fn, tmp_path):
    """同一模板與資料：Book 開啟→全簿渲染→存檔，須與 render() 逐格等價。"""
    via_render = tmp_path / f"{name}_render.xlsx"
    via_book = tmp_path / f"{name}_book.xlsx"

    render_quietly(lambda: do_render(template, via_render, data_fn()))

    def _book_render():
        with Book(template) as bk:
            bk.render(data_fn())
            bk.save(via_book)

    render_quietly(_book_render)

    diffs = compare_signatures(
        workbook_signature(via_render),
        workbook_signature(via_book),
    )
    assert not diffs, "Book 輸出與 render() 不一致:\n" + "\n".join(diffs)


def test_book_render_returns_warnings(simple_template, tmp_path):
    """資料缺鍵時，未解析標籤以非致命警告回報。"""
    with Book(simple_template) as bk:
        result = render_quietly(lambda: bk.render({"df": pd.DataFrame({"欄": ["a"]})}))
    assert any("name" in w for w in result.warnings)


def test_book_render_with_report(simple_template):
    """with_report 啟用時回傳渲染報告。"""
    with Book(simple_template) as bk:
        result = render_quietly(lambda: bk.render(_simple_data(), with_report=True))
    assert result.report is not None
    assert result.report.summary["total_worksheets"] >= 1


def test_book_render_default_no_report(simple_template):
    with Book(simple_template) as bk:
        result = render_quietly(lambda: bk.render(_simple_data()))
    assert result.report is None


def test_escape_formulas_passthrough(simple_template, tmp_path):
    """escape_formulas 預設中和資料帶入的公式前綴字串。"""
    out = tmp_path / "out.xlsx"
    with Book(simple_template) as bk:
        render_quietly(lambda: bk.render({"name": "=2+5", "df": pd.DataFrame({"欄": ["x"]})}))
        bk.save(out)

    ws = load_workbook(out).active
    formulas = [c.coordinate for row in ws.iter_rows() for c in row if c.data_type == "f"]
    assert formulas == []


# ===================================================================
# 生命週期
# ===================================================================


def test_template_not_found_raises():
    with pytest.raises(TemplateNotFoundError):
        Book("definitely_missing.xlsx")


def test_unsupported_format_raises(tmp_path):
    bad = tmp_path / "x.txt"
    bad.write_text("x")
    with pytest.raises(FileFormatError):
        Book(bad)


def test_invalid_zip_rejected(tmp_path):
    fake = tmp_path / "fake.xlsx"
    fake.write_text("not a zip")
    with pytest.raises(FileFormatError):
        Book(fake)


def test_zip_bomb_rejected(tmp_path):
    bomb = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/huge.bin", b"\x00" * (64 * 1024 * 1024))
    with pytest.raises(TemplateResourceError):
        Book(bomb)


def test_no_save_produces_no_output(simple_template, tmp_path):
    """未呼叫 save()：磁碟上不產生任何輸出檔。"""
    out = tmp_path / "never_written.xlsx"
    with Book(simple_template) as bk:
        render_quietly(lambda: bk.render(_simple_data()))
    assert not out.exists()
    assert list(tmp_path.glob("*.xlsx")) == [simple_template]


def test_exit_does_not_autosave(simple_template, tmp_path):
    """with 區塊結束不自動存檔（落地必須顯式）。"""
    before = sorted(p.name for p in tmp_path.iterdir())
    with Book(simple_template) as bk:
        render_quietly(lambda: bk.render(_simple_data()))
    after = sorted(p.name for p in tmp_path.iterdir())
    assert before == after


def test_exit_propagates_exception(simple_template):
    """區塊內的例外原樣傳播，不被 __exit__ 攔截。"""
    class Boom(Exception):
        pass

    with pytest.raises(Boom):
        with Book(simple_template):
            raise Boom()


def test_use_after_close_raises(simple_template):
    bk = Book(simple_template)
    bk.close()
    with pytest.raises(RenderError):
        bk.workbook


def test_close_is_idempotent(simple_template):
    bk = Book(simple_template)
    bk.close()
    bk.close()


def test_save_returns_path(simple_template, tmp_path):
    out = tmp_path / "out.xlsx"
    with Book(simple_template) as bk:
        render_quietly(lambda: bk.render(_simple_data()))
        assert bk.save(out) == str(out)
    assert out.exists()


def test_failed_save_leaves_no_partial_output(simple_template, tmp_path, monkeypatch):
    """存檔中途失敗：目標維持舊內容、無暫存檔殘留。"""
    out = tmp_path / "out.xlsx"
    out.write_bytes(b"OLD_CONTENT")

    def boom(self, filename):
        raise OSError("disk full")

    monkeypatch.setattr(OpenpyxlWorkbook, "save", boom)
    with Book(simple_template) as bk:
        with pytest.raises(RenderError):
            bk.save(out)

    assert out.read_bytes() == b"OLD_CONTENT"
    assert not list(tmp_path.glob("*.tmp"))


# ===================================================================
# sheet 存取與逃生門
# ===================================================================


def test_sheets_dual_index(simple_template):
    """int 與名稱兩種鍵取得等價的 Sheet 物件。"""
    with Book(simple_template) as bk:
        by_index = bk.sheets[0]
        by_name = bk.sheets[by_index.title]
        assert by_index == by_name
        assert by_index.ws is by_name.ws


def test_sheets_missing_name_raises_keyerror(simple_template):
    with Book(simple_template) as bk:
        with pytest.raises(KeyError):
            bk.sheets["不存在的工作表"]


def test_sheets_index_out_of_range_raises(simple_template):
    with Book(simple_template) as bk:
        with pytest.raises(IndexError):
            bk.sheets[99]


def test_sheets_rejects_other_key_types(simple_template):
    with Book(simple_template) as bk:
        with pytest.raises(TypeError):
            bk.sheets[1.5]


def test_sheets_iteration_and_len(simple_template):
    with Book(simple_template) as bk:
        assert len(bk.sheets) == 1
        assert [s.title for s in bk.sheets] == bk.sheets.names
        assert bk.sheets.names[0] in bk.sheets


def test_escape_hatch_edits_reach_output(simple_template, tmp_path):
    """經 Sheet.ws 的 openpyxl 原生操作會反映在存檔結果。"""
    out = tmp_path / "out.xlsx"
    with Book(simple_template) as bk:
        render_quietly(lambda: bk.render(_simple_data()))
        bk.sheets[0].ws.merge_cells(start_row=10, start_column=1, end_row=10, end_column=3)
        bk.save(out)

    ws = load_workbook(out).active
    assert "A10:C10" in {str(r) for r in ws.merged_cells.ranges}


def test_workbook_escape_hatch_exposes_openpyxl(simple_template):
    with Book(simple_template) as bk:
        assert isinstance(bk.workbook, OpenpyxlWorkbook)


def test_independent_books_do_not_share_state(simple_template, tmp_path):
    """兩個 Book 實例各自渲染不同資料，輸出互不影響。"""
    out_a, out_b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"

    def _run(out, name):
        with Book(simple_template) as bk:
            bk.render({"name": name, "df": pd.DataFrame({"欄": [name]})})
            bk.save(out)

    render_quietly(lambda: _run(out_a, "AAA"))
    render_quietly(lambda: _run(out_b, "BBB"))

    values_a = {c.value for row in load_workbook(out_a).active.iter_rows() for c in row}
    values_b = {c.value for row in load_workbook(out_b).active.iter_rows() for c in row}
    assert "AAA" in values_a and "AAA" not in values_b
    assert "BBB" in values_b and "BBB" not in values_a
