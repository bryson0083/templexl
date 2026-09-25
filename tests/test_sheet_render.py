"""sheet 級渲染契約測試（對應 specs/sheet-level-render）。

驗證範圍限定（同名標籤跨工作表隔離）、sheet 範圍 warnings、重複渲染
no-op，以及「逐 sheet 渲染 ≡ 全簿渲染」的等價性。 
"""

from __future__ import annotations

import logging

import pandas as pd
import pytest
from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl import load_workbook

import templexl
from templexl import Book, SheetRenderResult

from tests.fixtures import SCENARIOS
from tests.golden_compare import (
    compare_signatures,
    render_quietly,
    workbook_signature,
)


@pytest.fixture
def two_sheet_template(tmp_path):
    """兩張工作表帶「同名」標籤——用以驗證 sheet 級渲染的範圍限定。"""
    wb = OpenpyxlWorkbook()
    first = wb.active
    first.title = "甲"
    first["A1"] = "{{who}}"
    first["A3"] = "#{{df}}"
    second = wb.create_sheet("乙")
    second["A1"] = "{{who}}"
    second["A3"] = "#{{df}}"
    path = tmp_path / "two_sheet.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def sample_template(tmp_path):
    """單一 [Sample] 工作表模板——對應「複製 N 份、每份不同資料」的實務形狀。"""
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws.title = "Sample"
    ws["A1"] = "{{vehicle_desc}}"
    ws["A3"] = "#{{detail_df | noheader}}"
    path = tmp_path / "sample.xlsx"
    wb.save(path)
    return path


def _values(worksheet) -> set:
    return {cell.value for row in worksheet.iter_rows() for cell in row}


# ===================================================================
# 範圍限定
# ===================================================================


def test_other_sheets_untouched(two_sheet_template):
    """只渲染其一時，另一張工作表的同名標籤原樣保留。"""
    with Book(two_sheet_template) as bk:
        render_quietly(
            lambda: bk.sheets["甲"].render({"who": "AAA", "df": pd.DataFrame({"欄": ["a"]})})
        )
        assert "AAA" in _values(bk.sheets["甲"].ws)
        assert "{{who}}" in _values(bk.sheets["乙"].ws)
        assert "#{{df}}" in _values(bk.sheets["乙"].ws)


def test_each_sheet_gets_its_own_data(two_sheet_template, tmp_path):
    """兩張工作表各以不同資料渲染，內容互不串接。"""
    out = tmp_path / "out.xlsx"
    with Book(two_sheet_template) as bk:
        render_quietly(
            lambda: bk.sheets["甲"].render({"who": "AAA", "df": pd.DataFrame({"欄": ["a1"]})})
        )
        render_quietly(
            lambda: bk.sheets["乙"].render({"who": "BBB", "df": pd.DataFrame({"欄": ["b1"]})})
        )
        bk.save(out)

    wb = load_workbook(out)
    first, second = _values(wb["甲"]), _values(wb["乙"])
    assert "AAA" in first and "a1" in first
    assert "BBB" not in first and "b1" not in first
    assert "BBB" in second and "b1" in second
    assert "AAA" not in second and "a1" not in second


def test_copied_sheets_render_independently(sample_template, tmp_path):
    """複製 Sample 工作表 N 份、逐份填不同資料（各運具彙總表的形狀）。"""
    out = tmp_path / "out.xlsx"
    vehicles = {
        "公車": pd.DataFrame({"日期": ["2026/07/01"], "金額": [100]}),
        "捷運": pd.DataFrame({"日期": ["2026/07/02"], "金額": [200]}),
        "渡輪": pd.DataFrame({"日期": ["2026/07/03"], "金額": [300]}),
    }

    def _build():
        with Book(sample_template) as bk:
            sample = bk.workbook["Sample"]
            for vehicle, detail in vehicles.items():
                copied = bk.workbook.copy_worksheet(sample)  # openpyxl 原生
                copied.title = vehicle
                bk.sheets[vehicle].render(
                    {"vehicle_desc": vehicle, "detail_df": detail}
                )
            del bk.workbook["Sample"]
            bk.save(out)

    render_quietly(_build)

    wb = load_workbook(out)
    assert wb.sheetnames == list(vehicles)
    for vehicle in vehicles:
        values = _values(wb[vehicle])
        assert vehicle in values
        assert vehicles[vehicle]["金額"].iloc[0] in values
        # 其他運具的金額不應外溢到本工作表
        for other, detail in vehicles.items():
            if other != vehicle:
                assert detail["金額"].iloc[0] not in values


# ===================================================================
# warnings 範圍
# ===================================================================


def test_warnings_scoped_to_sheet(two_sheet_template):
    """尚未渲染的其他工作表仍帶原始標籤，不得產生假警告。"""
    with Book(two_sheet_template) as bk:
        result = render_quietly(
            lambda: bk.sheets["甲"].render({"who": "AAA", "df": pd.DataFrame({"欄": ["a"]})})
        )
    assert isinstance(result, SheetRenderResult)
    assert result.sheet_name == "甲"
    assert result.warnings == []


def test_unresolved_tag_in_own_sheet_reported(two_sheet_template):
    """本工作表缺鍵的標籤要入列 warnings（含名稱與座標）。"""
    with Book(two_sheet_template) as bk:
        result = render_quietly(
            lambda: bk.sheets["甲"].render({"df": pd.DataFrame({"欄": ["a"]})})
        )
    assert any("who" in w for w in result.warnings)
    assert any("甲" in w for w in result.warnings)


# ===================================================================
# 重複渲染
# ===================================================================


def test_second_render_is_noop(two_sheet_template, caplog):
    """已渲染的工作表再次渲染：不改動儲存格、不拋例外、發出 warning 日誌。"""
    with Book(two_sheet_template) as bk:
        sheet = bk.sheets["甲"]
        render_quietly(
            lambda: sheet.render({"who": "AAA", "df": pd.DataFrame({"欄": ["a"]})})
        )
        before = workbook_signature_of_sheet(bk)

        with caplog.at_level(logging.WARNING, logger="templexl.book"):
            result = render_quietly(lambda: sheet.render({"who": "ZZZ"}))

        after = workbook_signature_of_sheet(bk)

    assert before == after
    assert result.warnings == []
    assert any("no-op" in rec.message or "no-op" in rec.getMessage() for rec in caplog.records)


def workbook_signature_of_sheet(book) -> list:
    """取本測試用的輕量快照（值 + 合併範圍），供 no-op 前後比對。"""
    worksheet = book.sheets["甲"].ws
    return [
        [(c.coordinate, c.value) for row in worksheet.iter_rows() for c in row],
        sorted(str(r) for r in worksheet.merged_cells.ranges),
    ]


# ===================================================================
# 逐 sheet ≡ 全簿
# ===================================================================


@pytest.mark.parametrize(
    "name,template,data_fn",
    SCENARIOS,
    ids=[s[0] for s in SCENARIOS],
)
def test_per_sheet_equals_whole_book(name, template, data_fn, tmp_path):
    """相同資料下，逐 sheet 渲染的輸出須與全簿渲染逐格等價。"""
    whole = tmp_path / f"{name}_whole.xlsx"
    per_sheet = tmp_path / f"{name}_per_sheet.xlsx"

    def _whole():
        with Book(template) as bk:
            bk.render(data_fn())
            bk.save(whole)

    def _per_sheet():
        with Book(template) as bk:
            data = data_fn()
            for sheet in list(bk.sheets):
                sheet.render(data)
            bk.save(per_sheet)

    render_quietly(_whole)
    render_quietly(_per_sheet)

    diffs = compare_signatures(
        workbook_signature(whole),
        workbook_signature(per_sheet),
    )
    assert not diffs, "逐 sheet 渲染與全簿渲染不一致:\n" + "\n".join(diffs)


def test_sheet_render_exported_type():
    assert templexl.SheetRenderResult is SheetRenderResult
