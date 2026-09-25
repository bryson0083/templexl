"""表格渲染樣式契約測試（對應 specs/table-render-styling）。

無 Excel Table 物件的 ``#{{表格}}`` 渲染，資料列樣式的唯一來源是模板列
樣式複製；套件不得添加模板中不存在的樣式（框線亦然）。

模板形狀取自實務案例（PL003 票差補助報表）：明細工作表僅一格標籤、
無框線的 raw data dump；彙總工作表則為畫好格線的表格列。
"""

from __future__ import annotations

import pandas as pd
import pytest
from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl import load_workbook
from openpyxl.styles import Border, PatternFill, Side

from templexl import render

from tests.golden_compare import render_quietly

THIN = Side(style="thin")
MEDIUM = Side(style="medium")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YELLOW = PatternFill(fill_type="solid", fgColor="FFFFFF00")

SIDES = ("left", "right", "top", "bottom")


def _data() -> dict:
    return {
        "df": pd.DataFrame(
            {
                "路線": ["0001", "0002", "0003"],
                "筆數": [11, 22, 33],
                "金額": [100, 200, 300],
            }
        )
    }


def _border_styles(cell) -> list:
    """儲存格四邊的框線樣式（未設定為 None）。

    邊本身可能是 ``None``（如 ``Border(bottom=...)`` 只給單邊），
    故與 ``golden_compare._border_sig`` 同樣採兩層 getattr。
    """
    return [getattr(getattr(cell.border, side), "style", None) for side in SIDES]


def _rendered_rows(worksheet, row_count: int = 3) -> tuple:
    """定位渲染結果：回傳 (表頭列號, [資料列號...])。

    以表頭文字定位而非寫死列號，使本測試只針對「樣式」把關，
    不因日後表頭/位移策略調整而假性失敗。
    """
    header_row = None
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value == "路線":
                header_row = cell.row
                break
        if header_row is not None:
            break
    assert header_row is not None, "找不到渲染後的表頭列"
    return header_row, list(range(header_row + 1, header_row + 1 + row_count))


def _render(template_path, tmp_path):
    out = tmp_path / "out.xlsx"
    render_quietly(lambda: render(str(template_path), str(out), data=_data()))
    return load_workbook(out).active


@pytest.fixture
def plain_template(tmp_path):
    """無任何框線的模板列（實務上的明細 raw data dump）。"""
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws["A1"] = "#{{ df }}"
    path = tmp_path / "plain.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def bordered_template(tmp_path):
    """模板列已畫滿格線（實務上的彙總表資料列）。"""
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws["A1"] = "#{{ df }}"
    for col in range(1, 4):
        ws.cell(row=1, column=col).border = BOX
    path = tmp_path / "bordered.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def mixed_template(tmp_path):
    """模板列各欄樣式互異——驗證逐欄複製而非整區統一覆蓋。"""
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws["A1"] = "#{{ df }}"
    ws.cell(row=1, column=1).border = BOX                      # A: 四邊 thin
    ws.cell(row=1, column=2).fill = YELLOW                     # B: 無框線、有填色
    ws.cell(row=1, column=3).border = Border(bottom=MEDIUM)    # C: 僅下框線 medium
    path = tmp_path / "mixed.xlsx"
    wb.save(path)
    return path


def test_plain_template_renders_without_borders(plain_template, tmp_path):
    """模板列無框線 → 表頭列與所有資料列皆無框線。"""
    ws = _render(plain_template, tmp_path)
    header_row, data_rows = _rendered_rows(ws)

    for row in [header_row, *data_rows]:
        for col in range(1, 4):
            cell = ws.cell(row=row, column=col)
            assert _border_styles(cell) == [None, None, None, None], (
                f"{cell.coordinate} 模板未設框線，輸出不應有框線"
            )


def test_bordered_template_preserves_borders(bordered_template, tmp_path):
    """模板列有框線 → 所有資料列保留模板框線。"""
    ws = _render(bordered_template, tmp_path)
    _, data_rows = _rendered_rows(ws)

    for row in data_rows:
        for col in range(1, 4):
            cell = ws.cell(row=row, column=col)
            assert _border_styles(cell) == ["thin"] * 4, (
                f"{cell.coordinate} 應保留模板列的四邊框線"
            )


def test_mixed_template_copies_styles_per_column(mixed_template, tmp_path):
    """模板列各欄樣式互異 → 資料列逐欄比照，不被統一覆蓋。"""
    ws = _render(mixed_template, tmp_path)
    _, data_rows = _rendered_rows(ws)

    for row in data_rows:
        assert _border_styles(ws.cell(row=row, column=1)) == ["thin"] * 4
        assert _border_styles(ws.cell(row=row, column=2)) == [None, None, None, None]
        assert ws.cell(row=row, column=2).fill.fgColor.rgb == "FFFFFF00"
        assert _border_styles(ws.cell(row=row, column=3)) == [None, None, None, "medium"]
