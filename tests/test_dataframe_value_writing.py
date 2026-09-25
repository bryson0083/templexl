"""DataFrame 資料值寫入契約測試（對應 specs/dataframe-value-writing）。

鎖住「值寫成什麼」的語意，使寫值路徑得以自由重構（迭代方式、快路徑）
而不改變輸出：缺失值呈現、原生型別保真、公式前綴中和。 

缺失值三態取自實務：DuckDB 的 NULL 整數欄經 ``.df()`` 物化為
``Int64Dtype`` + ``pd.NA``；openpyxl 不接受 ``pd.NA``，而 ``str(pd.NA)``
會在報表上留下 ``<NA>`` 字樣——兩種失敗都必須不發生。 
"""

from __future__ import annotations

from datetime import date, datetime

import pandas as pd
import pytest
from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl import load_workbook

from templexl import render

from tests.golden_compare import render_quietly

BAD_MARKERS = {"nan", "NaN", "<NA>", "NaT", "None"}


@pytest.fixture
def template(tmp_path):
    """最小單表模板：A1 標籤，無合併、無框線（走無合併快路徑）。"""
    wb = OpenpyxlWorkbook()
    wb.active["A1"] = "#{{df}}"
    path = tmp_path / "tpl.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def merged_template(tmp_path):
    """帶合併儲存格的模板——確保快路徑與原路徑語意一致。"""
    wb = OpenpyxlWorkbook()
    ws = wb.active
    ws["A1"] = "#{{df}}"
    ws.merge_cells("E1:F1")
    path = tmp_path / "tpl_merged.xlsx"
    wb.save(path)
    return path


def _render(template, tmp_path, df, **kwargs):
    out = tmp_path / "out.xlsx"
    render_quietly(lambda: render(str(template), str(out), data={"df": df}, **kwargs))
    return load_workbook(out).active


def _column(ws, col: int, first_data_row: int = 2, n: int = 3) -> list:
    return [ws.cell(row=first_data_row + i, column=col).value for i in range(n)]


# ===================================================================
# Requirement: 缺失值寫為空白儲存格
# ===================================================================


def test_float_nan_renders_blank(template, tmp_path):
    """float 欄含 NaN → 空白儲存格，其餘數值正常。"""
    df = pd.DataFrame({"amount": [1.5, float("nan"), 3.5]})
    ws = _render(template, tmp_path, df)
    assert _column(ws, 1) == [1.5, None, 3.5]


def test_nullable_int_pd_na_renders_blank(template, tmp_path):
    """Int64Dtype 含 pd.NA（DuckDB NULL 整數欄的典型產出）→ 空白，非缺失值為整數。"""
    df = pd.DataFrame({"cnt": pd.array([10, None, 30], dtype="Int64")})
    assert df["cnt"].dtype == pd.Int64Dtype()

    ws = _render(template, tmp_path, df)
    values = _column(ws, 1)
    assert values == [10, None, 30]
    assert isinstance(values[0], int) and not isinstance(values[0], bool)


def test_datetime_nat_renders_blank(template, tmp_path):
    """datetime64 含 NaT → 空白，非缺失值為日期時間型別。"""
    df = pd.DataFrame({"ts": pd.to_datetime(["2026-01-02", None, "2026-03-04"])})
    ws = _render(template, tmp_path, df)
    values = _column(ws, 1)
    assert values[1] is None
    assert values[0] == datetime(2026, 1, 2)
    assert values[2] == datetime(2026, 3, 4)


def test_no_missing_value_markers_leak_into_output(template, tmp_path):
    """任何缺失值都不得以 nan / <NA> / NaT 等字樣出現在輸出。"""
    df = pd.DataFrame(
        {
            "f": [1.0, float("nan"), 3.0],
            "i": pd.array([1, None, 3], dtype="Int64"),
            "t": pd.to_datetime(["2026-01-01", None, "2026-01-03"]),
            "s": ["a", None, "c"],
        }
    )
    ws = _render(template, tmp_path, df)
    rendered = {
        str(cell.value)
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    }
    assert not (rendered & BAD_MARKERS), f"輸出出現缺失值字樣: {rendered & BAD_MARKERS}"


def test_missing_values_blank_on_merged_path(merged_template, tmp_path):
    """帶合併儲存格的模板（非快路徑）缺失值語意相同。"""
    df = pd.DataFrame({"f": [1.0, float("nan"), 3.0]})
    ws = _render(merged_template, tmp_path, df)
    assert _column(ws, 1) == [1.0, None, 3.0]


# ===================================================================
# Requirement: 原生型別保真
# ===================================================================


def test_native_types_are_preserved(template, tmp_path):
    """int/float/bool/datetime 不被字串化。"""
    df = pd.DataFrame(
        {
            "i": [1, 2, 3],
            "f": [1.5, 2.5, 3.5],
            "b": [True, False, True],
            "d": pd.to_datetime(["2026-01-01", "2026-01-02", "2026-01-03"]),
        }
    )
    ws = _render(template, tmp_path, df)

    ints, floats, bools, dates = (_column(ws, c) for c in (1, 2, 3, 4))
    assert ints == [1, 2, 3]
    assert all(isinstance(v, int) and not isinstance(v, bool) for v in ints)
    assert floats == [1.5, 2.5, 3.5]
    assert all(isinstance(v, float) for v in floats)
    assert bools == [True, False, True]
    assert all(isinstance(v, bool) for v in bools)
    assert all(isinstance(v, (datetime, date)) for v in dates)


def test_int_column_not_upcast_by_sibling_float_column(template, tmp_path):
    """int 欄不因同列存在 float 欄而被提升為 float。

    型別保真必須只取決於該欄自身，而非「同一個 DataFrame 裡還有哪些欄」。
    （``DataFrame.iterrows()`` 會把整列統一為共同 dtype，全數值表的 int 欄
    因此被提升為 float——本測試釘住不受該實作細節影響的語意。）
    """
    df = pd.DataFrame({"cnt": [1, 2, 3], "ratio": [0.5, 0.25, 0.125]})
    ws = _render(template, tmp_path, df)

    counts = _column(ws, 1)
    assert counts == [1, 2, 3]
    assert all(isinstance(v, int) and not isinstance(v, bool) for v in counts), (
        f"int 欄應保持整數，實際型別: {[type(v).__name__ for v in counts]}"
    )
    assert _column(ws, 2) == [0.5, 0.25, 0.125]


def test_numeric_strings_stay_strings(template, tmp_path):
    """字串型別的數字（如前導零的路線代碼）不得被轉成數值。"""
    df = pd.DataFrame({"route": ["0001", "0002", "0010"]})
    ws = _render(template, tmp_path, df)
    values = _column(ws, 1)
    assert values == ["0001", "0002", "0010"]
    assert all(isinstance(v, str) for v in values)


# ===================================================================
# Requirement: 字串公式前綴中和
# ===================================================================


@pytest.mark.parametrize("payload", ["=SUM(A1)", "+1+1", "-2-2", "@cmd"])
def test_formula_prefix_neutralised(template, tmp_path, payload):
    """危險前綴字串於快路徑寫入後仍為文字型別，值原樣保留。"""
    df = pd.DataFrame({"note": [payload, "safe", payload]})
    ws = _render(template, tmp_path, df)

    for row in (2, 4):
        cell = ws.cell(row=row, column=1)
        assert cell.value == payload
        assert cell.data_type == "s", f"{cell.coordinate} 應為文字型別，實際 {cell.data_type}"


def test_formula_prefix_preserved_when_escaping_disabled(template, tmp_path):
    """escape_formulas=False 維持舊行為（= 開頭視為公式）。"""
    df = pd.DataFrame({"note": ["=1+1"]})
    ws = _render(template, tmp_path, df, escape_formulas=False)
    cell = ws.cell(row=2, column=1)
    assert cell.value == "=1+1"
    assert cell.data_type == "f"
