"""多工作表渲染的基線特徵化測試（sheet 級渲染的比對錨點）。

whole-workbook 黃金測試（test_golden.py）逐格鎖住整份輸出；本檔專門鎖住
「多 sheet」這個維度的現況行為——工作表名稱與順序、各 sheet 標籤獨立解析、
完整資料下無未解析標籤。 

用途：sheet 級渲染入口（Sheet.render）落地後，「逐 sheet 渲染」必須與此處
鎖定的「全簿渲染」行為等價；本檔先在重構前把該行為釘住。 
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from tests.fixtures import (
    TEMPLATE_NON_TABLE,
    TEMPLATE_TABLE,
    non_table_data,
    table_data,
)
from tests.golden_compare import render_quietly, workbook_signature
from tests.render_adapter import do_render

# 各 sheet 的專屬資料標記：該值只應出現在對應的工作表，
# 用以確認「同一次渲染中，每張工作表各自吃到自己的標籤資料」。
NON_TABLE_SHEET_MARKERS = [
    ("工作表6", "紅線"),
    ("工作表8", "R001"),
    ("工作表9", "65-74"),
    # 「台北市」同時落在工作表10（資料標籤）與新增的兩張浮水印工作表
    # 11、12（與工作表10 共用同一份靜態內容）。
    (["工作表10", "工作表11", "工作表12"], "台北市"),
]

SCENARIO_PARAMS = [
    ("non_table", TEMPLATE_NON_TABLE, non_table_data, 12),
    ("table", TEMPLATE_TABLE, table_data, 9),
]


@pytest.fixture(scope="module")
def rendered(tmp_path_factory):
    """兩個模板各以完整資料全簿渲染一次，供本檔各斷言共用。"""
    base = tmp_path_factory.mktemp("multisheet")
    outputs = {}
    for name, template, data_fn, _ in SCENARIO_PARAMS:
        out = base / f"{name}.xlsx"
        result = render_quietly(lambda: do_render(template, out, data_fn()))
        outputs[name] = (out, result)
    return outputs


@pytest.mark.parametrize(
    "name,template,data_fn,sheet_count",
    SCENARIO_PARAMS,
    ids=[p[0] for p in SCENARIO_PARAMS],
)
def test_sheet_names_and_order_preserved(
    name, template, data_fn, sheet_count, rendered
):
    """渲染不得新增、刪除或重排工作表。"""
    out, _ = rendered[name]
    template_names = load_workbook(template).sheetnames
    output_names = load_workbook(out).sheetnames

    assert output_names == template_names
    assert len(output_names) == sheet_count


@pytest.mark.parametrize(
    "name,template,data_fn,sheet_count",
    SCENARIO_PARAMS,
    ids=[p[0] for p in SCENARIO_PARAMS],
)
def test_no_unresolved_tags_with_complete_data(
    name, template, data_fn, sheet_count, rendered
):
    """資料齊備時，全簿所有工作表的標籤都應解析完畢（warnings 為空）。"""
    _, result = rendered[name]
    assert result.warnings == []


@pytest.mark.parametrize("sheet_name,marker", NON_TABLE_SHEET_MARKERS)
def test_each_sheet_renders_its_own_data(sheet_name, marker, rendered):
    """各工作表的專屬資料只落在自己的工作表，不外溢至其他工作表。"""
    out, _ = rendered["non_table"]
    wb = load_workbook(out)
    hits = [
        title
        for title in wb.sheetnames
        if any(cell.value == marker for row in wb[title].iter_rows() for cell in row)
    ]
    expected = sheet_name if isinstance(sheet_name, list) else [sheet_name]
    assert hits == expected


def test_every_sheet_has_content(rendered):
    """全簿渲染後每張工作表都有內容（無 sheet 被整張略過）。"""
    for name, _, _, _ in SCENARIO_PARAMS:
        out, _ = rendered[name]
        sheets = workbook_signature(out)["sheets"]
        empty = [title for title, sig in sheets.items() if not sig["cells"]]
        assert not empty, f"{name}: 以下工作表渲染後無任何內容: {empty}"
